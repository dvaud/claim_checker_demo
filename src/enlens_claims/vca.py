"""Load the Value Chain Assessment (Peyerl & van der Zwaan 2024, supplementary table T3) as gold labels.

Layout of sheet "(T3) SDGs & Targets - Analyses" (data from row 6):
  col 2 goal, 3 target, 4 description,
  cols 5-7 direct influence on production / transportation / end-use,
  cols 8-10 indirect influence on the same segments,
  col 11 temporality, 12 reciprocal interdependence, 13 justification, 14 references.
The influence cells are empty; the sign is the cell colour: red = negative, green = positive.
"""
from __future__ import annotations

import json
import re
from dataclasses import dataclass, field
from pathlib import Path

from .store import Workspace
from .taxonomy import SEGMENTS, normalise_target, targets
from .textutil import clean_doi, find_dois

SHEET = "(T3) SDGs & Targets - Analyses"
FIRST_ROW = 6
DIRECT_COLS = {"production": 5, "transportation": 6, "end_use": 7}
INDIRECT_COLS = {"production": 8, "transportation": 9, "end_use": 10}
GREEN_THEMES = {9}  # theme colour 9 is the green used for positive influences in this workbook
ORG_RE = re.compile(r"^\s*(IEA|IRENA|International|FSR|Green Hydrogen Organi|World Bank|OECD|European|EU\b|UN\b|United Nations|"
                    r"Hydrogen Council|Hydrogen Europe|Clean Hydrogen|BloombergNEF|McKinsey|Deloitte|PwC|Agora|GIZ|DNV|"
                    r"Clifford Chance|Fuel Cells and Hydrogen|Research HE|ILO|UNIDO|UNEP|UNDP|WHO|FAO|Africa)", re.I)


@dataclass
class VcaTarget:
    target: str
    goal: int
    description: str
    cells: dict[str, str] = field(default_factory=dict)  # segment -> "direct_negative" | "indirect_positive" | ...
    horizon: str = ""
    interdependence: str = ""
    justification: str = ""
    references_text: str = ""


def _fill_sign(cell) -> str | None:
    """Return 'positive', 'negative' or None from a cell's fill colour."""
    fill = cell.fill
    if fill is None or fill.fill_type in (None, "none"):
        return None
    c = fill.fgColor
    if c is None:
        return None
    if c.type == "theme":
        if c.theme in GREEN_THEMES:
            return "positive"
        return None
    if c.type == "rgb" and isinstance(c.rgb, str) and len(c.rgb) >= 6:
        rgb = c.rgb[-6:]
        try:
            r, g, b = (int(rgb[i:i + 2], 16) for i in (0, 2, 4))
        except ValueError:
            return None
        if r > 200 and g > 200 and b > 200 and abs(r - g) < 25:
            return None  # white or near-white
        if r >= g + 25 and r >= b:
            return "negative"
        if g >= r + 15 and g >= b - 10:
            return "positive"
    return None


def _horizon(text: str) -> str:
    t = (text or "").lower()
    if "short" in t:
        return "short_term"
    if "medium" in t or "long" in t:
        return "medium_long_term"
    return ""


def load_workbook_targets(xlsx: str | Path) -> list[VcaTarget]:
    import openpyxl

    wb = openpyxl.load_workbook(xlsx)
    if SHEET not in wb.sheetnames:
        raise ValueError(f"Sheet {SHEET!r} not found in {xlsx}; sheets are {wb.sheetnames}")
    ws = wb[SHEET]
    out: list[VcaTarget] = []
    for r in range(FIRST_ROW, ws.max_row + 1):
        code = normalise_target(ws.cell(r, 3).value)
        if not code:
            continue
        t = targets()[code]
        vt = VcaTarget(code, t.goal, t.description)
        for seg in SEGMENTS:
            parts = []
            d = _fill_sign(ws.cell(r, DIRECT_COLS[seg]))
            i = _fill_sign(ws.cell(r, INDIRECT_COLS[seg]))
            if d:
                parts.append(f"direct_{d}")
            if i:
                parts.append(f"indirect_{i}")
            vt.cells[seg] = "+".join(parts)
        text = lambda c: re.sub(r"\s+", " ", str(ws.cell(r, c).value or "")).strip()
        vt.horizon = _horizon(text(11))
        vt.interdependence = text(12)
        vt.justification = text(13)
        vt.references_text = str(ws.cell(r, 14).value or "").strip()
        out.append(vt)
    return out


def split_references(text: str) -> list[str]:
    """Split a numbered reference block '(1) A. ... (2) B. ...' into items."""
    if not text or text.strip() in ("0", "nan"):
        return []
    parts = re.split(r"(?:^|\s)\(\d{1,3}\)\s*", text)
    items = []
    for p in parts:
        p = re.sub(r"\s+", " ", p).strip(" .|")
        if len(p) >= 15:
            items.append(p)
    return items


def ref_key(reference: str) -> tuple[str, str | None]:
    dois = find_dois(reference)
    doi = clean_doi(dois[0]) if dois else None
    if doi:
        return "doi:" + doi, doi
    slug = re.sub(r"[^a-z0-9]", "", reference.lower())[:48]
    return "ref:" + slug, None


def ref_kind(reference: str, doi: str | None) -> str:
    if doi:
        return "doi"
    if ORG_RE.search(reference):
        return "report"
    return "no_doi"


def summary(vts: list[VcaTarget]) -> dict[str, int]:
    s = {"targets": len(vts), "justified": 0, "direct": 0, "indirect": 0, "positive": 0, "negative": 0,
         "short_term": 0, "medium_long_term": 0}
    for v in vts:
        if v.justification:
            s["justified"] += 1
        for cell in v.cells.values():
            for part in filter(None, cell.split("+")):
                kind, sign = part.split("_")
                s[kind] += 1
                s[sign] += 1
        if v.horizon:
            s[v.horizon] += 1
    return s


def _title_words(reference: str) -> set[str]:
    from .textutil import STOPWORDS, normalise

    text = re.sub(r"https?://\S+", " ", normalise(reference))
    return {w for w in re.findall(r"[a-z]{4,}", text) if w not in STOPWORDS}


def load_fixes() -> list[dict]:
    import csv
    from importlib import resources

    path = resources.files("enlens_claims").joinpath("data/reference_fixes.csv")
    with path.open(encoding="utf-8") as f:
        return list(csv.DictReader(f))


def _norm_ref(s: str) -> str:
    from .textutil import normalise

    return normalise(s).replace("ü", "u").replace("é", "e")


def merge_references(refs: dict[str, dict], fixes: list[dict] | None = None) -> dict[str, str]:
    """Apply curated fixes and merge duplicate references. Returns alias key -> canonical key.

    Fixes: 'doi' sets or removes a DOI; 'same_as' points a reference at another. Then references whose
    title words overlap strongly (Jaccard >= 0.7, at least 5 words) are merged, keeping the one with a DOI.
    """
    fixes = load_fixes() if fixes is None else fixes
    for key, e in refs.items():
        n = _norm_ref(e["reference"])
        for f in fixes:
            if f["action"] == "doi" and _norm_ref(f["match"]) in n:
                e["doi"] = None if f["value"].strip().lower() == "none" else f["value"].strip().lower()
                e["note"] = f["note"]
    alias: dict[str, str] = {}
    keys = list(refs)
    for f in fixes:
        if f["action"] != "same_as":
            continue
        src = [k for k in keys if _norm_ref(f["match"]) in _norm_ref(refs[k]["reference"])]
        dst = [k for k in keys if _norm_ref(f["value"]) in _norm_ref(refs[k]["reference"]) and k not in src]
        for s_ in src:
            if dst:
                alias[s_] = dst[0]
    words = {k: _title_words(refs[k]["reference"]) for k in keys}
    years = {k: set(re.findall(r"\b((?:19|20)\d{2})[a-z]?\b", re.sub(r"https?://\S+", " ", refs[k]["reference"])))
             for k in keys}
    for i, a in enumerate(keys):
        for b in keys[i + 1:]:
            if a in alias or b in alias:
                continue
            wa, wb = words[a], words[b]
            if len(wa) < 5 or len(wb) < 5:
                continue
            if years[a] and years[b] and not (years[a] & years[b]):
                continue  # same wording, different years: an organisation's reports of different years are different works
            if len(wa & wb) / len(wa | wb) >= 0.7:
                keep, drop = (a, b) if (refs[a]["doi"] or not refs[b]["doi"]) else (b, a)
                alias[drop] = keep
    # resolve chains and fold targets and DOIs into the canonical entry
    def root(k):
        seen = set()
        while k in alias and k not in seen:
            seen.add(k)
            k = alias[k]
        return k

    for a in list(alias):
        alias[a] = root(a)
        c = refs[alias[a]]
        for t in refs[a]["targets"]:
            if t not in c["targets"]:
                c["targets"].append(t)
        if not c["doi"] and refs[a]["doi"]:
            c["doi"] = refs[a]["doi"]
    return alias


def load_into(ws: Workspace, xlsx: str | Path) -> dict[str, int]:
    vts = load_workbook_targets(xlsx)
    refs: dict[str, dict] = {}
    with ws.connect() as cx:
        cx.execute("DELETE FROM vca_targets")
        cx.execute("DELETE FROM vca_refs")
        for v in vts:
            cx.execute(
                "INSERT INTO vca_targets VALUES (?,?,?,?,?,?,?,?,?,?)",
                (v.target, v.goal, v.description, v.cells.get("production", ""), v.cells.get("transportation", ""),
                 v.cells.get("end_use", ""), v.horizon, v.interdependence, v.justification, v.references_text),
            )
            if not v.justification:
                continue
            for item in split_references(v.references_text):
                key, doi = ref_key(item)
                e = refs.setdefault(key, {"reference": item, "doi": doi, "targets": []})
                if v.target not in e["targets"]:
                    e["targets"].append(v.target)
        alias = merge_references(refs)
        for key, e in refs.items():
            cx.execute("INSERT INTO vca_refs(ref_key, reference, doi, kind, targets_json, alias_of, note) VALUES (?,?,?,?,?,?,?)",
                       (key, e["reference"], e["doi"], ref_kind(e["reference"], e["doi"]), json.dumps(e["targets"]),
                        alias.get(key), e.get("note")))
    s = summary(vts)
    s["references"] = len(refs) - len(alias)
    s["duplicate_references_merged"] = len(alias)
    s["references_fixed"] = sum(1 for e in refs.values() if e.get("note"))
    return s
