"""Check hypotheses and claims, typed by a person, against the documents of a project.

A hypothesis is checked in these steps:

0. (optional) A model splits the hypothesis into the claims that must all hold for it to be true ("decompose"),
   so each part is searched and graded on its own. Without this step the hypothesis is one claim.
1. (optional) For each claim, a model writes search terms: synonyms, related technical words and the words a source
   would use to state the opposite. Keyword ranking misses evidence worded differently; this narrows that gap.
2. The best passages are picked by BM25 on the claim and its terms (fused with semantic search when the project
   has embeddings), at most `per_doc` from one document first so one long report cannot crowd out the others.
3. A model grades each passage that bears on the claim: supports, partially supports or refutes, with a quote.
   Quotes are checked against the passage like every other answer.
4. Each claim gets a status from its grades; a decomposed hypothesis gets one from its parts (see `aggregate`).

Everything is stored like any other task (tasks, results, verifications), so a check can be audited later.
A check started by a model outside this program (Claude in a chat, through the connector) uses `open_check`,
`submit_grades` and `more_passages` instead of a backend.
"""
from __future__ import annotations

import bisect
import csv
import inspect
import io
import json
import uuid
from typing import Any, Callable

from . import frameworks, tasks
from .search import neighbours, passage_index, rank_passages  # noqa: F401  (re-exported for callers)
from .store import Workspace, now
from .textutil import find_quote

Progress = Callable[..., None]  # progress(step, message) or progress(step, message, info)

STATUS_TEXT = {
    "supported": "Supported by your documents",
    "partially_supported": "Partly supported",
    "contested": "Mixed evidence",
    "contradicted": "Contradicted by your documents",
    "unsupported": "No evidence found",
    "uncertain": "Only uncertain evidence",
    "not_checked": "Not finished yet",
    "no_passages": "Nothing in your documents matched",
    "failed": "The check did not work",
    "waiting": "Waiting for Claude Code",
}
STATUS_HELP = {
    "supported": "At least one excerpt says what the hypothesis says, and none says the opposite.",
    "partially_supported": "No excerpt says it all, but at least one says an important part of it. Read the excerpts to "
                           "judge whether the rest follows.",
    "contested": "Some excerpts support the hypothesis and others say the opposite. Read both sides before drawing a "
                 "conclusion.",
    "contradicted": "At least one excerpt says the opposite, and none supports it.",
    "unsupported": "None of the excerpts read was about this hypothesis. That does not prove your documents never say "
                   "it: they may use other words, or the evidence may be in parts not read yet.",
    "uncertain": "The AI found excerpts that may be about this hypothesis, but it was not sure enough about any of "
                 "them for your strictness setting, so none is counted.",
    "not_checked": "The excerpts have not been read yet.",
    "no_passages": "No part of your documents uses words close to this hypothesis. Try other wording, or add documents.",
    "failed": "The AI's answers could not be used, even after three tries. Try again, or choose another AI model in "
              "Settings.",
    "waiting": "The check is prepared and waits in the queue for Claude Code, which reads the excerpts on your own "
               "Claude plan. The answer appears here as soon as Claude Code has read them.",
}
STATUS_NEXT = {  # what a reader can usefully do next
    "supported": "Read the supporting excerpts to see whether they say exactly what you mean.",
    "partially_supported": "Read the excerpts, or look for more evidence in your documents.",
    "contested": "Compare the excerpts on both sides: do they concern the same place, group or period?",
    "contradicted": "Read the contradicting excerpt: it may concern a different situation than yours.",
    "unsupported": "Look for more evidence, or search again with other words your documents might use.",
    "uncertain": "Read the excerpts and judge them yourself, or make the strictness setting more lenient.",
    "no_passages": "Rephrase the hypothesis with words your documents use, or add documents.",
    "failed": "Check again, or choose another AI model in Settings.",
    "not_checked": "",
    "waiting": "Type /enlens-run and the project's name in Claude Code (the queue panel has the command to copy).",
}
GRADES = ("supports", "partially_supports", "refutes")
CHAT_MODEL = "claude (chat)"


def _new_id(prefix: str) -> str:
    return prefix + uuid.uuid4().hex[:12]


def _adapt(progress: Progress | None) -> Callable[[str, str, dict], None]:
    """Accept both progress(step, message) and progress(step, message, info)."""
    if progress is None:
        return lambda step, msg, info: None
    try:
        params = inspect.signature(progress).parameters.values()
        three = sum(p.kind in (p.POSITIONAL_ONLY, p.POSITIONAL_OR_KEYWORD) for p in params) >= 3 or \
            any(p.kind == p.VAR_POSITIONAL for p in params)
    except (TypeError, ValueError):
        three = False
    def call(s: str, m: str, i: dict) -> None:
        try:
            progress(s, m, i) if three else progress(s, m)
        except Exception:  # a display problem must never stop a check
            pass

    return call


# ---------------------------------------------------------------------------------------------
# statuses
# ---------------------------------------------------------------------------------------------
USER_VERDICTS = ("supports", "partially_supports", "refutes", "not_relevant")


def effective(v) -> dict:
    """A verification row with a person's judgement applied: theirs replaces the AI's and always counts."""
    d = dict(v)
    uv = d.get("user_verdict")
    if uv in ("supports", "partially_supports", "refutes"):
        d.update(verdict=uv, confidence=1.0)
    elif uv == "not_relevant":
        d.update(verdict="not_enough_info", confidence=None)
    return d


def changed_by_person(v) -> bool:
    """Whether a person's judgement differs from the AI's (agreeing with the AI is a review, not a change)."""
    v = dict(v)
    uv = v.get("user_verdict")
    if not uv:
        return False
    ai = v.get("verdict") if v.get("verdict") in ("supports", "partially_supports", "refutes") else "not_relevant"
    return uv != ai


def _graded(vs, thr: float) -> tuple[list, list, list]:
    strong = [v for v in vs if v["confidence"] is not None and v["confidence"] >= thr]
    return ([v for v in strong if v["verdict"] == "supports"], [v for v in strong if v["verdict"] == "partially_supports"],
            [v for v in strong if v["verdict"] == "refutes"])


def claim_status(vs, thr: float) -> str:
    """A claim's status from its grades at or above the threshold. Support and contradiction together: contested.

    Grades that exist but all fall below the threshold give "uncertain" (not "nothing found": something was found).
    Rows carrying a person's judgement must be passed through `effective` first.
    """
    if not vs:
        return "not_checked"
    sup, part, ref = _graded(vs, thr)
    if (sup or part) and ref:
        return "contested"
    if sup or part or ref:
        return "supported" if sup else "partially_supported" if part else "contradicted"
    if any(v["verdict"] in GRADES for v in vs):
        return "uncertain"
    return "unsupported"


def balance(vs, thr: float) -> dict:
    """Documents on each side, and a lean from -1 (only contradicted) to +1 (only supported).

    Each document counts once per side, with its strongest grade's confidence; a partial grade counts half. A document
    with excerpts on both sides is counted on both (`docs_both`); `docs_counted` counts each document once.
    """
    sup: dict[str, float] = {}
    part: dict[str, float] = {}
    ref: dict[str, float] = {}
    uncounted = 0
    for v in vs:
        c = v["confidence"]
        if v["verdict"] in GRADES and (c is None or c < thr):
            uncounted += 1
        if c is None or c < thr:
            continue
        side = {"supports": sup, "partially_supports": part, "refutes": ref}.get(v["verdict"])
        if side is not None:
            side[v["doc_id"]] = max(side.get(v["doc_id"], 0.0), float(c))
    pos = sum(sup.values()) + 0.5 * sum(c for d, c in part.items() if d not in sup)
    neg = sum(ref.values())
    lean = round((pos - neg) / (pos + neg), 2) if pos + neg else None
    positive = set(sup) | set(part)
    return {"docs_supporting": len(sup), "docs_partly": len(set(part) - set(sup)), "docs_refuting": len(ref),
            "docs_both": len(positive & set(ref)), "docs_counted": len(positive | set(ref)),
            "excerpts_uncounted": uncounted, "lean": lean}


def who_says_what(items: list[dict]) -> list[dict]:
    """The documents behind an answer, by side, by the rule of `balance`: a document is listed as supporting if an
    excerpt that counts supports, as partly supporting only if none supports, and as contradicting if one contradicts
    (so a document can be on both sides). Each entry lists the locations of its excerpts that count on that side."""
    docs: dict[str, dict] = {}
    for x in items:
        if not x.get("counted") or x.get("verdict") not in GRADES:
            continue
        d = docs.setdefault(x["doc_id"], {"doc_id": x["doc_id"], "title": x["title"], "supports": [],
                                          "partially_supports": [], "refutes": []})
        d[x["verdict"]].append(x.get("location") or "")
    out = []
    for d in docs.values():
        sides = [("supports", d["supports"])] if d["supports"] else (
            [("partly", d["partially_supports"])] if d["partially_supports"] else [])
        if d["refutes"]:
            sides.append(("refutes", d["refutes"]))
        for side, locs in sides:
            out.append({"doc_id": d["doc_id"], "title": d["title"], "side": side, "excerpts": len(locs),
                        "locations": list(dict.fromkeys(l for l in locs if l)), "both_sides": len(sides) > 1 and
                        any(sd == "refutes" for sd, _ in sides)})
    order = {"supports": 0, "partly": 1, "refutes": 2}
    return sorted(out, key=lambda e: (order[e["side"]], -e["excerpts"], e["title"]))


def aggregate(statuses: list[str]) -> str:
    """The status of a hypothesis from the statuses of its parts (which must all hold for it to hold).

    - a part contradicted: contradicted (a hypothesis whose necessary part is contradicted does not hold, even if
      other parts are supported; "contested" is kept for sources that disagree about the same claim)
    - otherwise a part contested: contested
    - every part supported: supported
    - some parts (partly) supported: partially_supported
    - only uncertain grades: uncertain
    - nothing found for any part: unsupported (or no_passages when no part matched any passage)
    Parts not graded yet, or whose grading failed, are left out; when no part is left, so is the hypothesis.
    """
    if "waiting" in statuses:
        return "waiting"  # a check by queue not finished yet (0.8)
    decided = [s for s in statuses if s not in ("failed", "not_checked")]
    if not decided:
        return "failed" if "failed" in statuses else "not_checked"
    if "contradicted" in decided:
        return "contradicted"
    if "contested" in decided:
        return "contested"
    if all(s == "supported" for s in decided) and len(decided) == len(statuses):
        return "supported"
    positive = any(s in ("supported", "partially_supported") for s in decided)
    if positive:
        return "partially_supported"
    if "uncertain" in decided:
        return "uncertain"
    if all(s == "no_passages" for s in decided):
        return "no_passages"
    return "unsupported"


def _queued_open(r) -> bool:
    """A round waiting in the queue for Claude Code (0.8)."""
    try:
        return bool(r["queued"]) and r["status"] in ("pending", "retry")
    except (IndexError, KeyError):
        return False


def _status_of_rounds(round_statuses: list[str], vs, thr: float, stage: str | None = None,
                      waiting: bool = False) -> str:
    if stage in ("terms", "rank", "judge") or waiting:
        return "waiting"
    if not round_statuses:
        return "no_passages"
    if "done" in round_statuses:
        return claim_status(vs, thr)
    return "failed" if round_statuses[-1] in ("failed", "retry") else "not_checked"


# ---------------------------------------------------------------------------------------------
# creating checks
# ---------------------------------------------------------------------------------------------
def _clean(text: str) -> str:
    text = " ".join((text or "").split())
    if len(text) < 10:
        raise ValueError("Write the claim as a sentence (at least 10 characters).")
    if len(text) > 2000:
        raise ValueError("Keep the hypothesis under 2,000 characters; split a longer text into several hypotheses.")
    return text


def new_hypothesis(ws: Workspace, text: str, framework: str | None = None, code: str | None = None,
                   options: dict | None = None, batch_id: str | None = None, hyp_id: str | None = None) -> str:
    text = _clean(text)
    fw, code, _desc, _label = frameworks.resolve(ws, framework, code)
    hid = hyp_id or _new_id("h")
    with ws.connect() as cx:
        cx.execute("INSERT INTO hypotheses(hyp_id, text, framework, code, options_json, decomposed, created_at, batch_id) "
                   "VALUES (?,?,?,?,?,0,?,?)", (hid, text, fw, code, json.dumps(options or {}), now(), batch_id))
    return hid


def new_claim(ws: Workspace, text: str, target: str | None = None, options: dict | None = None,
              hyp_id: str | None = None, part: int = 0, role: str | None = None, framework: str | None = None,
              claim_id: str | None = None) -> str:
    """One claim to check. Without `hyp_id` it is a hypothesis of its own (same id), as in 0.4."""
    text = _clean(text)
    fw, code, _desc, _label = frameworks.resolve(ws, framework or ("sdg" if target else None), target)
    cid = claim_id or _new_id("u")
    with ws.connect() as cx:
        if hyp_id is None:
            cx.execute("INSERT INTO hypotheses(hyp_id, text, framework, code, options_json, decomposed, created_at) "
                       "VALUES (?,?,?,?,?,0,?)", (cid, text, fw, code, json.dumps(options or {}), now()))
            hyp_id = cid
        cx.execute("INSERT INTO user_claims(claim_id, text, target, options_json, created_at, hyp_id, part, role, framework) "
                   "VALUES (?,?,?,?,?,?,?,?,?)", (cid, text, code, json.dumps(options or {}), now(), hyp_id, part, role, fw))
    return cid


def _add_task(ws: Workspace, task_id: str, kind: str, model: str, payload: dict) -> None:
    with ws.connect() as cx:
        tasks._insert_task(cx, task_id, kind, model, payload)


def _used(backend, model: str) -> str:
    return "baseline" if getattr(backend, "name", "") == "baseline" else model


def _stopper(should_stop: Callable[[], bool] | None) -> Callable[[], None]:
    def check() -> None:
        if should_stop and should_stop():
            from .llm import BackendError

            raise BackendError("Stopped.")
    return check


MAX_PARTS = 6


def clean_terms(terms) -> list[str]:
    """Search words as a person or a model gives them: trimmed, at most 80 characters each, no repeats, at most 40."""
    if isinstance(terms, str):
        terms = terms.replace(";", ",").split(",")
    out: list[str] = []
    for t in terms or []:
        t = " ".join(str(t).split()).strip("\"'")
        if t and len(t) <= 80 and t.lower() not in {x.lower() for x in out}:
            out.append(t)
    return out[:40]


def _given_parts(parts, text: str) -> list[dict]:
    """Parts a person reviewed before the check ("Look before checking"): each a sentence of its own."""
    out = []
    for i, p in enumerate(parts or []):
        t = p.get("text") if isinstance(p, dict) else p
        t = " ".join(str(t or "").split())
        if not t:
            continue
        try:
            t = _clean(t)
        except ValueError:
            raise ValueError(f"Part {i + 1} is too short: write each part as a sentence (at least 10 characters).") from None
        role = p.get("role") if isinstance(p, dict) else None
        from .schema import PART_ROLES

        out.append({"text": t, "role": role if role in PART_ROLES else None})
    if not out:
        raise ValueError("Give at least one part to check.")
    if len(out) > MAX_PARTS:
        raise ValueError(f"At most {MAX_PARTS} parts.")
    if len(out) == 1:
        return [{"text": out[0]["text"], "role": None, "part": 0}]
    return [{**p, "part": i} for i, p in enumerate(out, 1)]


def check_hypothesis(ws: Workspace, text: str, backend, model: str, framework: str | None = None,
                     code: str | None = None, k: int = 8, per_doc: int = 2, expand: bool = True,
                     decompose: bool = False, doc_ids: list[str] | None = None, progress: Progress | None = None,
                     embedder=None, batch_id: str | None = None, should_stop: Callable[[], bool] | None = None,
                     parts: list | None = None, terms: list | None = None, edited: dict | None = None,
                     second_pass: bool = False) -> str:
    """Run a full check and return the hypothesis id; read the outcome with `hypothesis_result`.

    `parts` (sentences, or {"text", "role"}) are the parts to check, as a person saw and edited them with
    `prepare_check`: the model does not split the hypothesis then. `terms` (one list of search words per part; a
    flat list for a hypothesis checked as one claim) replace the model's extra search words; an empty list
    searches with the hypothesis's own words only. `edited` ({"parts": bool, "terms": bool}) records whether the
    person changed what the AI proposed. `second_pass` sorts the best candidates again (see search.rerank).

    Backend errors propagate, and a check stopped by one (or by `should_stop`, checked before every model call)
    leaves nothing behind.
    """
    say = _adapt(progress)
    stop = _stopper(should_stop)
    if not ws.scalar("SELECT COUNT(*) FROM passages WHERE section='body'"):
        raise ValueError("This project has no readable documents yet. Add sources first.")
    text = _clean(text)
    fw, code, _desc, _label = frameworks.resolve(ws, framework, code)
    given = _given_parts(parts, text) if parts is not None else None
    if terms is not None and terms and not isinstance(terms[0], (list, tuple)):
        terms = [terms]  # one flat list: the words of a hypothesis checked as one claim
    given_terms = [clean_terms(t) if t is not None else None for t in terms] if terms is not None else None
    used = _used(backend, model)
    options = {"k": k, "per_doc": per_doc, "expand": expand, "decompose": decompose, "doc_ids": doc_ids or [],
               "backend": getattr(backend, "name", ""), "model": used,
               "search": "keywords+meaning" if embedder is not None else "keywords",
               "embed_model": getattr(embedder, "model", None), "second_pass": bool(second_pass)}
    if given is not None or given_terms is not None:  # "Look before checking": what the person saw and chose
        options["prepared"] = {"parts": given is not None, "terms": given_terms is not None,
                               "edited_parts": bool((edited or {}).get("parts")),
                               "edited_terms": bool((edited or {}).get("terms"))}
        if given is not None:
            options["decompose"] = len(given) > 1
    hid = _new_id("h")
    new_hypothesis(ws, text, fw, code, options, batch_id, hyp_id=hid)
    try:
        plan = [{"text": text, "role": None, "part": 0}]
        if given is not None:
            plan = given
            if len(plan) > 1:
                with ws.connect() as cx:
                    cx.execute("UPDATE hypotheses SET parts_json=?, decomposed=1 WHERE hyp_id=?",
                               (json.dumps([{"text": p["text"], "role": p["role"]} for p in plan], ensure_ascii=False), hid))
            say("decompose_done", f"Checked in the {len(plan)} parts you chose" if len(plan) > 1
                else "Checked as written, as you chose", {"parts": [p["text"] for p in plan]})
        elif decompose:
            stop()
            say("decompose", "Splitting your hypothesis into parts that can be checked one by one", {})
            got = _decompose(ws, hid, text, backend, model, used)
            if len(got) > 1:
                plan = [{"text": p["text"], "role": p.get("role"), "part": i} for i, p in enumerate(got, 1)]
                say("decompose_done", f"Split into {len(plan)} parts", {"parts": [p["text"] for p in plan]})
            else:
                say("decompose_done", "It is already one simple statement, so it is checked as written" if got is not None
                    else "It could not be split, so it is checked as written", {"parts": [text]})
        for i, p in enumerate(plan):
            cid = new_claim(ws, p["text"], code, options, hyp_id=hid, part=p["part"], role=p["role"], framework=fw,
                            claim_id=hid if p["part"] == 0 else None)
            info = {"part": p["part"], "parts": len(plan), "claim": p["text"]} if len(plan) > 1 else {}
            mine = given_terms[i] if given_terms is not None and i < len(given_terms) else None
            _run_check(ws, cid, backend, model, used, k, per_doc, expand, doc_ids,
                       lambda s, m, extra=None, _i=info: say(s, m, {**_i, **(extra or {})}), embedder,
                       context=text if p["part"] else None, stop=stop, given_terms=mine, second_pass=second_pass)
        stop()
        with ws.connect() as cx:
            cx.execute("UPDATE hypotheses SET model=? WHERE hyp_id=?", (used, hid))
    except BaseException:
        delete_hypothesis(ws, hid)  # a check stopped by a connection or key problem should not linger in the history
        raise
    return hid


def _decompose(ws: Workspace, hid: str, text: str, backend, model: str, used: str) -> list[dict] | None:
    """The parts the model splits a hypothesis into ([] when it gives one part); None when its answers failed."""
    tid = f"decompose:{hid}"
    _add_task(ws, tid, "decompose", used, {"hyp_id": hid, "hypothesis": text})
    ok, _errs = tasks.run_one(ws, tid, backend, model)
    if not ok:
        return None
    got = json.loads(ws.scalar("SELECT parts_json FROM hypotheses WHERE hyp_id=?", (hid,)) or "[]")
    return got if len(got) > 1 else []


def _expand(ws: Workspace, cid: str, claim: str, backend, model: str, used: str, context: str | None) -> list[str]:
    payload = {"claim_id": cid, "claim": claim}
    if context:
        payload["context"] = context
    _add_task(ws, f"expand:{cid}", "expand", used, payload)
    ok, _errs = tasks.run_one(ws, f"expand:{cid}", backend, model)
    return json.loads(ws.scalar("SELECT terms_json FROM user_claims WHERE claim_id=?", (cid,)) or "[]") if ok else []


def prepare_check(ws: Workspace, text: str, backend, model: str, decompose: bool = False, expand: bool = True,
                  framework: str | None = None, code: str | None = None, progress: Progress | None = None,
                  should_stop: Callable[[], bool] | None = None) -> dict:
    """What a check would do before it reads anything ("Look before checking"): the parts the model splits the
    hypothesis into and the extra search words it proposes for each, for a person to edit and pass to
    `check_hypothesis` as `parts` and `terms`. Nothing is kept in the project.

    Returns {"parts": [{"text", "role"}], "terms": [[...] per part], "split": bool}.
    """
    say = _adapt(progress)
    stop = _stopper(should_stop)
    text = _clean(text)
    frameworks.resolve(ws, framework, code)
    used = _used(backend, model)
    for (old,) in ws.query("SELECT hyp_id FROM hypotheses WHERE hyp_id LIKE 'prep%'"):
        delete_hypothesis(ws, old)  # left by an earlier one that was interrupted
    hid = _new_id("prep")
    new_hypothesis(ws, text, None, None, {"prepare": True}, None, hyp_id=hid)
    try:
        parts = [{"text": text, "role": None}]
        split = False
        if decompose:
            stop()
            say("decompose", "Splitting your hypothesis into parts that can be checked one by one", {})
            got = _decompose(ws, hid, text, backend, model, used)
            if got:
                parts, split = [{"text": p["text"], "role": p.get("role")} for p in got], True
            say("decompose_done", f"Split into {len(parts)} parts" if split else "It is one simple statement",
                {"parts": [p["text"] for p in parts]})
        terms: list[list[str]] = []
        for i, p in enumerate(parts, 1):
            if not expand:
                terms.append([])
                continue
            stop()
            info = {"part": i, "parts": len(parts), "claim": p["text"]} if split else {}
            say("terms", "Thinking of other words your documents might use", info)
            cid = new_claim(ws, p["text"], None, {"prepare": True}, hyp_id=hid, part=i if split else 0)
            got = _expand(ws, cid, p["text"], backend, model, used, text if split else None)
            new = [t for t in got if t.lower() not in p["text"].lower()]
            terms.append(new)
            say("terms_done", f"Proposed: {', '.join(new[:8])}{'…' if len(new) > 8 else ''}" if new
                else "No other words proposed", {**info, "terms": got, "new_terms": new})
        return {"parts": parts, "terms": terms, "split": split}
    finally:
        delete_hypothesis(ws, hid)


# ---------------------------------------------------------------------------------------------
# 0.8: a check by queue. The search and the choice of excerpts run here at once; only the model's steps wait in the
# queue (the split into parts, the search words, the judging), for Claude Code. Each step is taken whenever the
# answer before it arrives (`advance_queued`), in whichever program stores it: the app, or the connector Claude
# Code talks to. Nothing is lost when the app is closed meanwhile.
# ---------------------------------------------------------------------------------------------
def queue_check(ws: Workspace, text: str, router, framework: str | None = None, code: str | None = None, k: int = 8,
                per_doc: int = 2, expand: bool = True, decompose: bool = False, doc_ids: list[str] | None = None,
                embed_model: str | None = None, batch_id: str | None = None, parts: list | None = None,
                terms: list | None = None, edited: dict | None = None, second_pass: bool = False) -> str:
    """Prepare a check whose model steps Claude Code answers through the queue; returns the hypothesis id. Its
    status is "waiting" until the excerpts are judged; read it with `hypothesis_result` as any other check."""
    if not ws.scalar("SELECT COUNT(*) FROM passages WHERE section='body'"):
        raise ValueError("This project has no readable documents yet. Add sources first.")
    text = _clean(text)
    fw, code, _desc, _label = frameworks.resolve(ws, framework, code)
    given = _given_parts(parts, text) if parts is not None else None
    if terms is not None and terms and not isinstance(terms[0], (list, tuple)):
        terms = [terms]
    given_terms = [clean_terms(t) if t is not None else None for t in terms] if terms is not None else None
    model = (router.model("check") if router is not None else None) or "opus"
    options = {"k": k, "per_doc": per_doc, "expand": expand, "decompose": decompose, "doc_ids": doc_ids or [],
               "backend": "claude_code", "model": model, "queue": True,
               "search": "keywords+meaning" if embed_model else "keywords", "embed_model": embed_model,
               "second_pass": bool(second_pass)}
    if given is not None or given_terms is not None:
        options["prepared"] = {"parts": given is not None, "terms": given_terms is not None,
                               "edited_parts": bool((edited or {}).get("parts")),
                               "edited_terms": bool((edited or {}).get("terms"))}
        if given is not None:
            options["decompose"] = len(given) > 1
    if given_terms is not None:
        options["queue_terms"] = given_terms
    hid = _new_id("h")
    new_hypothesis(ws, text, fw, code, options, batch_id, hyp_id=hid)
    with ws.connect() as cx:
        cx.execute("UPDATE hypotheses SET queue='waiting', model=? WHERE hyp_id=?", (model, hid))
        if given is not None and len(given) > 1:
            cx.execute("UPDATE hypotheses SET parts_json=?, decomposed=1 WHERE hyp_id=?",
                       (json.dumps([{"text": p["text"], "role": p["role"]} for p in given], ensure_ascii=False), hid))
    if given is None and decompose:
        with tasks.queueing():
            _add_task(ws, f"decompose:{hid}", "decompose", (router.model("decompose") if router else None) or "sonnet",
                      {"hyp_id": hid, "hypothesis": text})
    advance_queued(ws, [hid])
    return hid


def _queue_embedder(ws: Workspace, model: str | None):
    if not model:
        return None
    try:
        from . import embed
        from .home import load_settings, make_embedder

        emb = make_embedder(load_settings())
        if emb is None or emb.model != model or not embed.status(ws, model)["embedded"]:
            return None
        return emb
    except Exception:  # noqa: BLE001 - search by words alone when the embedding server is not there
        return None


def advance_queued(ws: Workspace, hyp_ids: list[str] | None = None) -> list[str]:
    """Take every next step of the checks by queue that their answers allow. Returns the ids still waiting.

    Steps per part: `terms` (search words waiting), `rank` (search and choose the excerpts, here and now), `judge`
    (the judging waits), `done`. A step is taken by one program only: the change of stage and the task it creates
    are written together, and a program that finds the stage already changed leaves it."""
    sql = "SELECT * FROM hypotheses WHERE queue='waiting'"
    args: list = []
    if hyp_ids:
        sql += f" AND hyp_id IN ({','.join('?' * len(hyp_ids))})"
        args = list(hyp_ids)
    waiting = []
    with tasks.queueing():
        for h in ws.query(sql, args):
            try:
                if _advance_one(ws, dict(h)):
                    waiting.append(h["hyp_id"])
            except (KeyError, ValueError):
                continue  # the hypothesis was deleted meanwhile
    return waiting


def _advance_one(ws: Workspace, h: dict) -> bool:
    from .routing import from_settings

    hid, text = h["hyp_id"], h["text"]
    opts = json.loads(h["options_json"] or "{}")
    try:
        from .home import load_settings

        r = from_settings(load_settings(), route="claude_code")
    except Exception:  # noqa: BLE001
        r = None
    claims = ws.query("SELECT * FROM user_claims WHERE hyp_id=? ORDER BY part, rowid", (hid,))
    if not claims:
        plan = [{"text": text, "role": None, "part": 0}]
        if opts.get("prepared", {}).get("parts") and h.get("parts_json"):
            got = json.loads(h["parts_json"])
            plan = [{"text": p["text"], "role": p.get("role"), "part": i} for i, p in enumerate(got, 1)]
        elif opts.get("decompose") and not opts.get("prepared", {}).get("parts"):
            t = ws.query("SELECT status FROM tasks WHERE task_id=?", (f"decompose:{hid}",))
            if t and t[0]["status"] in ("pending", "retry"):
                return True
            got = json.loads(ws.scalar("SELECT parts_json FROM hypotheses WHERE hyp_id=?", (hid,)) or "[]")
            if t and t[0]["status"] == "done" and len(got) > 1:
                plan = [{"text": p["text"], "role": p.get("role"), "part": i} for i, p in enumerate(got, 1)]
        given_terms = opts.get("queue_terms")
        with ws.connect() as cx:
            if cx.execute("SELECT 1 FROM user_claims WHERE hyp_id=?", (hid,)).fetchone():
                return True  # another program made the parts just now
            for i, p in enumerate(plan):
                cid = new_claim(ws, p["text"], h["code"], opts, hyp_id=hid, part=p["part"], role=p["role"],
                                framework=h["framework"], claim_id=hid if p["part"] == 0 else None)
                mine = given_terms[i] if given_terms is not None and i < len(given_terms) else None
                if mine is not None:
                    cx.execute("UPDATE user_claims SET terms_json=?, stage='rank' WHERE claim_id=?",
                               (json.dumps(mine, ensure_ascii=False), cid))
                else:
                    cx.execute("UPDATE user_claims SET stage=? WHERE claim_id=?",
                               ("terms" if opts.get("expand", True) else "rank", cid))
        claims = ws.query("SELECT * FROM user_claims WHERE hyp_id=? ORDER BY part, rowid", (hid,))
    open_ = False
    for c in claims:
        cid, stage = c["claim_id"], c["stage"]
        if stage == "terms":
            tid = f"expand:{cid}"
            st = ws.scalar("SELECT status FROM tasks WHERE task_id=?", (tid,))
            if st is None:
                payload = {"claim_id": cid, "claim": c["text"]}
                if c["part"]:
                    payload["context"] = text
                _add_task(ws, tid, "expand", (r.model("expand") if r else None) or "sonnet", payload)
                open_ = True
                continue
            if st in ("pending", "retry"):
                open_ = True
                continue
            with ws.connect() as cx:
                if st == "failed":
                    cx.execute("UPDATE user_claims SET terms_json=COALESCE(terms_json, '[]') WHERE claim_id=?", (cid,))
                cx.execute("UPDATE user_claims SET stage='rank' WHERE claim_id=? AND stage='terms'", (cid,))
            stage = "rank"
        if stage == "rank":
            terms = json.loads(ws.scalar("SELECT terms_json FROM user_claims WHERE claim_id=?", (cid,)) or "[]")
            picked = rank_passages(ws, c["text"], terms, k=int(opts.get("k", 8)), per_doc=int(opts.get("per_doc", 2)),
                                   doc_ids=opts.get("doc_ids") or None,
                                   embedder=_queue_embedder(ws, opts.get("embed_model")),
                                   second_pass=bool(opts.get("second_pass")))
            with ws.connect() as cx:
                moved = cx.execute("UPDATE user_claims SET stage=? WHERE claim_id=? AND stage='rank'",
                                   ("judge" if picked else "done", cid)).rowcount
                if moved and picked:
                    payload = _payload(ws, cid, c["text"], c["target"], c["framework"], picked, terms)
                    tasks._insert_task(cx, f"check:{cid}", "check", (r.model("check") if r else None) or "opus", payload)
            if picked:
                open_ = True
            continue
        if stage == "judge":
            st = ws.scalar("SELECT status FROM tasks WHERE task_id=?", (f"check:{cid}",))
            if st in ("pending", "retry"):
                open_ = True
                continue
            with ws.connect() as cx:
                cx.execute("UPDATE user_claims SET stage='done', model=COALESCE(model, ?) WHERE claim_id=? AND "
                           "stage='judge'", (opts.get("model"), cid))
    if not open_:
        with ws.connect() as cx:
            cx.execute("UPDATE hypotheses SET queue='done' WHERE hyp_id=? AND queue='waiting'", (hid,))
    return open_


def queue_stage(ws: Workspace, hyp_id: str) -> dict:
    """Where a check by queue stands, for its progress line: the step that waits and for how many parts."""
    h = ws.query("SELECT queue, options_json FROM hypotheses WHERE hyp_id=?", (hyp_id,))
    if not h:
        return {"state": "gone"}
    claims = ws.query("SELECT part, stage, text FROM user_claims WHERE hyp_id=? ORDER BY part", (hyp_id,))
    if h[0]["queue"] != "waiting":
        return {"state": "done", "parts": len(claims)}
    if not claims:
        return {"state": "waiting", "step": "decompose", "parts": 0}
    stages = [c["stage"] for c in claims]
    step = "terms" if "terms" in stages else "grade" if "judge" in stages else "rank"
    return {"state": "waiting", "step": step, "parts": len(claims), "stages": stages,
            "texts": [c["text"] for c in claims]}


def check_claim(ws: Workspace, text: str, backend, model: str, target: str | None = None, k: int = 8, per_doc: int = 2,
                expand: bool = True, doc_ids: list[str] | None = None, progress: Progress | None = None,
                framework: str | None = None, embedder=None) -> str:
    """Check one claim as written (no decomposition) and return its id; read the outcome with `result`."""
    return check_hypothesis(ws, text, backend, model, framework=framework or ("sdg" if target else None), code=target,
                            k=k, per_doc=per_doc, expand=expand, doc_ids=doc_ids, progress=progress, embedder=embedder)


def _run_check(ws, cid, backend, model, used, k, per_doc, expand, doc_ids, say, embedder=None, context=None,
               stop=lambda: None, given_terms: list[str] | None = None, second_pass: bool = False) -> None:
    uc = ws.query("SELECT text, target, framework FROM user_claims WHERE claim_id=?", (cid,))[0]
    claim = uc["text"]
    terms: list[str] = []
    if given_terms is not None:  # the words a person chose before the check
        terms = list(given_terms)
        with ws.connect() as cx:
            cx.execute("UPDATE user_claims SET terms_json=? WHERE claim_id=?", (json.dumps(terms, ensure_ascii=False), cid))
        new = [t for t in terms if t.lower() not in claim.lower()]
        say("terms_done", f"Also searching for the words you chose: {', '.join(new[:8])}{'…' if len(new) > 8 else ''}"
            if new else "Searching with the words of the hypothesis only, as you chose", {"terms": terms, "new_terms": new})
    elif expand:
        stop()
        say("terms", "Thinking of other words your documents might use")
        terms = _expand(ws, cid, claim, backend, model, used, context)
        new = [t for t in terms if t.lower() not in claim.lower()]  # the hypothesis's own words are searched anyway
        say("terms_done", f"Also searching for: {', '.join(new[:8])}{'…' if len(new) > 8 else ''}" if new
            else "Searching with the words of the hypothesis only", {"terms": terms, "new_terms": new})
    say("rank", "Looking through your documents for relevant excerpts")
    picked = rank_passages(ws, claim, terms, k=k, per_doc=per_doc, doc_ids=doc_ids, embedder=embedder,
                           second_pass=second_pass)
    if not picked:
        say("rank_done", "No excerpt in your documents matches this hypothesis", {"passages": 0, "documents": 0})
        return
    n_docs = len({p["doc_id"] for p in picked})
    say("rank_done", f"Found {len(picked)} excerpts in {n_docs} {'document' if n_docs == 1 else 'documents'}", {"passages": len(picked), "documents": n_docs})
    stop()
    _add_task(ws, f"check:{cid}", "check", used, _payload(ws, cid, claim, uc["target"], uc["framework"], picked, terms))
    say("grade", f"The AI is reading the {len(picked)} excerpts")
    ok, _errs = tasks.run_one(ws, f"check:{cid}", backend, model)
    graded = ws.scalar("SELECT COUNT(*) FROM verifications WHERE subject_kind='user_claim' AND subject_id=? "
                       "AND verdict != 'not_enough_info'", (cid,)) or 0
    attempts = ws.scalar("SELECT attempts FROM tasks WHERE task_id=?", (f"check:{cid}",)) or 1
    say("grade_done", (f"{graded} of {len(picked)} excerpts are about the hypothesis; every quote was checked against "
                       "the document" if ok else "The AI's answers could not be used, even after three tries")
        + (f" ({attempts} tries)" if attempts > 1 and ok else ""),
        {"graded": graded, "ok": ok})


def _payload(ws: Workspace, cid: str, claim: str, code: str | None, framework: str | None, picked: list[dict],
             terms: list[str] | None = None) -> dict:
    desc = label = None
    if code:
        _fw, code, desc, label = frameworks.resolve(ws, framework or "sdg", code)
    from .references import context_before

    passages = []
    for p in picked:
        item = {"passage_id": p["passage_id"], "doc_id": p["doc_id"], "title": p["title"], "page": p["page"],
                "text": p["text"], **({"found_by": p["found_by"]} if p.get("found_by") else {}),
                **({"matched": p["matched"]} if p.get("matched") else {})}
        ctx, ctx_start = context_before(ws, p["passage_id"])  # check-v4: the text before, for refers_to
        if ctx:
            item.update(context=ctx, context_start=ctx_start)
        passages.append(item)
    out = {"subject_kind": "user_claim", "subject_id": cid, "claim": claim, "target": code or None,
           "target_description": desc, "passages": passages}
    if label and label != "SDG target":
        out["framework_label"] = label
    if terms is not None:
        out["terms"] = list(terms)  # the words this round searched for (shown with the round's excerpts)
    return out


def continue_check(ws: Workspace, cid: str, backend, model: str, k: int = 8, terms: list[str] | None = None,
                   embedder=None, progress: Progress | None = None, should_stop: Callable[[], bool] | None = None,
                   doc_ids: list[str] | None = None) -> dict:
    """Read the next best passages for a claim (none shown before) and grade them, as a new round of its check.
    `doc_ids` limits the round to those documents (see `more_passages`)."""
    say = _adapt(progress)
    _stopper(should_stop)()
    used = _used(backend, model)
    say("rank", "Looking for further excerpts not read yet", {})
    r = more_passages(ws, cid, k=k, model=used, terms=terms, embedder=embedder, doc_ids=doc_ids)
    if not r["passages"]:
        say("rank_done", "No further excerpts match", {"passages": 0})
        return r
    say("rank_done", f"Found {len(r['passages'])} more excerpts", {"passages": len(r["passages"])})
    tid = f"check:{cid}" if r["round"] == 1 else f"check:{cid}:r{r['round']}"
    say("grade", f"The AI is reading the {len(r['passages'])} excerpts", {})
    try:
        ok, _errs = tasks.run_one(ws, tid, backend, model)
    except BaseException:
        with ws.connect() as cx:  # a round stopped by a backend problem is removed, so it can be asked again
            cx.execute("DELETE FROM results WHERE task_id=?", (tid,))
            cx.execute("DELETE FROM tasks WHERE task_id=?", (tid,))
        raise
    say("grade_done", "Read; every quote was checked against the document" if ok else "The AI's answers could not be used, even after three tries", {"ok": ok})
    return {**r, "ok": ok}


# ---------------------------------------------------------------------------------------------
# a check graded step by step by a model outside this program (Claude in a chat, through the connector)
# ---------------------------------------------------------------------------------------------
def chat_rules() -> str:
    """The grading rules of the `check` prompt, without its answer format (the connector's tool defines that)."""
    from . import prompts

    return prompts.system_prompt("check").split("\nOutput format")[0].strip()


def _check_tasks(ws: Workspace, cid: str) -> list:
    return ws.query("SELECT * FROM tasks WHERE kind='check' AND (task_id=? OR task_id LIKE ?) ORDER BY created_at, rowid",
                    (f"check:{cid}", f"check:{cid}:r%"))


def _shown(payload: dict) -> list[dict]:
    return [{"n": i, "document": p["title"] or "untitled", "page": p["page"], "text": p["text"]}
            for i, p in enumerate(payload["passages"], 1)]


def open_check(ws: Workspace, text: str, terms: list[str] | None = None, target: str | None = None, k: int = 8,
               per_doc: int = 2, doc_ids: list[str] | None = None, model: str = CHAT_MODEL,
               framework: str | None = None) -> dict:
    """Pick the passages for a claim and return them; the grades come back through `submit_grades`."""
    if not ws.scalar("SELECT COUNT(*) FROM passages WHERE section='body'"):
        raise ValueError("This project has no readable documents yet. Add sources first.")
    terms = [str(t) for t in (terms or []) if str(t).strip()][:40]
    cid = new_claim(ws, text, target, {"k": k, "per_doc": per_doc, "expand": bool(terms), "doc_ids": doc_ids or [],
                                       "backend": "chat", "model": model}, framework=framework)
    with ws.connect() as cx:
        cx.execute("UPDATE user_claims SET terms_json=?, model=? WHERE claim_id=?", (json.dumps(terms), model, cid))
        cx.execute("UPDATE hypotheses SET model=? WHERE hyp_id=?", (model, cid))
    uc = ws.query("SELECT text, target, framework FROM user_claims WHERE claim_id=?", (cid,))[0]
    picked = rank_passages(ws, uc["text"], terms, k=k, per_doc=per_doc, doc_ids=doc_ids)
    if not picked:
        return {"check_id": cid, "round": 1, "passages": []}
    payload = _payload(ws, cid, uc["text"], uc["target"], uc["framework"], picked, terms)
    _add_task(ws, f"check:{cid}", "check", model, payload)
    return {"check_id": cid, "round": 1, "passages": _shown(payload)}


def more_passages(ws: Workspace, cid: str, k: int = 8, model: str = CHAT_MODEL, terms: list[str] | None = None,
                  embedder=None, doc_ids: list[str] | None = None) -> dict:
    """The next best passages, none of them shown before, as a new round of the same check.

    With `terms`, the ranking uses these new search terms (with the claim) instead of the earlier ones, and only
    passages containing at least one of them qualify, so a person or a model can reformulate the search; the terms
    are added to the check's record. With `doc_ids`, only those documents are read (for example the documents
    added since the check); otherwise the documents the check was limited to, if any.
    """
    uc = ws.query("SELECT * FROM user_claims WHERE claim_id=?", (cid,))
    if not uc:
        raise KeyError(cid)
    uc = dict(uc[0])
    rounds = _check_tasks(ws, cid)
    if rounds and rounds[-1]["status"] not in ("done", "failed"):
        raise ValueError("Grade the passages already shown first (submit_grades), then ask for more.")
    seen = {p["passage_id"] for r in rounds for p in json.loads(r["payload_json"])["passages"]}
    opts = json.loads(uc["options_json"] or "{}")
    old_terms = json.loads(uc["terms_json"] or "[]")
    new_terms = [str(t) for t in (terms or []) if str(t).strip()][:40]
    if new_terms:
        with ws.connect() as cx:
            merged = old_terms + [t for t in new_terms if t not in old_terms]
            cx.execute("UPDATE user_claims SET terms_json=? WHERE claim_id=?", (json.dumps(merged), cid))
    picked = rank_passages(ws, uc["text"], new_terms or old_terms, k=k, per_doc=int(opts.get("per_doc", 2)),
                           doc_ids=(doc_ids if doc_ids is not None else opts.get("doc_ids")) or None, exclude=seen,
                           embedder=embedder, require_terms=bool(new_terms), second_pass=bool(opts.get("second_pass")))
    if not picked:
        return {"check_id": cid, "round": len(rounds) + 1, "passages": []}
    payload = _payload(ws, cid, uc["text"], uc["target"], uc.get("framework"), picked, new_terms or old_terms)
    n = len(rounds) + 1
    _add_task(ws, f"check:{cid}:r{n}" if rounds else f"check:{cid}", "check", model, payload)
    return {"check_id": cid, "round": n, "passages": _shown(payload)}


def submit_grades(ws: Workspace, cid: str, grades: list[dict], summary: str | None = None,
                  model: str = CHAT_MODEL) -> dict:
    """Validate and store the grades for the open round. Problems come back so they can be fixed and resent."""
    rounds = _check_tasks(ws, cid)
    if not rounds:
        if not ws.query("SELECT 1 FROM user_claims WHERE claim_id=?", (cid,)):
            raise KeyError(cid)
        raise ValueError("This check has no passages to grade.")
    t = rounds[-1]
    if t["status"] == "done":
        raise ValueError("These passages are already graded. Ask for more_passages to read further.")
    if t["status"] == "failed":
        raise ValueError("This round failed three times. Start a new check.")
    ok, errors = tasks.ingest_result(ws, t["task_id"], {"evidence": grades, "rationale": summary}, model=model)
    if not ok:
        left = tasks.MAX_ATTEMPTS - int(ws.scalar("SELECT attempts FROM tasks WHERE task_id=?", (t["task_id"],)) or 0)
        return {"accepted": False, "problems": errors, "attempts_left": max(0, left)}
    return {"accepted": True}


# ---------------------------------------------------------------------------------------------
# reading results
# ---------------------------------------------------------------------------------------------
def _threshold(ws: Workspace, threshold: float | None) -> float:
    return float(ws.config["thresholds"]["support_confidence"]) if threshold is None else float(threshold)


def certainty_word(conf: float | None) -> str:
    """The AI's certainty in words, as the app shows it."""
    if conf is None:
        return ""
    return "high" if conf >= 0.8 else "medium" if conf >= 0.6 else "low"


def _locations(ws: Workspace, ids: list[str], at: dict[str, int] | None = None) -> dict[str, dict]:
    """Where each passage is, as a reader cites it (see `locate`), keyed by passage id."""
    return locate(ws, [(pid, pid, (at or {}).get(pid)) for pid in dict.fromkeys(ids)])


def locate(ws: Workspace, items: list[tuple]) -> dict:
    """Where each of several places is, as a reader cites it: `items` are (key, passage_id, offset or None), the
    offset in document characters (a quote, a concordance hit); the result is keyed by `key`, so one passage can be
    located at several offsets in one call.

    - PDFs (and other paged documents) stored from 0.6 on: the page and the paragraph on that page, of the quote
      when `at` gives its offset ("page 2, paragraph 3"); stored earlier: the page;
    - documents without pages: the paragraph ("paragraph 6 of 8"), or for documents stored before paragraphs were
      kept, how far into the document it starts.
    `span` says what the whole excerpt covers ("page 2, paragraphs 3–4", "pages 2–3", "paragraphs 1–6 of 8").
    """
    import bisect

    out: dict = {}
    if not items:
        return out
    ids = list(dict.fromkeys(pid for _k, pid, _o in items))
    by_id = {}
    for i in range(0, len(ids), 900):  # SQLite's limit on bound parameters
        chunk = ids[i:i + 900]
        for r in ws.query("SELECT p.passage_id, p.doc_id, p.start, p.end, p.page, d.n_chars FROM passages p JOIN "
                          f"documents d ON d.doc_id=p.doc_id WHERE p.passage_id IN ({','.join('?' * len(chunk))})", chunk):
            by_id[r["passage_id"]] = r
    rows = list(by_id.values())
    doc_ids = sorted({r["doc_id"] for r in rows})
    paras: dict[str, list[int]] = {}                          # documents without pages: paragraph starts
    on_page: dict[tuple[str, int], list[int]] = {}            # paged documents: paragraph starts per page
    if doc_ids:
        for r in ws.query(f"SELECT doc_id, start, page FROM paragraphs WHERE doc_id IN ({','.join('?' * len(doc_ids))}) "
                          "ORDER BY doc_id, n", doc_ids):
            if r["page"] is None:
                paras.setdefault(r["doc_id"], []).append(r["start"])
            else:
                on_page.setdefault((r["doc_id"], r["page"]), []).append(r["start"])
    heads: dict[str, list[tuple[int, int]]] = {}
    if doc_ids:
        for r in ws.query(f"SELECT doc_id, start, end FROM headings WHERE doc_id IN ({','.join('?' * len(doc_ids))}) "
                          "ORDER BY doc_id, start", doc_ids):
            heads.setdefault(r["doc_id"], []).append((r["start"], r["end"]))

    def past_heading(doc: str, offset: int) -> int:
        """A quote that starts with a heading is cited at the paragraph the heading introduces."""
        for h0, h1 in heads.get(doc, ()):
            if h0 <= offset < h1:
                return h1 + 1
        return offset

    page_ends: dict[str, list[int]] = {}
    paged = sorted({r["doc_id"] for r in rows if r["page"]})
    if paged:
        for r in ws.query(f"SELECT doc_id, end FROM pages WHERE doc_id IN ({','.join('?' * len(paged))}) "
                          "ORDER BY doc_id, page", paged):
            page_ends.setdefault(r["doc_id"], []).append(r["end"] + 2)

    def page_at(doc: str, offset: int, default: int) -> int:
        ends = page_ends.get(doc)
        return min(len(ends), bisect.bisect_right(ends, offset) + 1) if ends else default

    def para_on(doc: str, page: int, offset: int) -> int | None:
        starts = on_page.get((doc, page))
        return max(1, bisect.bisect_right(starts, offset)) if starts else None

    for key, pid, quote_at in items:
        r = by_id.get(pid)
        if r is None:
            continue
        pct = round(100 * (r["start"] or 0) / max(1, r["n_chars"] or 1))
        starts = paras.get(r["doc_id"])
        a0, b0 = r["start"] or 0, max((r["end"] or 1) - 1, r["start"] or 0)
        span = None  # what the whole excerpt covers (the location says where its quote is)
        if r["page"]:
            first, last = page_at(r["doc_id"], a0, r["page"]), page_at(r["doc_id"], b0, r["page"])
            pa, pb = para_on(r["doc_id"], first, a0), para_on(r["doc_id"], first, b0) if last == first else None
            if last > first:
                span = f"pages {first}–{last}"
            elif pa is not None and pb is not None:
                span = f"page {first}, paragraph {pa}" if pa == pb else f"page {first}, paragraphs {pa}–{pb}"
            else:
                span = f"page {first}"
        elif starts:
            a = max(1, bisect.bisect_right(starts, a0))
            b = max(a, bisect.bisect_right(starts, b0))
            span = f"paragraph {a} of {len(starts)}" if a == b else f"paragraphs {a}–{b} of {len(starts)}"
        if quote_at is not None:
            quote_at = past_heading(r["doc_id"], quote_at)
        if r["page"]:
            off = quote_at if quote_at is not None else a0
            pg = page_at(r["doc_id"], off, r["page"])
            n = para_on(r["doc_id"], pg, off)
            where = (f"page {pg}, paragraph {n}" if n is not None and quote_at is not None
                     else span if quote_at is None else f"page {pg}")
        elif starts and quote_at is not None:  # the paragraph of the quote
            where = f"paragraph {max(1, bisect.bisect_right(starts, quote_at))} of {len(starts)}"
        elif starts:  # no quote: the paragraphs the excerpt spans
            where = span
        else:
            where = "at the start of the document" if pct < 5 else f"about {pct}% into the document"
        out[key] = {"start": r["start"], "end": r["end"], "doc_id": r["doc_id"], "position_pct": pct,
                    "location": where, "span": span or where}
    return out


def _breaks(ws: Workspace, items: list[tuple[str, str, int, int]]) -> dict[str, list[dict]]:
    """Where new paragraphs and headings begin inside each excerpt (offsets local to the excerpt), from the layout
    found when its document was added: {passage_id: [{"at", "kind": "paragraph" | "heading", "end"?}]}.
    `items` are (passage_id, doc_id, start, end) in document offsets. Documents added before 0.6 have no layout."""
    out: dict[str, list[dict]] = {}
    docs = sorted({d for _p, d, _a, _b in items})
    if not docs:
        return out
    q = ",".join("?" * len(docs))
    paras: dict[str, list[int]] = {}
    for r in ws.query(f"SELECT doc_id, start FROM paragraphs WHERE doc_id IN ({q}) ORDER BY doc_id, start", docs):
        paras.setdefault(r["doc_id"], []).append(r["start"])
    heads: dict[str, list[tuple[int, int]]] = {}
    for r in ws.query(f"SELECT doc_id, start, end FROM headings WHERE doc_id IN ({q}) ORDER BY doc_id, start", docs):
        heads.setdefault(r["doc_id"], []).append((r["start"], r["end"]))
    import bisect

    for pid, doc, a, b in items:
        if a is None or b is None:
            continue
        got: list[dict] = []
        hs = heads.get(doc, [])
        for h0, h1 in hs[bisect.bisect_left(hs, (a, -1)):]:
            if h0 >= b:
                break
            got.append({"at": h0 - a, "kind": "heading", "end": min(h1, b) - a})
        starts = paras.get(doc, [])
        taken = {g["at"] for g in got}
        for p0 in starts[bisect.bisect_right(starts, a):]:
            if p0 >= b:
                break
            if p0 - a not in taken:
                got.append({"at": p0 - a, "kind": "paragraph"})
        if got:
            out[pid] = sorted(got, key=lambda g: (g["at"], g["kind"] != "heading"))
    return out


def _round_notes(ws: Workspace, rounds: list) -> list[str | None]:
    """The AI's own short note on each round of excerpts (from its accepted answer)."""
    notes = []
    for r in rounds:
        raw = ws.scalar("SELECT raw_json FROM results WHERE task_id=? AND valid=1 ORDER BY attempt DESC LIMIT 1",
                        (r["task_id"],))
        try:
            notes.append((tasks.parse_json(raw).get("rationale") or None) if raw else None)
        except ValueError:
            notes.append(None)
    return notes


def _opens_with_pronoun(q: str | None) -> bool:
    from .references import opens_with_pronoun

    return opens_with_pronoun(q)


def _reference_locations(ws: Workspace, items: list[dict]) -> None:
    """Where each `refers_to` is: "page 4, paragraph 2" (in the excerpt) or "the paragraph before"."""
    todo = []
    for x in items:
        if x.get("refers_start") is not None:
            pid = ws.scalar("SELECT passage_id FROM passages WHERE doc_id=? AND start<=? ORDER BY start DESC LIMIT 1",
                            (x["doc_id"], x["refers_start"]))
            if pid:
                todo.append((x["n"], pid, x["refers_start"]))
    where = locate(ws, todo)
    for x in items:
        if x.get("refers_to"):
            loc = (where.get(x["n"]) or {}).get("location")
            x["refers_location"] = (loc or "") + (" (the text before the excerpt)" if x.get("refers_where") == "before"
                                                  else "")


def result(ws: Workspace, claim_id: str, threshold: float | None = None) -> dict:
    """One claim's check: status, grades with the quote's position in each passage, search terms, rounds.

    Each passage carries the AI's grade (`ai_verdict`), a person's own judgement if they gave one (`user_verdict`),
    the judgement that counts (`verdict`), and whether it counts at this threshold (`counted`).
    """
    thr = _threshold(ws, threshold)
    rows = ws.query("SELECT * FROM user_claims WHERE claim_id=?", (claim_id,))
    if not rows:
        raise KeyError(claim_id)
    uc = dict(rows[0])
    rounds = _check_tasks(ws, claim_id)
    terms_all = json.loads(uc["terms_json"] or "[]")
    out = {"claim_id": claim_id, "hyp_id": uc.get("hyp_id") or claim_id, "part": uc.get("part") or 0,
           "role": uc.get("role"), "claim": uc["text"], "target": uc["target"], "framework": uc.get("framework"),
           "created_at": uc["created_at"], "terms": terms_all, "rationale": uc["rationale"],
           "model": uc["model"], "options": json.loads(uc["options_json"] or "{}"), "threshold": thr, "passages": [],
           "errors": [], "rounds": len(rounds), "rounds_info": []}
    if uc["target"]:
        try:
            _fw, _c, desc, label = frameworks.resolve(ws, uc.get("framework") or "sdg", uc["target"])
            out["target_description"], out["target_label"] = desc, label
        except ValueError:
            pass
    raw_vs = [dict(v) for v in ws.query("SELECT * FROM verifications WHERE subject_kind='user_claim' AND subject_id=?",
                                        (claim_id,))]
    vs = [effective(v) for v in raw_vs]
    out["balance"] = balance(vs, thr)
    stage = uc.get("stage")
    if not rounds:
        out["status"] = out["status_any"] = "waiting" if stage in ("terms", "rank", "judge") else "no_passages"
        return out
    payloads = [json.loads(r["payload_json"]) for r in rounds]
    notes = _round_notes(ws, rounds)
    for r_i, (r, pl, note) in enumerate(zip(rounds, payloads, notes), 1):
        out["rounds_info"].append({"round": r_i, "status": r["status"], "passages": len(pl["passages"]),
                                   "terms": pl.get("terms", terms_all if r_i == 1 else None), "ai_note": note,
                                   # 0.8: which task it is, and whether a stronger model read it (or will)
                                   "task_id": r["task_id"], "stepped_up": bool(r["step"]) and not r["again"],
                                   "stronger_waiting": bool(r["again"]), "waiting": _queued_open(r)})
    shown = [(r_i, p) for r_i, pl in enumerate(payloads, 1) for p in pl["passages"]]
    by_passage = {v["passage_id"]: v for v in raw_vs}
    locs = _locations(ws, list({p["passage_id"] for _r, p in shown}),
                      {v["passage_id"]: (v["user_start"] if v.get("user_start") is not None else v["start"])
                       for v in raw_vs if v.get("start") is not None or v.get("user_start") is not None})
    breaks = _breaks(ws, [(pid, lc["doc_id"], lc["start"], lc["end"]) for pid, lc in locs.items()])
    seconds = _second_opinions(ws, [claim_id]).get(claim_id, {})
    for n, (r_i, p) in enumerate(shown, 1):
        v = by_passage.get(p["passage_id"])
        loc = locs.get(p["passage_id"], {})
        item = {"n": n, "round": r_i, "passage_id": p["passage_id"], "doc_id": p["doc_id"], "title": p["title"],
                "page": p["page"], "location": loc.get("location") or (f"page {p['page']}" if p["page"] else None),
                "position_pct": loc.get("position_pct"), "doc_start": loc.get("start"),
                "span": loc.get("span") or (f"page {p['page']}" if p["page"] else None), "text": p["text"],
                "found_by": p.get("found_by") or "keywords",
                "verdict": None, "ai_verdict": None, "confidence": None, "certainty": "", "user_verdict": None,
                "user_note": None, "counted": False, "quote": None, "quote_start": None, "quote_end": None,
                "in_corpus": p["passage_id"] in locs, "reviewed": False, "changed": False, "quote_by": None,
                "ai_quote": None, "breaks": breaks.get(p["passage_id"], []), "scope_note": None, "ai_history": [],
                "reread": False, "reread_info": None,
                # 0.7: the words that found it (F6), and what its quote refers to (F5)
                "matched": p.get("matched") or [], "refers_to": None, "refers_where": None, "refers_location": None,
                "refers_start": None, "reference_missing": False}
        if v:
            ai = v["verdict"] if v["verdict"] != "not_enough_info" else None
            eff = effective(v)
            item.update(ai_verdict=ai, user_verdict=v.get("user_verdict"), user_note=v.get("user_note"),
                        verdict=eff["verdict"] if eff["verdict"] != "not_enough_info" else None,
                        reviewed=bool(v.get("user_verdict")), changed=changed_by_person(v),
                        scope_note=v.get("scope_note") if ai else None)
            if v.get("reread_json"):
                item["reread_info"] = json.loads(v["reread_json"])
            hist = json.loads(v.get("ai_history_json") or "[]")
            if hist:  # the AI was asked to read this excerpt again: its earlier judgements, oldest first
                item.update(reread=True, ai_history=[{"verdict": h.get("verdict") if h.get("verdict") in GRADES else None,
                                                      "confidence": h.get("confidence") if h.get("verdict") in GRADES else None,
                                                      "quote": h.get("quote"), "scope_note": h.get("scope_note"),
                                                      "model": h.get("model"), "at": h.get("at"),
                                                      "replaced_at": h.get("replaced_at"), "why": h.get("why")}
                                                     for h in hist])
            base = loc.get("start")
            if ai:
                item.update(confidence=v["confidence"], certainty=certainty_word(v["confidence"]), quote=v["quote"],
                            ai_quote=v["quote"], quote_by="the AI" if v["quote"] else None)
                if v.get("refers_to"):
                    item.update(refers_to=v["refers_to"], refers_where=v.get("refers_where"),
                                refers_start=v.get("refers_start"))
                elif _opens_with_pronoun(v["quote"]) and (v.get("prompt_version") or "") >= "check-v4":
                    item["reference_missing"] = True  # check-v4 could have said what "it" is, and did not
                if base is not None and v["start"] is not None:
                    item.update(quote_start=v["start"] - base, quote_end=v["end"] - base)
                elif v["quote"] and v["quote"] in p["text"]:  # the document was removed later: find the quote in the copy
                    s0 = p["text"].index(v["quote"])
                    item.update(quote_start=s0, quote_end=s0 + len(v["quote"]))
            if v.get("user_quote") and v.get("user_verdict"):  # the sentence the person points to
                item.update(quote=v["user_quote"], quote_by="you")
                if base is not None and v.get("user_start") is not None:
                    item.update(quote_start=v["user_start"] - base, quote_end=v["user_end"] - base)
                elif v["user_quote"] in p["text"]:
                    s0 = p["text"].index(v["user_quote"])
                    item.update(quote_start=s0, quote_end=s0 + len(v["user_quote"]))
            item["counted"] = bool(item["verdict"]) and (eff["confidence"] or 0) >= thr
        so = seconds.get(p["passage_id"])
        item["second"] = ({**so, "agrees": so["verdict"] == item["ai_verdict"]} if so else None)
        out["passages"].append(item)
    statuses = [r["status"] for r in rounds]
    waiting = _queued_open(rounds[-1])
    out["status"] = _status_of_rounds(statuses, vs, thr, stage, waiting)
    out["status_any"] = _status_of_rounds(statuses, vs, 0.0, stage, waiting)
    out["waiting"] = out["status"] == "waiting"
    out["rounds_task_ids"] = [r["task_id"] for r in rounds]
    out["stepped_up"] = {r["task_id"]: {"step": int(r["step"] or 0), "again": bool(r["again"])} for r in rounds}
    out["round_status"] = statuses[-1]
    if "done" not in statuses or statuses[-1] in ("failed", "retry"):
        err = ws.query("SELECT errors_json FROM results WHERE task_id=? ORDER BY attempt DESC LIMIT 1", (rounds[-1]["task_id"],))
        out["errors"] = json.loads(err[0][0]) if err else []
    _reference_locations(ws, out["passages"])
    out["who"] = who_says_what(out["passages"])
    out["disagreements"] = sum(1 for x in out["passages"] if x["second"] and not x["second"]["agrees"])
    out["references_missing"] = sum(1 for x in out["passages"] if x["reference_missing"] and not x["reviewed"])
    order = {"supports": 0, "partially_supports": 1, "refutes": 2, None: 3}
    out["passages"].sort(key=lambda x: (order[x["verdict"]], not x["counted"], -(x["confidence"] or 0), x["n"]))
    return out


def hypothesis_result(ws: Workspace, hyp_id: str, threshold: float | None = None) -> dict:
    """A hypothesis with the check of each of its parts and its overall status."""
    thr = _threshold(ws, threshold)
    rows = ws.query("SELECT * FROM hypotheses WHERE hyp_id=?", (hyp_id,))
    if not rows:
        raise KeyError(hyp_id)
    h = dict(rows[0])
    part_ids = [r[0] for r in ws.query("SELECT claim_id FROM user_claims WHERE hyp_id=? ORDER BY part, rowid", (hyp_id,))]
    parts = [result(ws, cid, thr) for cid in part_ids]
    out = {"hyp_id": hyp_id, "text": h["text"], "framework": h["framework"], "code": h["code"],
           "created_at": h["created_at"], "options": json.loads(h["options_json"] or "{}"),
           "decomposed": len(parts) > 1, "model": h["model"] or (parts[0]["model"] if parts else None),
           "batch_id": h["batch_id"], "note": h["note"], "threshold": thr, "parts": parts}
    if h["code"]:
        try:
            _fw, _c, out["code_description"], out["code_label"] = frameworks.resolve(ws, h["framework"] or "sdg", h["code"])
        except ValueError:
            pass
    if not parts and h.get("queue") == "waiting":
        out["status"] = out["status_any"] = "waiting"
    elif len(parts) == 1:
        out["status"], out["status_any"] = parts[0]["status"], parts[0]["status_any"]
    else:
        out["status"] = aggregate([p["status"] for p in parts])
        out["status_any"] = aggregate([p["status_any"] for p in parts])
    vs = [effective(v) for v in ws.query(
        "SELECT verdict, confidence, doc_id, user_verdict FROM verifications WHERE subject_kind='user_claim' AND subject_id IN "
        f"({','.join('?' * len(part_ids))})", part_ids)] if part_ids else []
    out["balance"] = balance(vs, thr)
    out["documents_in_project"] = int(ws.scalar("SELECT COUNT(*) FROM documents") or 0)
    out["passages_in_project"] = int(ws.scalar("SELECT COUNT(*) FROM passages") or 0)
    out["excerpts_read"] = sum(len(p["passages"]) for p in parts)
    out["incomplete"] = any(p["status"] in ("failed", "not_checked", "waiting") for p in parts)
    out["queue"] = h.get("queue")
    out["status_text"] = STATUS_TEXT.get(out["status"], out["status"])
    out["rationale"] = parts[0]["rationale"] if len(parts) == 1 else None
    items = [x for p in parts for x in p["passages"]]
    out["corrected"] = sum(1 for x in items if x["changed"])       # judgements a person changed
    out["reviewed"] = sum(1 for x in items if x["reviewed"])       # judgements a person looked at (agreed or changed)
    out["to_review"] = sum(1 for x in items if not x["reviewed"])  # excerpts read that no person has judged yet
    out["to_review_relevant"] = sum(1 for x in items if x["ai_verdict"] and not x["reviewed"])  # of them, the AI's
    out["who"] = who_says_what(items)                                                           # grades for or against
    with_second = [x for x in items if x["second"]]
    out["second_models"] = sorted({x["second"]["model"] for x in with_second})
    out["second_read"] = len(with_second)
    out["disagreements"] = sum(1 for x in with_second if not x["second"]["agrees"])
    out["disagreements_to_review"] = sum(1 for x in with_second if not x["second"]["agrees"] and not x["reviewed"])
    out["docs_added_since"] = new_documents(ws, hyp_id, parts=parts)
    out["docs_added_since_matching"] = sum(1 for d in out["docs_added_since"] if d["matching"])
    return out


def _second_opinions(ws: Workspace, claim_ids: list[str]) -> dict[str, dict[str, dict]]:
    """The latest second-model grade of each excerpt: {claim_id: {passage_id: {model, verdict, confidence, ...}}}.
    `verdict` is None for an excerpt the second model found not about the claim."""
    out: dict[str, dict[str, dict]] = {}
    if not claim_ids:
        return out
    for r in ws.query(f"SELECT * FROM second_opinions WHERE subject_id IN ({','.join('?' * len(claim_ids))}) "
                      "ORDER BY at, rowid", claim_ids):
        graded = r["verdict"] in GRADES
        out.setdefault(r["subject_id"], {})[r["passage_id"]] = {
            "model": r["model"], "verdict": r["verdict"] if graded else None,
            "confidence": r["confidence"] if graded else None, "quote": r["quote"] if graded else None,
            "scope_note": r["scope_note"] if graded else None, "at": r["at"]}
    return out


def _shown_docs(rounds: list) -> set[str]:
    return {p["doc_id"] for r in rounds for p in json.loads(r["payload_json"])["passages"]}


def new_documents(ws: Workspace, hyp_id: str, parts: list[dict] | None = None, matching: bool = True) -> list[dict]:
    """Documents added to the project after a hypothesis was checked (or after its new documents were last read)
    that none of its rounds read from: [{doc_id, title, added_at, matching}]. `matching` says whether the document
    has an excerpt that shares a word with the hypothesis or its search words (keyword search, nothing is read).
    A check limited to chosen documents has none."""
    h = ws.query("SELECT created_at, docs_checked_at, options_json FROM hypotheses WHERE hyp_id=?", (hyp_id,))
    if not h:
        raise KeyError(hyp_id)
    if json.loads(h[0]["options_json"] or "{}").get("doc_ids"):
        return []
    cutoff = max(h[0]["created_at"] or "", h[0]["docs_checked_at"] or "")
    claims = [dict(r) for r in ws.query("SELECT claim_id, text, terms_json FROM user_claims WHERE hyp_id=?", (hyp_id,))]
    rounds = [t for c in claims for t in _check_tasks(ws, c["claim_id"])]
    seen_docs = _shown_docs(rounds)
    rows = [d for d in ws.query("SELECT doc_id, title, added_at FROM documents WHERE added_at > ? "
                                "ORDER BY added_at, title", (cutoff,)) if d["doc_id"] not in seen_docs]
    hits: set[str] = set()
    if matching and rows:
        from .search import docs_matching

        for c in claims:
            hits |= docs_matching(ws, c["text"], json.loads(c["terms_json"] or "[]"), {d["doc_id"] for d in rows})
    return [{"doc_id": d["doc_id"], "title": d["title"] or "", "added_at": d["added_at"], "matching": d["doc_id"] in hits}
            for d in rows]


def read_new_documents(ws: Workspace, hyp_id: str, backend, model: str, k: int = 8, embedder=None,
                       progress: Progress | None = None, should_stop: Callable[[], bool] | None = None) -> dict:
    """Read the documents added since a check, for every part, as a further round restricted to them (the check
    is not run again, so a person's reviews are kept). Documents without a matching excerpt are only marked as
    looked at. Returns {"hyp_id", "documents", "matching", "passages", "rounds"}."""
    say = _adapt(progress)
    new = new_documents(ws, hyp_id)
    docs = [d["doc_id"] for d in new if d["matching"]]
    total, rounds = 0, []
    if docs:
        for c in ws.query("SELECT claim_id FROM user_claims WHERE hyp_id=? ORDER BY part, rowid", (hyp_id,)):
            r = continue_check(ws, c["claim_id"], backend, model, k=k, embedder=embedder, should_stop=should_stop,
                               doc_ids=docs, progress=lambda s, m, i=None: say(s, m, i or {}))
            total += len(r["passages"])
            rounds.append({"claim_id": c["claim_id"], "round": r["round"], "passages": len(r["passages"]),
                           "ok": r.get("ok")})
    with ws.connect() as cx:
        cx.execute("UPDATE hypotheses SET docs_checked_at=? WHERE hyp_id=?", (now(), hyp_id))
    return {"hyp_id": hyp_id, "documents": len(new), "matching": len(docs), "passages": total, "rounds": rounds}


def second_opinion(ws: Workspace, hyp_ids: list[str], backend, main_model: str, second_model: str,
                   progress: Progress | None = None, should_stop: Callable[[], bool] | None = None) -> dict:
    """Let a second model grade every excerpt the checks of these hypotheses already read, with the same prompt,
    to see where two models disagree. The answers are kept apart (they never change an answer).

    Only one model is in memory at a time: the main model is unloaded first (where the server allows it: Ollama),
    and the second model when it is done. Excerpts the second model already graded are not asked again.
    Returns {"hypotheses", "rounds", "graded", "skipped", "failed", "unloaded": [...]}.
    """
    say = _adapt(progress)
    stop = _stopper(should_stop)
    if not second_model:
        raise ValueError("Choose a second model in Settings first.")
    if second_model == main_model and getattr(backend, "name", "") not in ("claude_code", "anthropic"):
        # on Claude the auditor gives the second opinion; it may run on the same model as the checker (Opus)
        raise ValueError("The second model must be another model than the main one.")
    todo = []
    for hid in hyp_ids:
        for c in ws.query("SELECT claim_id FROM user_claims WHERE hyp_id=? ORDER BY part, rowid", (hid,)):
            for t in _check_tasks(ws, c["claim_id"]):
                todo.append((hid, t))
    out = {"hypotheses": len(hyp_ids), "rounds": len(todo), "graded": 0, "skipped": 0, "failed": 0, "unloaded": []}
    if not todo:
        return out
    unload = getattr(backend, "unload", None)
    if unload and main_model and unload(main_model):  # free its memory before the second model loads
        out["unloaded"].append(main_model)
    # every task first (0.8): with Claude Code they all wait in the queue at once
    plan = []
    for hid, t in todo:
        tid = f"second:{t['task_id']}:{second_model}"
        st = ws.scalar("SELECT status FROM tasks WHERE task_id=?", (tid,))
        if st == "done":
            out["skipped"] += 1
            continue
        payload = json.loads(t["payload_json"])
        payload["second_model"] = second_model
        with ws.connect() as cx:
            if st is not None:  # asked again after a failure: three new attempts
                cx.execute("UPDATE tasks SET status='pending', attempts=0 WHERE task_id=?", (tid,))
                if tasks.queue_mode():
                    tasks.mark_queued(cx, [tid])
            else:
                tasks._insert_task(cx, tid, "second", second_model, payload)
        plan.append((tid, payload))
    try:
        for n, (tid, payload) in enumerate(plan, 1):
            stop()
            say("grade", f"The second model is reading {len(payload['passages'])} excerpts ({n} of {len(plan)})",
                {"index": n, "total": len(plan)})
            ok, _errs = tasks.run_one(ws, tid, backend, second_model)
            out["graded" if ok else "failed"] += 1
    finally:
        if unload and unload(second_model):
            out["unloaded"].append(second_model)
    say("grade_done", f"The second model read the excerpts of {len(hyp_ids)} "
        f"{'hypothesis' if len(hyp_ids) == 1 else 'hypotheses'}", {"ok": not out["failed"]})
    return out


def evidence_map(ws: Workspace, threshold: float | None = None, batch_id: str | None = None,
                 hyp_ids: list[str] | None = None) -> dict:
    """Which documents say what about which hypotheses, as a grid (computed in a few queries).

    `cells[hyp_id][doc_id]` = {"side": "supports" | "partly" | "refutes" | "both", "excerpts": n}, by the rule of
    `who_says_what` (judgements that count at this strictness; a person's always count). A hypothesis that rests on
    one document has `single_document`; a document `decides` an answer when leaving out its excerpts would change
    that answer, and is `key` when it decides at least two answers and a fifth of them.
    """
    thr = _threshold(ws, threshold)
    hs = history(ws, thr, batch_id=batch_id)
    if hyp_ids is not None:
        keep = set(hyp_ids)
        hs = [h for h in hs if h["hyp_id"] in keep]
    want = {h["hyp_id"] for h in hs}
    claims: dict[str, list[str]] = {}
    for r in ws.query("SELECT claim_id, hyp_id FROM user_claims ORDER BY part, rowid"):
        if (r["hyp_id"] or r["claim_id"]) in want:
            claims.setdefault(r["hyp_id"] or r["claim_id"], []).append(r["claim_id"])
    rounds: dict[str, list[str]] = {}
    titles: dict[str, str] = {}
    for r in ws.query("SELECT task_id, status, payload_json FROM tasks WHERE kind='check' ORDER BY created_at, rowid"):
        cid = r["task_id"].split(":")[1]
        rounds.setdefault(cid, []).append(r["status"])
        for p in json.loads(r["payload_json"])["passages"]:
            titles.setdefault(p["doc_id"], p.get("title") or "")
    in_project = {r["doc_id"]: r["title"] or "" for r in ws.query("SELECT doc_id, title FROM documents")}
    titles.update(in_project)
    ver: dict[str, list[dict]] = {}
    for v in ws.query("SELECT subject_id, verdict, confidence, doc_id, user_verdict FROM verifications "
                      "WHERE subject_kind='user_claim'"):
        ver.setdefault(v["subject_id"], []).append(effective(v))

    def status_of(hid: str, leave_out: str | None = None) -> str:
        sts = []
        for cid in claims.get(hid, []):
            vs = [v for v in ver.get(cid, []) if v["doc_id"] != leave_out]
            sts.append(_status_of_rounds(rounds.get(cid, []), vs, thr))
        return sts[0] if len(sts) == 1 else aggregate(sts)

    cells: dict[str, dict[str, dict]] = {}
    answers: dict[str, int] = {}
    decides: dict[str, int] = {}
    hyps = []
    for h in hs:
        hid = h["hyp_id"]
        per_doc: dict[str, dict[str, int]] = {}
        for cid in claims.get(hid, []):
            for v in ver.get(cid, []):
                if v["verdict"] in GRADES and v["confidence"] is not None and v["confidence"] >= thr:
                    c = per_doc.setdefault(v["doc_id"], {"supports": 0, "partially_supports": 0, "refutes": 0})
                    c[v["verdict"]] += 1
        row = {}
        for doc, c in per_doc.items():
            positive = "supports" if c["supports"] else "partly" if c["partially_supports"] else None
            side = "both" if positive and c["refutes"] else positive or "refutes"
            row[doc] = {"side": side, "excerpts": sum(c.values())}
            answers[doc] = answers.get(doc, 0) + 1
            if status_of(hid, doc) != h["status"]:
                decides[doc] = decides.get(doc, 0) + 1
        cells[hid] = row
        hyps.append({"hyp_id": hid, "text": h["text"], "status": h["status"],
                     "status_text": STATUS_TEXT.get(h["status"], h["status"]), "docs_counted": len(row),
                     "single_document": len(row) == 1, "batch_id": h["batch_id"]})
    docs = []
    for doc in sorted(answers, key=lambda d: (-answers[d], -decides.get(d, 0), titles.get(d, "").lower())):
        n = decides.get(doc, 0)
        docs.append({"doc_id": doc, "title": titles.get(doc, ""), "answers": answers[doc], "decides": n,
                     "key": n >= 2 and n >= 0.2 * max(1, len(hyps)), "in_project": doc in in_project})
    return {"threshold": thr, "hypotheses": hyps, "documents": docs, "cells": cells,
            "documents_without_evidence": len(set(in_project) - set(answers))}


MAP_SYMBOL = {"supports": "+", "partly": "(+)", "refutes": "−", "both": "±"}


def map_table(m: dict) -> list[dict]:
    """The evidence map as rows for a spreadsheet: one row per hypothesis, one column per document."""
    rows = []
    heads = []
    for d in m["documents"]:
        name = (d["title"] or d["doc_id"])[:60]
        while name in heads or name in ("Hypothesis", "Answer", "Documents it rests on"):
            name += " "
        heads.append(name)
    for h in m["hypotheses"]:
        row = {"Hypothesis": h["text"], "Answer": h["status_text"],
               "Documents it rests on": h["docs_counted"] if not h["single_document"] else "1 (only one document)"}
        for d, name in zip(m["documents"], heads):
            c = m["cells"].get(h["hyp_id"], {}).get(d["doc_id"])
            row[name] = f"{MAP_SYMBOL[c['side']]} {c['excerpts']}" if c else ""
        rows.append(row)
    if rows:
        rows.append({"Hypothesis": "Key: + supports, (+) partly supports, − contradicts, ± both; the number is how many "
                                   "excerpts of that document count", "Answer": "", "Documents it rests on": ""})
    return rows


CALIBRATION_BINS = (0.0, 0.5, 0.6, 0.7, 0.8, 0.9)


def calibration(ws: Workspace) -> dict:
    """How often a person agreed with the AI, by how sure the AI was: from the excerpts a person judged that also
    have an AI judgement for or against. Agreeing means giving the same judgement; marking an AI "supports" as not
    about the hypothesis is a disagreement. For an excerpt the AI was asked to read again, its first judgement is
    used (the later ones may follow the person's own hint).

    Returns {"bins": [{"from", "to", "judged", "agreed", "rate"}], "total", "agreed", "suggestion": {"threshold",
    "reason"} | null, "note"}. The suggestion is the lowest strictness (a bin with at least 3 judgements) from which
    on you agreed with at least 80% of the AI's judgements, once at least 20 were judged.
    """
    rows = []
    for r in ws.query("SELECT verdict, confidence, user_verdict, ai_history_json FROM verifications "
                      "WHERE subject_kind='user_claim' AND user_verdict IS NOT NULL"):
        first = json.loads(r["ai_history_json"] or "[]")
        if first:  # read again (perhaps with the person's hint): the AI's own first judgement is the one to compare
            r = {"verdict": first[0].get("verdict"), "confidence": first[0].get("confidence"),
                 "user_verdict": r["user_verdict"]}
        if r["verdict"] in GRADES and r["confidence"] is not None:
            rows.append(r)
    edges = list(CALIBRATION_BINS) + [1.0000001]
    bins = [{"from": a, "to": min(1.0, b), "judged": 0, "agreed": 0, "rate": None} for a, b in zip(edges, edges[1:])]
    for r in rows:
        c = float(r["confidence"])
        for b, hi in zip(bins, edges[1:]):
            if b["from"] <= c < hi:
                b["judged"] += 1
                b["agreed"] += int(r["user_verdict"] == r["verdict"])
                break
    for b in bins:
        b["rate"] = round(b["agreed"] / b["judged"], 3) if b["judged"] else None
    total, agreed = sum(b["judged"] for b in bins), sum(b["agreed"] for b in bins)
    out = {"bins": bins, "total": total, "agreed": agreed, "suggestion": None, "minimum": 20, "note": ""}
    if total < 20:
        out["note"] = f"Judge at least 20 of the AI's judgements to get a suggestion ({total} so far)."
        return out
    for i, b in enumerate(bins):
        if b["from"] < 0.5 or b["judged"] < 3:  # start where there is something to go on
            continue
        judged = sum(x["judged"] for x in bins[i:])
        ok = sum(x["agreed"] for x in bins[i:])
        if judged and ok / judged >= 0.8:
            out["suggestion"] = {"threshold": b["from"],
                                 "reason": f"When the AI was at least {round(b['from'] * 100)}% sure, you agreed with it in "
                                           f"{round(100 * ok / judged)}% of the {judged} excerpts you judged."}
            break
    if out["suggestion"] is None:
        out["note"] = "Even when the AI was very sure, you agreed with it less than 80% of the time: keep reviewing."
    return out


def history(ws: Workspace, threshold: float | None = None, limit: int | None = None, query: str | None = None,
            status: str | None = None, batch_id: str | None = None) -> list[dict]:
    """Every hypothesis checked in the project, newest first, with its status (computed in a few queries).

    `supports`, `partly`, `refutes` count the excerpts whose judgement counts at this threshold; `uncounted` the
    judged excerpts below it; `docs_*` count documents as on the result page.
    """
    thr = _threshold(ws, threshold)
    hyps = ws.query("SELECT * FROM hypotheses ORDER BY created_at DESC, rowid DESC")
    claims: dict[str, list] = {}
    for r in ws.query("SELECT claim_id, hyp_id, part, model, options_json, stage FROM user_claims ORDER BY part, rowid"):
        claims.setdefault(r["hyp_id"] or r["claim_id"], []).append(r)
    rounds: dict[str, list] = {}
    for r in ws.query("SELECT task_id, status, payload_json, queued FROM tasks WHERE kind='check' ORDER BY created_at, rowid"):
        cid = r["task_id"].split(":")[1]
        rounds.setdefault(cid, []).append(r)
    ver: dict[str, list] = {}
    ai_of: dict[tuple[str, str], str | None] = {}
    for v in ws.query("SELECT subject_id, passage_id, verdict, confidence, doc_id, user_verdict FROM verifications "
                      "WHERE subject_kind='user_claim'"):
        e = effective(v)
        e["_changed"] = changed_by_person(v)
        e["_ai_judged"] = v["verdict"] in ("supports", "partially_supports", "refutes")
        ver.setdefault(v["subject_id"], []).append(e)
        ai_of[(v["subject_id"], v["passage_id"])] = v["verdict"] if v["verdict"] in GRADES else None
    second: dict[tuple[str, str], str | None] = {}  # the latest second-model grade of each excerpt
    for r in ws.query("SELECT subject_id, passage_id, verdict FROM second_opinions ORDER BY at, rowid"):
        second[(r["subject_id"], r["passage_id"])] = r["verdict"] if r["verdict"] in GRADES else None
    added = sorted((r["added_at"] or "", r["doc_id"]) for r in ws.query("SELECT doc_id, added_at FROM documents"))
    q = (query or "").lower().strip()
    labels: dict[str, str] = {}

    def label(fw: str | None) -> str:
        if fw not in labels:
            try:
                f = frameworks.get(fw or "sdg", ws)
                labels[fw] = f.label if f else ""
            except ValueError:
                labels[fw] = ""
        return labels[fw]

    out = []
    for h in hyps:
        if h["hyp_id"].startswith("prep"):  # a "Look before checking" in progress, not a check
            continue
        if q and q not in (h["text"] or "").lower():
            continue
        if batch_id and h["batch_id"] != batch_id:
            continue
        parts = claims.get(h["hyp_id"], [])
        st, st_any, counts, read = [], [], {"supports": 0, "partially_supports": 0, "refutes": 0}, 0
        all_vs = []
        uncounted = corrected = reviewed = to_review = disagree = 0
        shown_docs: set[str] = set()
        limited = False
        for c in parts:
            rs = rounds.get(c["claim_id"], [])
            vs = ver.get(c["claim_id"], [])
            all_vs += vs
            sts = [r["status"] for r in rs]
            wt = bool(rs) and _queued_open(rs[-1])
            st.append(_status_of_rounds(sts, vs, thr, c["stage"], wt))
            st_any.append(_status_of_rounds(sts, vs, 0.0, c["stage"], wt))
            for r in rs:
                shown = json.loads(r["payload_json"])["passages"]
                read += len(shown)
                for p in shown:
                    shown_docs.add(p["doc_id"])
                    key = (c["claim_id"], p["passage_id"])
                    if key in second and second[key] != ai_of.get(key):
                        disagree += 1
            limited = limited or bool(json.loads(c["options_json"] or "{}").get("doc_ids"))
            for v in vs:
                corrected += v["_changed"]
                reviewed += bool(v.get("user_verdict"))
                if v["verdict"] in counts:
                    if (v["confidence"] or 0) >= thr:
                        counts[v["verdict"]] += 1
                    else:
                        uncounted += 1
        if not parts and h["queue"] != "waiting":
            continue
        if not parts:  # a check by queue waiting for its parts (0.8)
            st, st_any = ["waiting"], ["waiting"]
        to_review = max(0, read - reviewed)
        s = st[0] if len(st) == 1 else aggregate(st)
        if status and s != status:
            continue
        first = parts[0] if parts else {"claim_id": h["hyp_id"], "model": None, "options_json": h["options_json"]}
        model = h["model"] or first["model"] or json.loads(first["options_json"] or "{}").get("model", "")
        bal = balance(all_vs, thr)
        cutoff = max(h["created_at"] or "", h["docs_checked_at"] or "")
        since = 0 if limited else sum(1 for at, d in added[bisect.bisect_right(added, (cutoff, "\uffff")):]
                                      if d not in shown_docs)
        out.append({"hyp_id": h["hyp_id"], "claim_id": first["claim_id"], "created_at": h["created_at"],
                    "claim": h["text"], "text": h["text"], "framework": h["framework"] or "", "target": h["code"] or "",
                    "code": h["code"] or "", "code_label": label(h["framework"]) if h["code"] else "", "status": s,
                    "status_any": st_any[0] if len(st_any) == 1 else aggregate(st_any),
                    "supports": counts["supports"], "partly": counts["partially_supports"], "refutes": counts["refutes"],
                    "uncounted": uncounted, "corrected": corrected, "reviewed": reviewed, "to_review": to_review,
                    "docs_supporting": bal["docs_supporting"],
                    "docs_partly": bal["docs_partly"], "docs_refuting": bal["docs_refuting"],
                    "passages_read": read, "parts": len(parts), "model": model, "batch_id": h["batch_id"] or "",
                    "lean": bal["lean"], "note": h["note"] or "", "docs_added_since": since,
                    "disagreements": disagree})
        if limit and len(out) >= limit:
            break
    if batch_id:
        out.reverse()  # a set reads in the order it was typed
    return out


def export_rows(ws: Workspace, threshold: float | None = None) -> list[dict]:
    """One row per graded passage of every check (one row for a check with no grade), for a spreadsheet."""
    rows = []
    for h in history(ws, threshold):
        res = hypothesis_result(ws, h["hyp_id"], threshold)
        for part in res["parts"]:
            graded = [p for p in part["passages"] if p["verdict"] or p["user_verdict"]] or [None]
            for p in graded:
                rows.append({"hyp_id": res["hyp_id"], "claim_id": part["claim_id"], "checked_at": res["created_at"],
                             "hypothesis": res["text"], "part": part["part"], "claim": part["claim"],
                             "framework": res["framework"] or "", "target": res["code"] or "", "status": res["status"],
                             "part_status": part["status"], "model": h["model"],
                             "verdict": (p["verdict"] or "") if p else "", "ai_verdict": (p["ai_verdict"] or "") if p else "",
                             "confidence": p["confidence"] if p and p["confidence"] is not None else "",
                             "counted": (int(p["counted"]) if p else ""),
                             "user_verdict": (p["user_verdict"] or "") if p else "",
                             "user_note": (p["user_note"] or "") if p else "",
                             "quote": (p["quote"] or "") if p else "", "source": p["title"] if p else "",
                             "page": (p["page"] or "") if p else "", "location": (p["location"] or "") if p else "",
                             "found_by": p["found_by"] if p else "", "note": res["note"] or ""})
    return rows


def hypothesis_rows(res: dict) -> list[dict]:
    """The rows of one hypothesis result for CSV: every passage read, graded or not."""
    out = []
    for part in res["parts"]:
        for p in part["passages"]:
            out.append({"hypothesis": res["text"], "status": res["status"], "part": part["part"], "claim": part["claim"],
                        "part_status": part["status"], "verdict": p["verdict"] or "", "ai_verdict": p["ai_verdict"] or "",
                        "confidence": p["confidence"] if p["confidence"] is not None else "", "counted": int(p["counted"]),
                        "user_verdict": p["user_verdict"] or "", "user_note": p["user_note"] or "",
                        "quote": p["quote"] or "", "source": p["title"], "page": p["page"] or "",
                        "location": p["location"] or "", "found_by": p["found_by"], "note": res["note"] or ""})
    return out


JUDGEMENT_TEXT = {"supports": "Supports", "partially_supports": "Partly supports", "refutes": "Contradicts",
                  "not_relevant": "Not about the hypothesis", None: "Not about the hypothesis"}


def _local_time(iso: str | None) -> str:
    import datetime as _dt

    try:
        return _dt.datetime.fromisoformat(iso).astimezone().strftime("%Y-%m-%d %H:%M")
    except (TypeError, ValueError):
        return iso or ""


def _local_dt(iso: str | None):
    """A local date and time (without zone) that Excel shows as a date; the text as given if it is not one."""
    import datetime as _dt

    try:
        return _dt.datetime.fromisoformat(iso).astimezone().replace(tzinfo=None, microsecond=0)
    except (TypeError, ValueError):
        return iso or ""


def export_table(ws: Workspace, threshold: float | None = None, hyp_ids: list[str] | None = None,
                 every_excerpt: bool = True) -> list[dict]:
    """Rows for a spreadsheet a person will read: the app's own words, local times, notes and corrections.

    One row per excerpt read, the ones judged not about the hypothesis included (without `every_excerpt`, only the
    judged ones); a hypothesis without any excerpt still gets one row.
    """
    thr = _threshold(ws, threshold)
    ids = hyp_ids if hyp_ids is not None else [h["hyp_id"] for h in history(ws, thr)]
    rows = []
    results = [hypothesis_result(ws, hid, thr) for hid in ids]
    seconds = any(r["second_read"] for r in results)  # columns for a second model only when one was asked
    for res in results:
        n_parts = len(res["parts"])
        for part in res["parts"]:
            chosen = [p for p in part["passages"] if every_excerpt or p["verdict"] or p["user_verdict"]] or [None]
            for p in chosen:
                judged_by = ("you" if p["user_verdict"] else "the AI") if p else ""
                cert = certainty_word(p["confidence"]) if p and p["ai_verdict"] else ""
                rows.append({
                    "Checked on": _local_dt(res["created_at"]),
                    "Hypothesis": res["text"],
                    "Answer": STATUS_TEXT.get(res["status"], res["status"]),
                    "Part": f"{part['part']} of {n_parts}" if n_parts > 1 else "",
                    "Part checked": part["claim"] if n_parts > 1 else "",
                    "Answer for this part": STATUS_TEXT.get(part["status"], part["status"]) if n_parts > 1 else "",
                    "Judgement": JUDGEMENT_TEXT.get(p["verdict"] if p else None, "") if p else "",
                    "Judged by": judged_by,
                    "AI's own judgement": JUDGEMENT_TEXT.get(p["ai_verdict"]) if p else "",
                    "AI certainty": f"{cert} ({round(p['confidence'] * 100)}%)" if cert else "",
                    "Counted in the answer": ("yes" if p["counted"] else "no") if p else "",
                    "Quote": (p["quote"] or "") if p else "",
                    "Quote chosen by": (p["quote_by"] or "") if p and p["quote"] else "",
                    "Document": p["title"] if p else "",
                    "Where": (p["location"] or "") if p else "",
                    "Your comment on this excerpt": (p["user_note"] or "") if p else "",
                    "About another group (AI's note)": (p["scope_note"] or "") if p else "",
                    "What the quote refers to (AI)": (f"{p['refers_to']} ({p['refers_location']})" if p["refers_location"]
                                                      else p["refers_to"]) if p and p.get("refers_to") else "",
                    "Words that found this excerpt": ", ".join(p.get("matched") or []) if p else "",
                    "AI's earlier judgement (before reading it again)": _earlier(p) if p else "",
                    "Excerpt (the text the AI read)": p["text"] if p else "",
                    "Your notes on this result": res["note"] or "",
                    "Code": f"{res.get('code_label') or ''} {res['code']}".strip() if res["code"] else "",
                })
                if seconds:
                    so = p["second"] if p else None
                    rows[-1].update({
                        "Second model": so["model"] if so else "",
                        "Second model's judgement": JUDGEMENT_TEXT.get(so["verdict"]) if so else "",
                        "Second model's certainty": (f"{certainty_word(so['confidence'])} ({round(so['confidence'] * 100)}%)"
                                                     if so and so["confidence"] is not None else ""),
                        "The two models agree": ("yes" if so["agrees"] else "no") if so else "",
                    })
    return rows


def _earlier(p: dict) -> str:
    """The AI's first judgement of an excerpt it was asked to read again: "Supports (75%)"."""
    if not p.get("ai_history"):
        return ""
    h = p["ai_history"][0]
    return JUDGEMENT_TEXT.get(h["verdict"]) + (f" ({round(h['confidence'] * 100)}%)" if h.get("confidence") is not None else "")


def strictness_text(thr: float) -> str:
    """The strictness as the app words it."""
    word = "relaxed" if thr <= 0.55 else "normal" if thr <= 0.65 else "strict" if thr <= 0.8 else "very strict"
    return f"{word}: judgements count if the AI was at least {round(thr * 100)}% sure; yours always count"


def summary_table(ws: Workspace, threshold: float | None = None, hyp_ids: list[str] | None = None) -> list[dict]:
    """One row per hypothesis, in the app's words: the answer, who says what, and how much a person reviewed."""
    thr = _threshold(ws, threshold)
    rows = []
    for h in history(ws, thr):
        if hyp_ids is not None and h["hyp_id"] not in hyp_ids:
            continue
        res = hypothesis_result(ws, h["hyp_id"], thr)
        parts = res["parts"]
        rows.append({
            "Checked on": _local_dt(res["created_at"]),
            "Hypothesis": res["text"],
            "Answer": STATUS_TEXT.get(res["status"], res["status"]),
            "Answer for each part": "; ".join(f"{i}. {p['claim']}: {STATUS_TEXT.get(p['status'], p['status'])}"
                                              for i, p in enumerate(parts, 1)) if len(parts) > 1 else "",
            "Documents that support it": res["balance"]["docs_supporting"],
            "Documents that partly support it": res["balance"]["docs_partly"],
            "Documents that contradict it": res["balance"]["docs_refuting"],
            "Excerpts that count": h["supports"] + h["partly"] + h["refutes"],
            "Excerpts too uncertain to count": h["uncounted"],
            "Excerpts read": h["passages_read"],
            "Judgements you reviewed": res["reviewed"],
            "Judgements you changed": res["corrected"],
            "Excerpts not reviewed by you yet": res["to_review"],
            "Excerpts where a second model disagrees": res["disagreements"] if res["second_read"] else "",
            "Documents added since the check": len(res["docs_added_since"]),
            "Your notes on this result": res["note"] or "",
            "Code": f"{res.get('code_label') or ''} {res['code']}".strip() if res["code"] else "",
            "Strictness": strictness_text(thr),
            "AI model used (for a methods section)": res["model"] or "",
        })
    return rows


def write_xlsx(rows: list[dict], path_or_buffer, title: str = "Results", summary: list[dict] | None = None,
               extra: list[tuple[str, list[dict]]] | None = None) -> None:
    """A tidy Excel file: bold header, wrapped long columns, frozen header row, filters. With `summary`, a first sheet
    "Summary" (one row per hypothesis) comes before the sheet of excerpts; `extra` sheets (name, rows) follow it."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    from openpyxl.utils import get_column_letter

    widths = {"Hypothesis": 50, "Part checked": 40, "Quote": 70, "Document": 40, "Your comment on this excerpt": 40,
              "Your notes on this result": 40, "Excerpt (the text the AI read)": 90, "Answer for each part": 60,
              "Strictness": 40, "About another group (AI's note)": 30, "What the quote refers to (AI)": 30,
              "Words that found this excerpt": 25, "Quote": 70, "Reason": 40, "Representative passages": 90,
              "Name": 30, "Key terms": 40, "Description": 50, "Left": 45, "Right": 45}
    wb = Workbook()
    sheets = ([("Summary", summary)] if summary is not None else []) + [(title, rows)] + list(extra or [])
    for i, (name, data) in enumerate(sheets):
        sh = wb.active if i == 0 else wb.create_sheet()
        sh.title = name[:31]
        headers = list(data[0]) if data else ["Hypothesis", "Answer"]
        sh.append(headers)
        for c in sh[1]:
            c.font = Font(bold=True)
        for r in data:
            sh.append([r.get(h, "") for h in headers])
        for j, h in enumerate(headers, 1):
            sh.column_dimensions[get_column_letter(j)].width = widths.get(h, max(12, min(30, len(h) + 2)))
            if h == "Checked on":
                sh.column_dimensions[get_column_letter(j)].width = 17
                for cell in sh[get_column_letter(j)][1:]:
                    cell.number_format = "yyyy-mm-dd hh:mm"
            if h in widths:
                for cell in sh[get_column_letter(j)][1:]:
                    cell.alignment = Alignment(wrap_text=True, vertical="top")
        sh.freeze_panes = "A2"
        if data:
            sh.auto_filter.ref = sh.dimensions
    wb.save(path_or_buffer)


_KEEP = object()


def set_judgement(ws: Workspace, claim_id: str, passage_id: str, verdict: str | None, note: str | None = None,
                  quote=_KEEP) -> None:
    """A person's own judgement of an excerpt the AI read. It replaces the AI's in the answer (and always counts);
    `verdict` None removes it. The AI's judgement is kept and shown beside it. With `verdict` None and a `note`,
    only the comment is kept: a comment alone is not a judgement (it does not count as agreeing with the AI).

    `quote` is the sentence the person points to (text of this excerpt, found as quotes are); it is shown and exported
    instead of the AI's quote. Left out, an earlier one is kept; None or "" removes it.
    """
    if verdict is not None and verdict not in USER_VERDICTS:
        raise ValueError(f"The judgement must be one of {', '.join(USER_VERDICTS)} (or empty to remove yours).")
    with ws.connect() as cx:
        row = cx.execute("SELECT v.user_quote, v.user_start, v.user_end, p.text, p.start FROM verifications v "
                         "LEFT JOIN passages p ON p.passage_id=v.passage_id WHERE v.subject_kind='user_claim' "
                         "AND v.subject_id=? AND v.passage_id=?", (claim_id, passage_id)).fetchone()
        if row is None:
            raise KeyError(f"{claim_id}/{passage_id}")
        uq, us, ue = row[0], row[1], row[2]
        if not verdict:
            uq = us = ue = None
        elif quote is not _KEEP:
            q = (quote or "").strip()
            if not q:
                uq = us = ue = None
            else:
                text, base = row[3], row[4]
                if text is None:  # the document was removed later: the excerpt's copy is in the check's record
                    for t in _check_tasks(ws, claim_id):
                        for p in json.loads(t["payload_json"])["passages"]:
                            if p["passage_id"] == passage_id:
                                text, base = p["text"], None
                hit = find_quote(text or "", q)
                if not hit:
                    raise ValueError("The sentence must be text from this excerpt (select it in the excerpt).")
                uq = (text or "")[hit[0]:hit[1]]
                us, ue = (base + hit[0], base + hit[1]) if base is not None else (None, None)
        comment = (note or "").strip()[:2000] or None
        cx.execute("UPDATE verifications SET user_verdict=?, user_note=?, user_at=?, user_quote=?, user_start=?, "
                   "user_end=? WHERE subject_kind='user_claim' AND subject_id=? AND passage_id=?",
                   (verdict, comment, now() if verdict or comment else None, uq, us, ue, claim_id, passage_id))


def reread_excerpt(ws: Workspace, claim_id: str, passage_id: str, backend, model: str, note: str | None = None,
                   progress: Progress | None = None, should_stop: Callable[[], bool] | None = None) -> dict:
    """Ask the AI to read one excerpt again, alone, with the check prompt; a person's comment (`note`, or the
    comment saved with their judgement) is added as a hint. The new judgement replaces the AI's for this excerpt;
    the earlier ones are kept (`ai_history` in `result`). A person's own judgement is not changed and still counts.

    Returns {"before": {verdict, confidence}, "after": {verdict, confidence, scope_note}, "ok": bool}.
    """
    say = _adapt(progress)
    _stopper(should_stop)()
    uc = ws.query("SELECT * FROM user_claims WHERE claim_id=?", (claim_id,))
    if not uc:
        raise KeyError(claim_id)
    uc = dict(uc[0])
    shown = None
    for t in _check_tasks(ws, claim_id):
        for p in json.loads(t["payload_json"])["passages"]:
            if p["passage_id"] == passage_id:
                shown = (t, p)
    if shown is None:
        raise KeyError(f"{claim_id}/{passage_id}")
    rnd, p = shown
    v = ws.query("SELECT * FROM verifications WHERE subject_kind='user_claim' AND subject_id=? AND passage_id=?",
                 (claim_id, passage_id))
    v = dict(v[0]) if v else None
    if note is None and v:
        note = v.get("user_note")
    used = _used(backend, model)
    payload = _payload(ws, claim_id, uc["text"], uc["target"], uc.get("framework"), [p])
    if (note or "").strip():
        payload["reader_note"] = " ".join(note.split())[:1000]
    payload["previous_at"] = ws.scalar("SELECT received_at FROM results WHERE task_id=? AND valid=1 ORDER BY attempt DESC "
                                       "LIMIT 1", (rnd["task_id"],))
    n = int(ws.scalar("SELECT COUNT(*) FROM tasks WHERE kind='reread' AND task_id LIKE ?",
                      (f"reread:{claim_id}:{passage_id}:%",)) or 0) + 1
    tid = f"reread:{claim_id}:{passage_id}:{n}"
    if n > 1:
        payload["previous_at"] = ws.scalar("SELECT received_at FROM results WHERE task_id=? AND valid=1 ORDER BY attempt "
                                           "DESC LIMIT 1", (f"reread:{claim_id}:{passage_id}:{n - 1}",)) or payload["previous_at"]
    _add_task(ws, tid, "reread", used, payload)
    say("grade", "The AI is reading the excerpt again", {})
    before = {"verdict": v["verdict"] if v and v["verdict"] in GRADES else None,
              "confidence": v["confidence"] if v and v["verdict"] in GRADES else None}
    try:
        ok, _errs = tasks.run_one(ws, tid, backend, model)
    except BaseException:
        with ws.connect() as cx:  # stopped: nothing changes, and it can be asked again
            cx.execute("DELETE FROM results WHERE task_id=?", (tid,))
            cx.execute("DELETE FROM tasks WHERE task_id=?", (tid,))
        raise
    now_v = ws.query("SELECT verdict, confidence, scope_note FROM verifications WHERE subject_kind='user_claim' AND "
                     "subject_id=? AND passage_id=?", (claim_id, passage_id))
    after = ({"verdict": now_v[0]["verdict"] if now_v[0]["verdict"] in GRADES else None,
              "confidence": now_v[0]["confidence"] if now_v[0]["verdict"] in GRADES else None,
              "scope_note": now_v[0]["scope_note"],
              "rationale": json.loads(ws.scalar("SELECT reread_json FROM verifications WHERE subject_kind='user_claim' "
                                                "AND subject_id=? AND passage_id=?", (claim_id, passage_id)) or "{}"
                                      ).get("rationale")} if now_v and ok else None)
    say("grade_done", "Read again; the quote was checked against the document" if ok else
        "The AI's answers could not be used, even after three tries; its earlier judgement stays", {"ok": ok})
    return {"claim_id": claim_id, "passage_id": passage_id, "ok": ok, "before": before, "after": after}


def more_available(ws: Workspace, claim_id: str, semantic: bool = False) -> bool:
    """Whether reading further would find any excerpt not read yet: with keywords, one that shares a word with the
    claim or the words searched so far; with semantic search (`semantic`), any readable excerpt not read yet."""
    uc = ws.query("SELECT * FROM user_claims WHERE claim_id=?", (claim_id,))
    if not uc:
        raise KeyError(claim_id)
    uc = dict(uc[0])
    seen = {p["passage_id"] for r in _check_tasks(ws, claim_id) for p in json.loads(r["payload_json"])["passages"]}
    opts = json.loads(uc["options_json"] or "{}")
    docs = set(opts.get("doc_ids") or [])
    if semantic:
        return any(r["passage_id"] not in seen and (not docs or r["doc_id"] in docs) for r in passage_index(ws).rows)
    return bool(rank_passages(ws, uc["text"], json.loads(uc["terms_json"] or "[]"), k=1,
                              doc_ids=list(docs) or None, exclude=seen))


def delete_claim(ws: Workspace, claim_id: str) -> None:
    with ws.connect() as cx:
        hyp = cx.execute("SELECT hyp_id FROM user_claims WHERE claim_id=?", (claim_id,)).fetchone()
        cx.execute("DELETE FROM verifications WHERE subject_kind='user_claim' AND subject_id=?", (claim_id,))
        for tid in (f"expand:{claim_id}", f"check:{claim_id}"):
            cx.execute("DELETE FROM results WHERE task_id=?", (tid,))
            cx.execute("DELETE FROM tasks WHERE task_id=?", (tid,))
        for pat in (f"check:{claim_id}:r%", f"reread:{claim_id}:%", f"second:check:{claim_id}:%"):
            cx.execute("DELETE FROM results WHERE task_id LIKE ?", (pat,))
            cx.execute("DELETE FROM tasks WHERE task_id LIKE ?", (pat,))
        cx.execute("DELETE FROM second_opinions WHERE subject_id=?", (claim_id,))
        cx.execute("DELETE FROM user_claims WHERE claim_id=?", (claim_id,))
        hid = hyp[0] if hyp and hyp[0] else claim_id
        if not cx.execute("SELECT 1 FROM user_claims WHERE hyp_id=?", (hid,)).fetchone():
            cx.execute("DELETE FROM results WHERE task_id=?", (f"decompose:{hid}",))
            cx.execute("DELETE FROM tasks WHERE task_id=?", (f"decompose:{hid}",))
            cx.execute("DELETE FROM hypotheses WHERE hyp_id=?", (hid,))


def delete_hypothesis(ws: Workspace, hyp_id: str) -> None:
    with ws.connect() as cx:
        for (cid,) in cx.execute("SELECT claim_id FROM user_claims WHERE hyp_id=?", (hyp_id,)).fetchall():
            delete_claim(ws, cid)  # nested: one transaction
        cx.execute("DELETE FROM results WHERE task_id=?", (f"decompose:{hyp_id}",))
        cx.execute("DELETE FROM tasks WHERE task_id=?", (f"decompose:{hyp_id}",))
        cx.execute("DELETE FROM hypotheses WHERE hyp_id=?", (hyp_id,))


def set_note(ws: Workspace, hyp_id: str, note: str) -> None:
    with ws.connect() as cx:
        if not cx.execute("UPDATE hypotheses SET note=? WHERE hyp_id=?", ((note or "").strip()[:4000], hyp_id)).rowcount:
            raise KeyError(hyp_id)


# ---------------------------------------------------------------------------------------------
# frameworks and batches
# ---------------------------------------------------------------------------------------------
def suggest_targets(text: str, n: int = 5, framework: str = "sdg", ws: Workspace | None = None) -> list[dict]:
    """Codes of a framework (by default the SDG targets) whose wording is closest to the claim: a starting point."""
    fw = frameworks.get(framework, ws)
    out = fw.suggest(text, n) if fw else []
    for x in out:
        x["goal"] = int(x["group"]) if framework == "sdg" and str(x.get("group", "")).isdigit() else x.get("group")
    return out


BATCH_COLUMNS = ("hypothesis", "claim", "text", "statement")


def parse_batch(text: str) -> list[dict]:
    """Hypotheses from CSV text (a column hypothesis, claim, text or statement; optional framework and code or
    target; comma, semicolon or tab separated, as spreadsheets save them) or from plain text, one hypothesis per
    line (blank lines and lines starting with # are skipped)."""
    text = (text or "").lstrip("\ufeff")
    lines = text.splitlines()
    first = lines[0] if lines else ""
    delim = max(",;\t", key=first.count) if first else ","
    head = [x.strip().strip('"').lower() for x in first.split(delim)] if first.count(delim) else []
    items: list[dict] = []
    if any(c in head for c in BATCH_COLUMNS):
        reader = csv.DictReader(io.StringIO(text), delimiter=delim)
        cols = {c.strip().lower(): c for c in (reader.fieldnames or []) if c}
        col = next(cols[c] for c in BATCH_COLUMNS if c in cols)
        for r in reader:
            t = " ".join((r.get(col) or "").split())
            if not t:
                continue
            code = (r.get(cols.get("code", ""), "") or r.get(cols.get("target", ""), "") or "").strip()
            fw = (r.get(cols.get("framework", ""), "") or "").strip()
            items.append({"text": t, "framework": fw or ("sdg" if code else None), "code": code or None})
    else:
        for line in lines:
            t = " ".join(line.split())
            if t and not t.startswith("#"):
                items.append({"text": t, "framework": None, "code": None})
    if len(items) > 500:
        raise ValueError("At most 500 hypotheses per batch.")
    return items


def check_batch(ws: Workspace, items: list[dict], backend, model: str, progress: Progress | None = None,
                should_stop: Callable[[], bool] | None = None, **options: Any) -> dict:
    """Check several hypotheses one after the other. A problem with one is recorded and the next is checked; a
    problem with the model service stops the batch (what was done is kept)."""
    from .llm import BackendError

    say = _adapt(progress)
    batch_id = _new_id("b")
    done: list[str] = []
    errors: list[dict] = []
    stopped = None
    for i, it in enumerate(items, 1):
        if should_stop and should_stop():
            stopped = "Stopped."
            break
        say("item", f"Hypothesis {i} of {len(items)}", {"index": i, "total": len(items), "text": it["text"]})
        try:
            hid = check_hypothesis(ws, it["text"], backend, model, framework=it.get("framework"), code=it.get("code"),
                                   batch_id=batch_id, progress=lambda s, m, info=None, _i=i: say(s, m, {**(info or {}),
                                                                                                    "index": _i}),
                                   should_stop=should_stop, **options)
            done.append(hid)
        except BackendError as e:
            stopped = str(e)
            break
        except ValueError as e:
            errors.append({"index": i, "text": it["text"], "error": str(e)})
    return {"batch_id": batch_id, "checked": done, "errors": errors, "stopped": stopped, "total": len(items)}
