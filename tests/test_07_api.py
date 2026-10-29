"""0.7 through the local API: search, profiles, concepts, themes, references, highlights, quick look, repeatability.

The model passes run against the fake Ollama (tests/fake_servers.py), which answers the new task kinds."""
import io
import json
import zipfile

import pytest

import fake_servers
from test_app import CLAIM, Api, api, home  # noqa: F401  (fixtures)


@pytest.fixture
def ollama():
    srv, url = fake_servers.start("ollama")
    yield srv, url
    srv.shutdown()
    fake_servers.WRONG_FIRST.clear()


def _project(api, corpus, url):
    assert api("PUT", "/api/settings", {"backend": "ollama", "ollama_url": url, "ollama_model": fake_servers.FAKE_MODEL,
                                        "embed_backend": "off"})[0] == 200
    assert api("POST", "/api/projects", {"name": "P"})[0] == 200
    for f in sorted(corpus.iterdir()):
        assert api("POST", f"/api/projects/P/uploads?batch=b1&name={f.name}", raw=f.read_bytes())[0] == 200
    st, r = api("POST", "/api/projects/P/ingest", {"batch": "b1"})
    job = api.wait(r["job"]["id"])
    assert job["state"] == "done", job["error"]
    # 0.7: adding documents builds the glossary and finds the codes by keywords
    assert job["result"]["glossary"]["terms"] > 0 and job["result"]["codes"]["passages"] > 0
    return {d["file"]: d["doc_id"] for d in api("GET", "/api/projects/P/documents")[1]["documents"]}


def test_search_concordance_and_downloads(api, corpus, ollama):
    docs = _project(api, corpus, ollama[1])
    st, r = api("GET", "/api/projects/P/search?q=water")
    assert st == 200 and r["total"] >= 3 and r["per_doc"][0]["hits"] >= 3
    line = r["lines"][0]
    assert line["node"].lower().startswith("water") and line["location"] and line["end"] > line["start"]
    exact = api("GET", "/api/projects/P/search?q=jobs&exact=1")[1]
    stemmed = api("GET", "/api/projects/P/search?q=jobs")[1]
    assert stemmed["total"] >= exact["total"] >= 1
    ph = api("GET", "/api/projects/P/search?q=%22water%20stress%22")[1]
    assert ph["total"] >= 1 and all("water stress" in x["node"].lower() for x in ph["lines"])
    near = api("GET", "/api/projects/P/search?q=hydrogen%20exported&mode=near&near=12")[1]
    assert near["total"] >= 1
    assert api("GET", "/api/projects/P/search?q=water&sort=sideways")[0] == 400
    assert api("GET", f"/api/projects/P/search?q=water&docs={docs['birds.html']}")[1]["total"] == 0
    refs = api("GET", "/api/projects/P/search?q=Smith&exact=1")[1]["total"]
    assert api("GET", "/api/projects/P/search?q=Smith&exact=1&all=1")[1]["total"] > refs  # the reference list
    col = api("GET", "/api/projects/P/collocates?q=hydrogen")[1]
    assert "words" in col and col["hits"] >= 1
    st, body = api("GET", f"/api/projects/P/search.csv?q=water&token={api.token}", token=False)
    assert st == 200 and b"Words found" in body
    st, body = api("GET", "/api/projects/P/search.xlsx?q=water")
    assert st == 200 and body[:2] == b"PK"
    assert api("GET", "/api/projects/P/search?q=water&mode=meaning")[0] == 400  # no embeddings here


def test_profile_concepts_themes_references_and_highlights(api, corpus, ollama):
    srv, url = ollama
    docs = _project(api, corpus, url)
    paper = docs["paper_a.pdf"]

    # F2: the keyword pass ran on adding; the model pass reads passages with a candidate
    prof = api("GET", "/api/projects/P/profile")[1]
    assert prof["pass1"]["done"] and prof["grid"]["documents"] and any(f["name"] == "sdg_goals" for f in prof["frameworks"])
    plan = api("GET", "/api/projects/P/profile/plan")[1]
    assert plan["to_read"] > 0 and plan["run"]["now"]["model"] == fake_servers.FAKE_MODEL
    fake_servers.WRONG_FIRST.add("label")   # a first answer with a wrong quote is asked again
    st, r = api("POST", "/api/projects/P/profile/label", {"limit": 3})
    job = api.wait(r["job"]["id"])
    assert job["state"] == "done" and job["result"]["read"] == 3, job
    plan2 = api("GET", "/api/projects/P/profile/plan")[1]
    assert plan2["already_read"] == 3 and plan2["run"]["last"]["model"] == fake_servers.FAKE_MODEL
    assert plan2["run"]["changed"] == []  # same model, digest, prompt, seed and settings
    dp = api("GET", f"/api/projects/P/documents/{paper}/profile")[1]
    assert dp["codes"] and dp["codes"][0]["items"]
    item = next((i for c in dp["codes"] for i in c["items"] if i["source"] == "model"), None) or dp["codes"][0]["items"][0]
    assert item["reason"]
    assert api("PUT", f"/api/projects/P/labels/{item['label_id']}", {"decision": "reject"})[0] == 200
    dp2 = api("GET", f"/api/projects/P/documents/{paper}/profile")[1]
    assert item["label_id"] not in {i["label_id"] for c in dp2["codes"] for i in c["items"]}
    assert api("PUT", f"/api/projects/P/labels/{item['label_id']}", {"decision": "maybe"})[0] == 400
    st, added = api("POST", "/api/projects/P/labels", {"framework": "sdg", "passage_id": item["passage_id"],
                                                       "code": "8.5", "quote": None})
    assert st == 200 and added["label_id"]
    st, body = api("GET", "/api/projects/P/profile.xlsx")
    assert st == 200 and body[:2] == b"PK"
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(body))
    assert "How these results were made" in wb.sheetnames and "Key terms per document" in wb.sheetnames
    how = list(wb["How these results were made"].values)
    assert any(fake_servers.FAKE_MODEL in (r or ()) for r in how)
    goals = api("GET", "/api/projects/P/profile?level=group")[1]
    assert all(c["code"].isdigit() for c in goals["grid"]["codes"])

    # F3: the glossary, a concept card, a link judged by the model and one by a person
    g = api("GET", "/api/projects/P/terms?q=hydrogen")[1]
    assert g["built"] and g["items"] and all(any("hydrogen" in w.lower() for w in e["wordings"]) for e in g["items"])
    tid = api("GET", "/api/projects/P/concepts?name=hydrogen")[1]["term_id"]
    card = api("GET", f"/api/projects/P/concepts/{tid}")[1]
    assert len(card["documents"]) >= 2 and card["lines"] and card["wordings"][0]["label"]
    others = [e["term_id"] for e in api("GET", "/api/projects/P/terms")[1]["items"] if e["term_id"] != tid]
    other = others[0]
    assert api("GET", "/api/projects/P/concepts?name=zzzz")[0] == 404
    st, r = api("POST", "/api/projects/P/concepts/judge", {"pairs": [[tid, other]]})
    j = api.wait(r["job"]["id"])
    assert j["state"] == "done" and j["result"]["judged"] == 1, j
    card = api("GET", f"/api/projects/P/concepts/{tid}")[1]
    assert any(x["term_id"] == other for x in card["related"]["links"])
    assert api("PUT", "/api/projects/P/concepts/links", {"a": tid, "b": other, "decision": "change",
                                                         "relation": "related"})[0] == 200
    assert api("PUT", "/api/projects/P/concepts/links", {"a": tid, "b": tid, "decision": "add"})[0] == 400
    kt = api("GET", f"/api/projects/P/documents/{paper}/terms")[1]
    assert kt["terms"] and kt["spans"]

    # F4: themes by words, a rename, a model pass for the names, the export (on more text: the demo corpus too)
    from pathlib import Path
    demo = Path(__file__).resolve().parent.parent / "tools" / "demo_corpus"
    st, r = api("POST", "/api/projects/P/ingest", {"path": str(demo)})
    assert api.wait(r["job"]["id"])["state"] == "done"
    st, r = api("POST", "/api/projects/P/themes", {"k": 5, "seed": 0})
    j = api.wait(r["job"]["id"])
    assert j["state"] == "done", j["error"]
    rid = j["result"]["run_id"]
    v = api("GET", f"/api/projects/P/themes/{rid}?framework=sdg")[1]
    assert v["themes"] and v["run"]["mode"] == "words" and v["doc_grid"]
    t1 = v["themes"][0]["theme_id"]
    assert api("PUT", f"/api/projects/P/themes/{rid}/edit", {"what": "rename", "theme_id": t1, "value": "Water"})[0] == 200
    assert api("GET", f"/api/projects/P/themes/{rid}")[1]["themes"][0]["name"] == "Water"
    tp = api("GET", f"/api/projects/P/themes/{rid}/{t1}")[1]
    assert tp["passages"] and tp["total"] == v["themes"][0]["size"]
    st, r = api("POST", f"/api/projects/P/themes/{rid}/name")
    j = api.wait(r["job"]["id"])
    assert j["state"] == "done" and j["result"]["named"] >= 1, j
    v2 = api("GET", f"/api/projects/P/themes/{rid}")[1]
    assert v2["themes"][0]["name"] == "Water"   # the person's name wins over the model's
    assert any(t["name_source"] == "model" and t["quotes"] for t in v2["themes"][1:])
    st, body = api("GET", f"/api/projects/P/themes/{rid}/export.xlsx")
    assert st == 200 and body[:2] == b"PK"
    assert len(api("GET", "/api/projects/P/themes")[1]["runs"]) == 1

    # F5: reference chains for one document
    st, r = api("POST", f"/api/projects/P/documents/{paper}/references", {})
    j = api.wait(r["job"]["id"])
    assert j["state"] == "done", j["error"]
    ch = api("GET", f"/api/projects/P/documents/{paper}/references")[1]
    assert ch["chains"] and ch["chains"][0]["mentions"] and ch["status"]["windows_read"] >= 1
    m = ch["chains"][0]["mentions"][-1]
    assert api("GET", f"/api/projects/P/documents/{paper}/references?start={m['start']}&end={m['end']}")[1]["chain"]

    # F7: a check, then its highlighted copies
    st, r = api("POST", "/api/projects/P/checks", {"text": CLAIM, "k": 3, "expand": False})
    j = api.wait(r["job"]["id"])
    assert j["state"] == "done", j["error"]
    hid = j["result"]["hyp_id"]
    h = api("GET", f"/api/projects/P/hypotheses/{hid}")[1]
    graded = [x for p in h["parts"] for x in p["passages"] if x["verdict"]]
    doc = graded[0]["doc_id"]
    st, pdf = api("GET", f"/api/projects/P/documents/{doc}/annotated.pdf?hyp={hid}&token={api.token}", token=False)
    assert st == 200 and (pdf[:4] == b"%PDF" or pdf.lstrip()[:9].lower() == b"<!doctype")
    st, z = api("GET", f"/api/projects/P/hypotheses/{hid}/annotated.zip")
    assert st == 200 and "README.txt" in zipfile.ZipFile(io.BytesIO(z)).namelist()
    assert api("GET", f"/api/projects/P/documents/{doc}/annotated.pdf?layers=nonsense")[0] == 400
    st, pdf = api("GET", f"/api/projects/P/documents/{paper}/annotated.pdf?layers=labels,references")
    assert st == 200 and pdf[:4] == b"%PDF"

    # reproducibility: a seed goes with every request; the server is repeatable; the exports say what made them
    chats = [b for path, b in srv.requests if path == "/api/chat"]
    assert chats and all(b["options"].get("seed") == 42 for b in chats)
    rep = api("POST", "/api/servers/repeatable")[1]
    assert rep["repeatable"] is True and rep["seed"] is True
    st, body = api("GET", f"/api/projects/P/hypotheses/{hid}/export.xlsx")
    wb = openpyxl.load_workbook(io.BytesIO(body))
    assert "How these results were made" in wb.sheetnames


def test_quick_look_is_temporary_until_kept(api, corpus, ollama, home):
    assert api("PUT", "/api/settings", {"backend": "baseline"})[0] == 200
    pdf = (corpus / "paper_a.pdf").read_bytes()
    st, r = api("POST", "/api/quicklook?name=paper_a.pdf", raw=pdf)
    assert st == 200 and r["project"]["temporary"]
    name = r["project"]["name"]
    job = api.wait(r["job"]["id"])
    assert job["state"] == "done" and job["result"]["documents_added"]
    assert api("GET", "/api/state")[1]["project"] == name
    assert api("POST", "/api/quicklook?name=notes.doc", raw=b"x")[0] == 400   # (0.8 reads .docx; .doc stays refused)
    # not kept: moved to the trash a week later
    import time
    data = json.loads((home.home() / "quicklook.json").read_text())
    data[name]["created_at"] = time.time() - 8 * 86400
    (home.home() / "quicklook.json").write_text(json.dumps(data))
    st, r2 = api("POST", "/api/quicklook?name=again.pdf", raw=pdf)
    api.wait(r2["job"]["id"])
    assert api("POST", f"/api/projects/{r2['project']['name']}/keep")[1]["kept"] is True
    projects = {p["name"]: p for p in api("GET", "/api/state")[1]["projects"]}
    assert name not in projects and projects[r2["project"]["name"]]["temporary"] is None


def test_flexiconc_download_and_quick_look_discard(api, corpus, ollama, home):
    srv, url = ollama
    _project(api, corpus, url)
    st, body = api("GET", "/api/projects/P/flexiconc.sqlite")
    assert st == 200 and body[:15] == b"SQLite format 3"
    # search: other forms of a word, and the words around it by count
    r = api("GET", "/api/projects/P/search?q=export")[1]
    assert "other_forms" in r
    c = api("GET", "/api/projects/P/collocates?q=hydrogen&order=often&window=8")[1]
    assert c["window"] == 8 and c["order"] == "often"
    assert api("GET", "/api/projects/P/collocates?q=hydrogen&order=odd")[0] == 400
    # a quick look remembers the project it came from; removing it opens another project
    f = next(p for p in sorted(corpus.iterdir()) if p.suffix == ".txt")
    st, r = api("POST", f"/api/quicklook?name={f.name}", raw=f.read_bytes())
    assert st == 200 and api.wait(r["job"]["id"])["state"] == "done"
    name = r["project"]["name"]
    assert r["project"]["temporary"]["from"] == "P"
    assert api("DELETE", f"/api/projects/{name}", {"confirm": name})[0] == 200
    assert api("GET", "/api/state")[1]["project"] == "P"
