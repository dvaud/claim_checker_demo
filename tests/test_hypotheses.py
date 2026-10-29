"""0.5: hypotheses, statuses, search, embeddings, frameworks, and the defects fixed since 0.4.1."""
import json
import re
import sqlite3
import zipfile

import pytest

from enlens_claims import claimcheck, embed, frameworks, ingest, tasks
from enlens_claims.llm import BaselineBackend
from enlens_claims.retrieve import BM25
from enlens_claims.search import passage_index, rank_passages
from enlens_claims.store import Workspace
from enlens_claims.textutil import find_quote, normalise_with_map, tokens, tokens_v1

CLAIM = "Green hydrogen production can increase water stress in arid regions."


# ---------------------------------------------------------------------------------------------
# statuses
# ---------------------------------------------------------------------------------------------
def _v(verdict, conf, doc="d1"):
    return {"verdict": verdict, "confidence": conf, "doc_id": doc}


def test_claim_status_and_balance():
    assert claimcheck.claim_status([], 0.6) == "not_checked"
    assert claimcheck.claim_status([_v("not_enough_info", None)], 0.6) == "unsupported"
    assert claimcheck.claim_status([_v("supports", 0.9)], 0.6) == "supported"
    assert claimcheck.claim_status([_v("supports", 0.9), _v("refutes", 0.7, "d2")], 0.6) == "contested"
    assert claimcheck.claim_status([_v("supports", 0.9), _v("refutes", 0.5, "d2")], 0.6) == "supported"  # weak refutation
    assert claimcheck.claim_status([_v("partially_supports", 0.7), _v("refutes", 0.8, "d2")], 0.6) == "contested"
    assert claimcheck.claim_status([_v("refutes", 0.8)], 0.6) == "contradicted"
    b = claimcheck.balance([_v("supports", 0.8), _v("supports", 0.6), _v("refutes", 0.8, "d2"),
                            _v("partially_supports", 0.8, "d3")], 0.6)
    assert (b["docs_supporting"], b["docs_partly"], b["docs_refuting"]) == (1, 1, 1)
    assert b["lean"] == round((0.8 + 0.4 - 0.8) / (0.8 + 0.4 + 0.8), 2)
    assert claimcheck.balance([], 0.6)["lean"] is None


def test_aggregate_over_parts():
    agg = claimcheck.aggregate
    assert agg(["supported", "supported"]) == "supported"
    assert agg(["supported", "unsupported"]) == "partially_supported"
    assert agg(["supported", "contradicted"]) == "contradicted"  # a necessary part does not hold
    assert agg(["unsupported", "contradicted"]) == "contradicted"
    assert agg(["contested", "supported"]) == "contested"
    assert agg(["contested", "contradicted"]) == "contradicted"
    assert agg(["no_passages", "no_passages"]) == "no_passages"
    assert agg(["no_passages", "unsupported"]) == "unsupported"
    assert agg(["failed", "supported"]) == "partially_supported"  # not every part could be judged
    assert agg(["failed", "failed"]) == "failed" and agg(["not_checked"]) == "not_checked"


# ---------------------------------------------------------------------------------------------
# hypotheses
# ---------------------------------------------------------------------------------------------
def test_decomposed_hypothesis_is_checked_part_by_part(loaded):
    events = []
    hid = claimcheck.check_hypothesis(loaded, CLAIM + " It also creates local jobs, and it needs training.",
                                      BaselineBackend(), "baseline", decompose=True, k=3,
                                      progress=lambda s, m, i: events.append((s, i.get("part"))))
    r = claimcheck.hypothesis_result(loaded, hid)
    assert r["decomposed"] and len(r["parts"]) >= 2 and [p["part"] for p in r["parts"]] == list(range(1, len(r["parts"]) + 1))
    assert r["status"] == claimcheck.aggregate([p["status"] for p in r["parts"]])
    assert events[0] == ("decompose", None) and ("terms", 2) in events
    h = claimcheck.history(loaded)[0]
    assert h["hyp_id"] == hid and h["parts"] == len(r["parts"]) and h["status"] == r["status"]
    rows = claimcheck.export_rows(loaded)
    assert {row["part"] for row in rows} == {p["part"] for p in r["parts"]}
    claimcheck.delete_hypothesis(loaded, hid)
    for table in ("hypotheses", "user_claims", "tasks", "results"):
        assert loaded.scalar(f"SELECT COUNT(*) FROM {table}") == 0, table
    assert not loaded.scalar("SELECT COUNT(*) FROM verifications WHERE subject_kind='user_claim'")


def test_a_simple_hypothesis_stays_one_claim(loaded):
    hid = claimcheck.check_hypothesis(loaded, CLAIM, BaselineBackend(), "baseline", decompose=True, k=2)
    r = claimcheck.hypothesis_result(loaded, hid)
    assert not r["decomposed"] and r["parts"][0]["claim"] == CLAIM and r["parts"][0]["claim_id"] == hid


def test_same_claim_twice_in_a_second_gets_two_ids(loaded):
    a = claimcheck.new_claim(loaded, CLAIM)
    b = claimcheck.new_claim(loaded, CLAIM)
    assert a != b


def test_continue_check_reads_further_passages(loaded):
    cid = claimcheck.check_claim(loaded, CLAIM, BaselineBackend(), "baseline", k=2)
    first = {p["passage_id"] for p in claimcheck.result(loaded, cid)["passages"]}
    r = claimcheck.continue_check(loaded, cid, BaselineBackend(), "baseline", k=2, terms=["jobs", "workers"])
    res = claimcheck.result(loaded, cid)
    assert r["round"] == 2 and res["rounds"] == 2 and "jobs" in res["terms"]
    assert not ({p["passage_id"] for p in res["passages"] if p["round"] == 2} & first)


def test_batch_parsing_and_run(loaded):
    items = claimcheck.parse_batch("hypothesis,code\n\"" + CLAIM + "\",6.4\n,\nshort,\n")
    assert items == [{"text": CLAIM, "framework": "sdg", "code": "6.4"}, {"text": "short", "framework": None, "code": None}]
    assert len(claimcheck.parse_batch("# a note\nFirst hypothesis here.\n\nSecond hypothesis here.\n")) == 2
    res = claimcheck.check_batch(loaded, items, BaselineBackend(), "baseline", k=2)
    assert len(res["checked"]) == 1 and res["errors"][0]["index"] == 2 and res["stopped"] is None
    assert [h["hyp_id"] for h in claimcheck.history(loaded, batch_id=res["batch_id"])] == res["checked"]


def test_04_checks_become_hypotheses(tmp_path):
    db = tmp_path / "old" / "enlens.sqlite"
    db.parent.mkdir()
    cx = sqlite3.connect(db)
    cx.executescript("CREATE TABLE user_claims(claim_id TEXT PRIMARY KEY, text TEXT, target TEXT, terms_json TEXT, "
                     "rationale TEXT, options_json TEXT, model TEXT, created_at TEXT);"
                     "INSERT INTO user_claims VALUES ('u1', 'An old claim about water stress.', '6.4', '[]', 'r', '{}', "
                     "'m', '2026-01-01T00:00:00+00:00');")
    cx.commit()
    cx.close()
    ws = Workspace.open(db.parent)
    h = claimcheck.history(ws)
    assert h[0]["hyp_id"] == "u1" and h[0]["code"] == "6.4" and h[0]["framework"] == "sdg" and h[0]["status"] == "no_passages"
    Workspace.open(db.parent)  # migrating twice changes nothing
    assert ws.scalar("SELECT COUNT(*) FROM hypotheses") == 1


# ---------------------------------------------------------------------------------------------
# text and search
# ---------------------------------------------------------------------------------------------
def test_tokens_keep_letters_of_every_language_and_stem_consistently():
    assert tokens("Économie circulaire, waterstofeconomie, Wasserstoff") == ["économi", "circulair", "waterstofeconomi",
                                                                          "wasserstoff"]
    for a, b in [("stress", "stresses"), ("tax", "taxes"), ("price", "prices"), ("increase", "increasing"),
                 ("policy", "policies"), ("analysis", "analysis")]:
        assert tokens(a) == tokens(b), (a, b)
    # the SDG pack keeps the 0.4 tokenizer exactly (its selections must not change)
    assert tokens_v1("Économie increases stresses") == ["conomie", "increas", "stress"]


def test_quote_matching_with_special_characters():
    t = "Projects near İzmir and İstanbul could create jobs."
    s, e, _ = find_quote(t, "create jobs")
    assert t[s:e] == "create jobs"
    assert find_quote("İzmir hydrogen hub creates jobs", "creates jobs")[:2] == (19, 31)  # ends at the end: no crash
    nfd = "Electrolysis in Querétaro raises water stress."
    s, e, _ = find_quote(nfd, "Querétaro raises water")
    assert nfd[s:e] == "Querétaro raises water"
    assert find_quote("Large-scale electro­lysis needs water.", "electrolysis needs water")
    norm, m = normalise_with_map("İİ abc")
    assert len(norm) == len(m)


def test_bm25_weighting_and_inverted_index():
    docs = [["water", "stress"], ["water", "water", "jobs"], ["jobs", "training"], []]
    bm = BM25(docs)
    dense = bm.score(["water", "jobs"])
    assert dense[3] == 0 and set(bm.sparse(["water", "jobs"])) == {0, 1, 2}
    assert bm.score(["water", "water"]) == bm.score(["water"])  # 0.4 behaviour: query repeats ignored
    assert bm.score(["water", "water"], weighted=True)[0] == pytest.approx(2 * bm.score(["water"])[0])
    assert bm.top(["jobs"], 5) == sorted(bm.top(["jobs"], 5), key=lambda x: -x[1])


def test_index_follows_the_corpus(loaded, tmp_path):
    before = passage_index(loaded)
    assert passage_index(loaded) is before  # cached
    doc = loaded.scalar("SELECT doc_id FROM documents WHERE kind='html'")
    loaded.remove_document(doc)
    after = passage_index(loaded)
    assert after is not before and all(r["doc_id"] != doc for r in after.rows)
    assert all(p["doc_id"] != doc for p in rank_passages(loaded, "swallows migrate birds", k=5))


class HashEmbedder:
    """Deterministic stand-in for an embedding model: hashed bag of words."""

    model = "hash-embed"

    def __init__(self):
        self.calls = 0

    def embed(self, texts):
        import hashlib

        self.calls += 1
        out = []
        for t in texts:
            v = [0.0] * 32
            for w in tokens(t):
                v[int(hashlib.md5(w.encode()).hexdigest(), 16) % 32] += 1
            out.append(v if any(v) else [1.0] + [0.0] * 31)
        return out


def test_semantic_index_and_hybrid_ranking(loaded, monkeypatch):
    emb = HashEmbedder()
    st = embed.status(loaded, emb.model)
    assert st["embedded"] == 0 and not st["complete"]
    res = embed.build(loaded, emb, batch=2)
    assert res["complete"] and res["embedded"] == st["passages"]
    assert embed.build(loaded, emb)["added"] == 0  # nothing left to do
    picked = rank_passages(loaded, CLAIM, k=3, embedder=emb)
    assert picked and {p["found_by"] for p in picked} <= {"keywords", "meaning", "both"} and "both" in {p["found_by"] for p in picked}
    # without numpy the same ranking comes from pure Python
    import builtins
    real_import = builtins.__import__

    def no_numpy(name, *a, **kw):
        if name == "numpy":
            raise ImportError
        return real_import(name, *a, **kw)

    embed._VECTORS.clear()
    monkeypatch.setattr(builtins, "__import__", no_numpy)
    assert [p["passage_id"] for p in rank_passages(loaded, CLAIM, k=3, embedder=emb)] == [p["passage_id"] for p in picked]


# ---------------------------------------------------------------------------------------------
# frameworks
# ---------------------------------------------------------------------------------------------
def test_frameworks(loaded):
    with pytest.raises(ValueError, match="code and description"):
        frameworks.parse_csv("name,text\na,b\n")
    with pytest.raises(ValueError, match="appears twice"):
        frameworks.parse_csv("code,description\nA,x\nA,y\n")
    fw = frameworks.save(loaded, "Energy Policy", "﻿code,description,group\nEP1,Carbon tax raises energy prices,tax\n"
                         "EP2,Subsidies for heat pumps reduce gas demand,heat\n", title="Energy policy claims", label="claim")
    assert fw.name == "energy_policy" and fw.codes["EP1"].group == "tax"
    assert frameworks.get("energy_policy", loaded).title == "Energy policy claims"
    assert frameworks.resolve(loaded, "energy_policy", "ep2") == ("energy_policy", "EP2",
                                                                  "Subsidies for heat pumps reduce gas demand", "claim")
    assert frameworks.get("energy_policy", loaded).suggest("heat pumps cut gas use")[0]["code"] == "EP2"
    hid = claimcheck.check_hypothesis(loaded, CLAIM, BaselineBackend(), "baseline", framework="energy_policy",
                                      code="EP1", k=2)
    r = claimcheck.hypothesis_result(loaded, hid)
    assert r["code_label"] == "claim" and r["code_description"].startswith("Carbon tax")
    payload = json.loads(loaded.scalar("SELECT payload_json FROM tasks WHERE task_id=?", (f"check:{hid}",)))
    assert payload["framework_label"] == "claim" and "about claim EP1" in tasks.prompts.user_prompt("check", payload)
    with pytest.raises(ValueError, match="Unknown claim 'EP9'"):
        claimcheck.new_claim(loaded, CLAIM, target="EP9", framework="energy_policy")
    with pytest.raises(ValueError, match="not 'sdg'"):
        frameworks.save(loaded, "sdg", "code,description\na,b\n")


# ---------------------------------------------------------------------------------------------
# defects fixed since 0.4.1
# ---------------------------------------------------------------------------------------------
def test_zip_cannot_write_outside_its_folder(ws, tmp_path):
    evil = tmp_path / "src.zip"
    with zipfile.ZipFile(evil, "w") as z:
        z.writestr("../src_evil/payload.txt", "x" * 200)  # a sibling folder sharing the prefix 'src'
        z.writestr("ok.txt", "Hydrogen production needs water. " * 10)
    inner = tmp_path / "outer.zip"
    with zipfile.ZipFile(inner, "w") as z:  # '...zip' would unpack into 'inputs/..', the workspace itself
        buf = tmp_path / "dots.zip"
        with zipfile.ZipFile(buf, "w") as d:
            d.writestr("enlens.sqlite", "not a database")
            d.writestr("config.json", "{}")
        z.write(buf, "...zip")
    ingest.ingest(ws, [evil, inner])
    assert not list((ws.root / "inputs").glob("src_evil*")) and not (ws.root / "src_evil").exists()
    assert (ws.root / "enlens.sqlite").read_bytes()[:15] == b"SQLite format 3"
    assert ws.scalar("SELECT COUNT(*) FROM documents") == 1


def test_two_zips_of_the_same_name_keep_both(ws, tmp_path):
    for sub, text in (("a", "Electrolysers in Chile use desalinated seawater. " * 5),
                      ("b", "Ammonia terminals in Namibia create jobs for local workers. " * 5)):
        (tmp_path / sub).mkdir()
        with zipfile.ZipFile(tmp_path / sub / "sources.zip", "w") as z:
            z.writestr("report.txt", text)
    st = ingest.ingest(ws, [tmp_path / "a" / "sources.zip", tmp_path / "b" / "sources.zip"])
    assert st["added"] == 2


def test_one_unreadable_file_does_not_stop_the_others(ws, tmp_path, monkeypatch):
    d = tmp_path / "c"
    d.mkdir()
    (d / "a.txt").write_text("Hydrogen production needs water in arid regions. " * 5)
    (d / "b.txt").write_text("Hydrogen projects create local jobs for workers. " * 5)
    real = ingest.read_file
    monkeypatch.setattr(ingest, "read_file", lambda p, *a: (_ for _ in ()).throw(PermissionError("denied"))
                        if p.name == "a.txt" else real(p, *a))
    st = ingest.ingest(ws, [d])
    assert st["added"] == 1 and st["failed"][0].startswith("a.txt")


def test_a_failed_answer_keeps_the_earlier_claims_and_reviews_survive(loaded):
    from enlens_claims import review
    from enlens_claims.retrieve import screen

    screen(loaded)
    tasks.make_extract(loaded, model="m")
    tasks.run(loaded, "extract", BaselineBackend())
    cid, tid, text = loaded.query("SELECT claim_id, task_id, claim FROM claims ORDER BY claim_id LIMIT 1")[0]
    n_before = loaded.scalar("SELECT COUNT(*) FROM claims WHERE task_id=?", (tid,))
    with loaded.connect() as cx:
        cx.execute("UPDATE claims SET status='rejected' WHERE claim_id=?", (cid,))
        cx.execute("INSERT INTO reviews(item_kind, item_id, field, decision) VALUES ('claim', ?, 'all', 'reject')", (cid,))
        cx.execute("UPDATE tasks SET status='pending', attempts=0 WHERE task_id=?", (tid,))
    bad = {"claims": [{"claim": "A paraphrase the passage does not hold.", "activity": "x", "segment": "production",
                       "targets": [{"target": "6.4", "direction": "negative"}], "confidence": 0.9,
                       "evidence": [{"field": "claim", "quote": "words that are not in the passage"}]}]}
    ok, _ = tasks.ingest_result(loaded, tid, bad)
    assert not ok and loaded.scalar("SELECT COUNT(*) FROM claims WHERE task_id=?", (tid,)) == n_before
    payload = json.loads(loaded.scalar("SELECT payload_json FROM tasks WHERE task_id=?", (tid,)))
    tasks.ingest_result(loaded, tid, BaselineBackend().complete("extract", "", "", "baseline", payload))
    assert loaded.scalar("SELECT status FROM claims WHERE claim_id=?", (cid,)) == "rejected"  # same claim, same decision
    assert loaded.scalar("SELECT claim FROM claims WHERE claim_id=?", (cid,)) == text
    assert review.export(loaded, str(loaded.root / "q.csv"))["rows"] >= 0


def test_resplit_drops_grades_of_the_claim_that_moved(loaded):
    tasks.make_split(loaded, model="m")
    tasks.run(loaded, "split", BaselineBackend())
    tasks.make_verify_vca(loaded, model="m", k=3)
    tasks.run(loaded, "verify_set", BaselineBackend())
    graded = loaded.scalar("SELECT subject_id FROM verifications WHERE subject_kind='vca_claim' LIMIT 1")
    target = graded.split("#")[0]
    just = loaded.scalar("SELECT justification FROM vca_targets WHERE target=?", (target,))
    with loaded.connect() as cx:
        cx.execute("UPDATE tasks SET status='pending', attempts=0 WHERE task_id=?", (f"split:{target}",))
    reordered = {"claims": [{"text": f"Claim number {i} of this justification.", "quote": just[:40], "direction": "mixed"}
                            for i in range(2)]}
    ok, errs = tasks.ingest_result(loaded, f"split:{target}", reordered)
    assert ok, errs
    assert not loaded.scalar("SELECT COUNT(*) FROM verifications WHERE subject_id=?", (graded,))
    assert not loaded.scalar("SELECT COUNT(*) FROM tasks WHERE task_id=?", (f"verify_set:{graded}",))


def test_refresh_keeps_the_model_of_audits(loaded):
    tasks.make_split(loaded, model="m")
    tasks.run(loaded, "split", BaselineBackend())
    tasks.make_verify_vca(loaded, model="main", k=3)
    tasks.run(loaded, "verify_set", BaselineBackend())
    tasks.make_audit_vca(loaded, "second-model", sample=1.0)
    with loaded.connect() as cx:
        cx.execute("UPDATE tasks SET prompt_version='old' WHERE kind='verify_set'")
    tasks.refresh(loaded, "verify_set", model="new-main")
    models = {r[0]: r[1] for r in loaded.query("SELECT task_id, model FROM tasks WHERE kind='verify_set'")}
    assert {m for t, m in models.items() if ":audit:" in t} == {"second-model"}
    assert {m for t, m in models.items() if ":audit:" not in t} == {"new-main"}


def test_export_never_overwrites_a_batch_and_import_skips_odd_lines(loaded):
    for i in range(3):
        claimcheck.new_claim(loaded, f"{CLAIM} ({i})")
    for cid in [r[0] for r in loaded.query("SELECT claim_id FROM user_claims")]:
        claimcheck._add_task(loaded, f"expand:{cid}", "expand", "m", {"claim_id": cid, "claim": CLAIM})
    files = tasks.export(loaded, "expand", batch_size=1)
    files[1].unlink()
    with loaded.connect() as cx:
        cx.execute("UPDATE tasks SET status='pending' WHERE task_id IN (SELECT task_id FROM tasks LIMIT 1)")
    new = tasks.export(loaded, "expand", batch_size=1)
    assert new[0].name == "batch_0004.jsonl" and files[2].exists()
    res = loaded.root / "r.jsonl"
    res.write_text('[1, 2]\n{"task_id": "expand:nope", "output": {}}\n')
    assert tasks.import_results(loaded, [res])["invalid"] == 1
    assert tasks.check_results(loaded, [res])["problems"][0]["errors"] == ["each line must be a JSON object"]


def test_reasoning_text_before_the_json_is_ignored():
    assert tasks.parse_json('<think>{"not": "this"}</think>\n{"terms": ["a"]}') == {"terms": ["a"]}
    assert tasks.parse_json('reasoning cut off...</think>{"terms": ["b"]}') == {"terms": ["b"]}


def test_nested_transactions_roll_back_only_their_own_writes(ws):
    with ws.connect() as cx:
        cx.execute("INSERT INTO meta VALUES ('a', '1')")
        with pytest.raises(RuntimeError):
            with ws.connect() as inner:
                inner.execute("INSERT INTO meta VALUES ('b', '2')")
                raise RuntimeError
    assert ws.meta("a") == 1 and ws.meta("b") is None


def test_stop_between_steps_leaves_nothing(loaded):
    from enlens_claims.llm import BackendError

    calls = []
    with pytest.raises(BackendError, match="Stopped"):
        claimcheck.check_hypothesis(loaded, CLAIM, BaselineBackend(), "baseline", k=2,
                                    progress=lambda s, m: calls.append(s), should_stop=lambda: "terms_done" in calls)
    assert claimcheck.history(loaded) == [] and not loaded.scalar("SELECT COUNT(*) FROM tasks")


def test_refresh_with_a_model_reaches_every_kind(loaded):
    from enlens_claims.retrieve import screen

    screen(loaded)
    n = tasks.make_extract(loaded, model="old")
    with loaded.connect() as cx:
        cx.execute("UPDATE tasks SET status='done'")
    assert tasks.refresh(loaded, "extract", model="new") == n
    assert {r[0] for r in loaded.query("SELECT model FROM tasks WHERE kind='extract'")} == {"new"}


def test_batch_csv_from_european_spreadsheets():
    items = claimcheck.parse_batch("hypothesis;framework;code\n" + CLAIM + ";sdg;6.4\n")
    assert items == [{"text": CLAIM, "framework": "sdg", "code": "6.4"}]
    assert claimcheck.parse_batch("hypothesis,code\n") == []
    assert claimcheck.parse_batch("Statement\tcode\nWater use rises with electrolysis.\t\n")[0]["text"].startswith("Water")


def test_parallel_writers_wait_instead_of_failing(loaded):
    import threading

    errors = []

    def work(i):
        try:
            for j in range(20):
                claimcheck.new_claim(loaded, f"{CLAIM} thread {i} item {j}")
        except Exception as e:  # pragma: no cover - what this test guards against
            errors.append(e)

    ts = [threading.Thread(target=work, args=(i,)) for i in range(4)]
    for t in ts:
        t.start()
    for t in ts:
        t.join()
    assert not errors and loaded.scalar("SELECT COUNT(*) FROM user_claims") == 80


def test_framework_names_are_never_paths(loaded, tmp_path):
    victim = loaded.root.parent / "victim.csv"
    victim.write_text("code,description\na,b\n")
    (loaded.root / "frameworks").mkdir(exist_ok=True)
    with pytest.raises(KeyError):
        frameworks.delete(loaded, "../../victim")
    assert victim.exists()


def test_the_pack_tokenizer_is_exactly_04():
    # 0.4 kept soft hyphens (splitting the word) and decomposed accents; the pack must not change
    assert tokens_v1("electro\u00adlysis") == ["electro", "lysi"]  # as 0.4 tokenised it
    assert tokens_v1("café Économie") == ["cafe", "conomie"]


def test_a_persons_judgement_overrides_the_ais(loaded):
    cid = claimcheck.check_claim(loaded, CLAIM, BaselineBackend(), "baseline", k=4)
    r = claimcheck.result(loaded, cid)
    graded = [p for p in r["passages"] if p["verdict"]]
    assert r["status"] == "supported" and graded and all(p["ai_verdict"] == p["verdict"] for p in graded)
    for p in graded:  # every supporting excerpt is, in the reader's view, not about the hypothesis
        claimcheck.set_judgement(loaded, cid, p["passage_id"], "not_relevant", "about water use in general")
    r2 = claimcheck.result(loaded, cid)
    assert r2["status"] == "unsupported" and claimcheck.history(loaded)[0]["corrected"] == len(graded)
    changed = next(p for p in r2["passages"] if p["passage_id"] == graded[0]["passage_id"])
    assert changed["verdict"] is None and changed["ai_verdict"] == "supports" and changed["user_note"].startswith("about")
    ungraded = next(p for p in r2["passages"] if not p["ai_verdict"])
    claimcheck.set_judgement(loaded, cid, ungraded["passage_id"], "refutes")
    r3 = claimcheck.result(loaded, cid, threshold=0.99)  # a person's judgement counts at any strictness
    assert r3["status"] == "contradicted"
    claimcheck.set_judgement(loaded, cid, ungraded["passage_id"], None)
    with pytest.raises(ValueError):
        claimcheck.set_judgement(loaded, cid, ungraded["passage_id"], "maybe")
    rows = claimcheck.export_table(loaded, hyp_ids=[cid], every_excerpt=True)
    assert {r["Judged by"] for r in rows} >= {"you"} and rows[0]["Answer"] == claimcheck.STATUS_TEXT["unsupported"]


def test_judgements_below_the_strictness_are_uncertain_not_absent(loaded):
    cid = claimcheck.check_claim(loaded, CLAIM, BaselineBackend(), "baseline", k=4)
    r = claimcheck.result(loaded, cid, threshold=1.0)
    assert r["status"] == "uncertain" and r["balance"]["excerpts_uncounted"] >= 1
    assert all(not p["counted"] for p in r["passages"])
    assert claimcheck.aggregate(["uncertain", "unsupported"]) == "uncertain"
    h = claimcheck.history(loaded, threshold=1.0)[0]
    assert h["supports"] == 0 and h["uncounted"] >= 1  # the list follows the strictness too


def test_locations_rounds_and_excel_export(loaded, tmp_path):
    cid = claimcheck.check_claim(loaded, CLAIM, BaselineBackend(), "baseline", k=2)
    claimcheck.continue_check(loaded, cid, BaselineBackend(), "baseline", k=2, terms=["jobs", "workers"])
    r = claimcheck.result(loaded, cid)
    assert [x["round"] for x in r["rounds_info"]] == [1, 2] and r["rounds_info"][1]["terms"] == ["jobs", "workers"]
    assert r["rounds_info"][0]["ai_note"]
    assert all(p["location"] for p in r["passages"])
    txt = [p for p in r["passages"] if not p["page"]]
    assert all(re.fullmatch(r"paragraphs? \d+(–\d+)? of \d+", p["location"]) for p in txt)
    claimcheck.set_note(loaded, cid, "check page 2 by hand")
    rows = claimcheck.export_table(loaded)
    assert rows[0]["Your notes on this result"] == "check page 2 by hand" and rows[0]["Where"]
    out = tmp_path / "r.xlsx"
    claimcheck.write_xlsx(rows, out, "Excerpts", summary=claimcheck.summary_table(loaded))
    import openpyxl
    wb = openpyxl.load_workbook(out)
    assert wb.sheetnames == ["Summary", "Excerpts"]
    assert wb["Excerpts"]["A1"].value == "Checked on" and wb["Excerpts"].max_row == len(rows) + 1
    assert wb["Summary"].max_row == 2 and wb["Summary"]["B2"].value == CLAIM
    summary = claimcheck.summary_table(loaded)[0]
    assert summary["Answer"] == claimcheck.STATUS_TEXT[r["status"]] and summary["Excerpts read"] == len(r["passages"])
    assert all(row["Excerpt (the text the AI read)"] for row in rows)
    assert claimcheck.more_available(loaded, cid) in (True, False)


def test_a_text_files_title_is_its_first_line(ws, tmp_path):
    d = tmp_path / "t"
    d.mkdir()
    (d / "interviews.txt").write_text('"You never really meet anyone": young employees on remote work\nInterview study, '
                                 "2024\n\nWe interviewed 42 employees. " * 3)
    ingest.ingest(ws, [d])
    assert ws.scalar("SELECT title FROM documents") == '"You never really meet anyone": young employees on remote work'


def test_agreeing_with_the_ai_is_a_review_not_a_change(loaded):
    cid = claimcheck.check_claim(loaded, CLAIM, BaselineBackend(), "baseline", k=4)
    graded = [p for p in claimcheck.result(loaded, cid)["passages"] if p["ai_verdict"]]
    assert graded
    before = claimcheck.hypothesis_result(loaded, cid)
    read = before["excerpts_read"]
    assert before["to_review"] == read and before["to_review_relevant"] == len(graded) and before["reviewed"] == 0
    claimcheck.set_judgement(loaded, cid, graded[0]["passage_id"], graded[0]["ai_verdict"])  # "I agree"
    after = claimcheck.hypothesis_result(loaded, cid)
    assert (after["reviewed"], after["corrected"], after["to_review"]) == (1, 0, read - 1)
    assert after["to_review_relevant"] == len(graded) - 1
    h = claimcheck.history(loaded)[0]
    assert (h["reviewed"], h["corrected"], h["to_review"]) == (1, 0, read - 1)


def test_who_says_what_follows_the_counting_rule():
    item = lambda doc, verdict, counted=True, loc="page 1": {"doc_id": doc, "title": doc.upper(), "verdict": verdict,
                                                           "counted": counted, "location": loc}
    who = claimcheck.who_says_what([item("a", "supports"), item("a", "partially_supports", loc="page 2"),
                                    item("b", "partially_supports"), item("b", "refutes", loc="page 3"),
                                    item("c", "supports", counted=False), item("d", None)])
    assert [(w["doc_id"], w["side"], w["both_sides"]) for w in who] == [
        ("a", "supports", False), ("b", "partly", True), ("b", "refutes", True)]
    vs = [{"doc_id": "a", "verdict": "supports", "confidence": 0.9}, {"doc_id": "a", "verdict": "partially_supports",
          "confidence": 0.9}, {"doc_id": "b", "verdict": "partially_supports", "confidence": 0.9},
          {"doc_id": "b", "verdict": "refutes", "confidence": 0.9}]
    bal = claimcheck.balance(vs, 0.6)  # the same documents on the same sides as the counts
    assert (bal["docs_supporting"], bal["docs_partly"], bal["docs_refuting"], bal["docs_both"]) == (1, 1, 1, 1)


def test_quotes_in_text_files_are_located_by_paragraph(ws, tmp_path):
    d = tmp_path / "t"
    d.mkdir()
    paras = ["Rivermouth survey, 2024", "Methods. We asked 4,000 workers about every trip in one week. " * 4,
             "Findings. Remote work reduces commuting emissions by eighteen percent in the region. " * 4]
    (d / "survey.txt").write_text("\n\n".join(paras))
    ingest.ingest(ws, [d])
    text = ws.scalar("SELECT text FROM documents")
    starts = [r[0] for r in ws.query("SELECT start FROM paragraphs ORDER BY n")]
    assert len(starts) == 3 and text[starts[2]:].startswith("Findings.")
    pid = ws.scalar("SELECT passage_id FROM passages ORDER BY seq LIMIT 1")
    at = text.index("Remote work reduces")
    locs = claimcheck._locations(ws, [pid], {pid: at})
    assert locs[pid]["location"] == "paragraph 3 of 3" and locs[pid]["span"].startswith("paragraph")


def test_a_search_with_other_words_finds_those_words(loaded):
    cid = claimcheck.check_claim(loaded, CLAIM, BaselineBackend(), "baseline", k=2)
    r = claimcheck.continue_check(loaded, cid, BaselineBackend(), "baseline", k=4, terms=["zzzq", "yyyq"])
    assert r["passages"] == []  # words found nowhere: nothing more of the claim's own words is read instead
    r = claimcheck.continue_check(loaded, cid, BaselineBackend(), "baseline", k=4, terms=["jobs"])
    assert r["passages"] and all("job" in p["text"].lower() for p in r["passages"])


def test_a_person_can_point_to_the_sentence_they_mean(loaded):
    cid = claimcheck.check_claim(loaded, CLAIM, BaselineBackend(), "baseline", k=4)
    p = next(x for x in claimcheck.result(loaded, cid)["passages"] if x["ai_verdict"])
    words = p["text"].split()
    mine = " ".join(words[-8:])  # a sentence other than the AI's quote
    claimcheck.set_judgement(loaded, cid, p["passage_id"], "refutes", "this one says the opposite", quote=mine)
    q = next(x for x in claimcheck.result(loaded, cid)["passages"] if x["passage_id"] == p["passage_id"])
    assert q["quote_by"] == "you" and q["ai_quote"] == p["quote"]
    assert q["text"][q["quote_start"]:q["quote_end"]] == q["quote"] and q["quote"].split()[-1] == words[-1]
    claimcheck.set_judgement(loaded, cid, p["passage_id"], "refutes", "still")  # the quote is kept when left out
    assert next(x for x in claimcheck.result(loaded, cid)["passages"] if x["passage_id"] == p["passage_id"])["quote_by"] == "you"
    with pytest.raises(ValueError, match="text from this excerpt"):
        claimcheck.set_judgement(loaded, cid, p["passage_id"], "refutes", quote="words that are not in the excerpt at all")
    rows = claimcheck.export_table(loaded, hyp_ids=[cid])
    row = next(r for r in rows if r["Judged by"] == "you")
    assert row["Quote chosen by"] == "you" and row["AI's own judgement"]
    claimcheck.set_judgement(loaded, cid, p["passage_id"], None)  # removing the judgement removes the sentence too
    r = next(x for x in claimcheck.result(loaded, cid)["passages"] if x["passage_id"] == p["passage_id"])
    assert r["quote_by"] == "the AI" and r["quote"] == p["quote"]


# ---------------------------------------------------------------------------------------------
# 0.6: layout (paragraphs, headings), scanned pages
# ---------------------------------------------------------------------------------------------
REPORT_WITH_HEADINGS = """Commuting and emissions in Rivermouth, a short report of 2024

Summary

Remote work reduces commuting emissions, but by less than often assumed. Workers who stay home two days a week cut
their commuting distance by 35 percent.

3. Discussion
The results suggest that the saving is real. Part of it is lost to short car trips.

A second paragraph of the discussion says that public transport lost riders.
"""


def test_layout_finds_paragraphs_and_headings_without_changing_the_text(ws, tmp_path):
    from helpers import make_pdf

    (tmp_path / "r.txt").write_text(REPORT_WITH_HEADINGS)
    (tmp_path / "w.html").write_text("<html><body><h1>Rush hour is back</h1><p>Dateline, 2025</p>"
                                     "<p>Traffic jams are long again.</p><h2>Trains</h2><p>Trains are full.</p></body></html>")
    (tmp_path / "m.md").write_text("# Findings\n\nHomes use more heat. Offices use less.\n\n## Method\n\nA survey.\n")
    make_pdf(tmp_path / "p.pdf", ["Wellbeing study of hybrid workers in the panel\n\nAbstract. " +
                                  "Hybrid workers were happier than office workers in every year of the panel. " * 2 +
                                  "\n\nResults\nLoneliness rose among young remote workers. It did not rise among "
                                  "hybrid workers.", "Second page. The effect was largest for parents."])
    for f in ("r.txt", "w.html", "m.md", "p.pdf"):
        raw = ingest.read_file(tmp_path / f)
        doc = ingest.build_document(raw)
        assert doc.text == ingest.assemble(raw.pages)[0]  # the stored text is the same as without layout
        heads = [t for _a, _b, t in doc.layout.headings]
        starts = [doc.text[s0:s0 + 12] for _pg, s0 in doc.layout.paragraphs]
        if f == "r.txt":
            assert heads == ["Commuting and emissions in Rivermouth, a short report of 2024", "Summary", "3. Discussion"]
            assert starts == ["Commuting an", "Summary Remo", "Remote work ", "3. Discussio", "A second par"]  # blocks
        elif f == "w.html":
            assert heads == ["Rush hour is back", "Trains"] and "Dateline, 20" in starts  # a dateline is no heading
        elif f == "m.md":
            assert heads == ["Findings", "Method"] and starts == ["Findings Hom", "Homes use mo", "Method A sur", "A survey."]
        else:
            assert heads == ["Wellbeing study of hybrid workers in the panel", "Results"]
            assert [pg for pg, _s in doc.layout.paragraphs] == [1, 1, 1, 2]  # numbered per page
            assert starts == ["Wellbeing st", "Abstract. Hy", "Results Lone", "Second page."]
    ingest.ingest(ws, [tmp_path / "r.txt", tmp_path / "p.pdf"])
    assert ws.scalar("SELECT COUNT(*) FROM headings") == 5
    assert ws.scalar("SELECT COUNT(*) FROM paragraphs WHERE page IS NOT NULL") == 4
    pdf = ws.scalar("SELECT doc_id FROM documents WHERE kind='pdf'")
    ws.remove_document(pdf)
    assert ws.scalar("SELECT COUNT(*) FROM headings") == 3 and not ws.scalar("SELECT COUNT(*) FROM paragraphs WHERE page=1")


def test_pdf_quotes_are_cited_by_page_and_paragraph_and_excerpts_keep_their_breaks(ws, tmp_path):
    from helpers import make_pdf

    make_pdf(tmp_path / "p.pdf", ["Wellbeing study of hybrid workers in the panel\n\nAbstract. Hybrid workers were "
                                  "happier than office workers in every year of the panel.\n\nResults\n"
                                  "Loneliness rose among young remote workers who stayed home. It did not rise among "
                                  "hybrid workers of any age in the panel.", "Second page. The effect was largest."])
    (tmp_path / "r.txt").write_text(REPORT_WITH_HEADINGS)
    ingest.ingest(ws, [tmp_path])
    hid = claimcheck.check_hypothesis(ws, "Loneliness rose among young remote workers.", BaselineBackend(), "baseline",
                                      k=4, expand=False)
    res = claimcheck.result(ws, hid)
    pdf_item = next(x for x in res["passages"] if x["page"])
    assert pdf_item["quote"] and "Loneliness rose" in pdf_item["quote"]
    assert pdf_item["location"] == "page 1, paragraph 3" and pdf_item["span"] == "page 1, paragraphs 1–3"
    kinds = [(b["kind"], pdf_item["text"][b["at"]:b.get("end", b["at"] + 10)]) for b in pdf_item["breaks"]]
    assert ("heading", "Results") in kinds and ("paragraph", "Abstract. ") in kinds
    txt = next(x for x in res["passages"] if not x["page"])
    assert any(b["kind"] == "heading" and txt["text"][b["at"]:b["end"]] == "Summary" for b in txt["breaks"])
    ctx = claimcheck.neighbours(ws, pdf_item["passage_id"], 0, 0)
    assert ctx["passages"][0]["breaks"] == pdf_item["breaks"]


def _scanned_pdf(path, text):
    """A PDF whose page is only an image of the text (as a scanner makes it)."""
    PIL = pytest.importorskip("PIL")
    from PIL import Image, ImageDraw, ImageFont

    img = Image.new("L", (1240, 1754), 255)
    try:
        font = ImageFont.load_default(size=44)
    except TypeError:  # an old Pillow without scalable default font
        pytest.skip("Pillow too old to draw readable text")
    d = ImageDraw.Draw(img)
    for i, line in enumerate(text.split("\n")):
        d.text((90, 120 + i * 70), line, fill=0, font=font)
    img.save(path, "PDF", resolution=150)
    return PIL


def test_scanned_pages_are_read_with_local_ocr_when_it_is_installed(ws, tmp_path):
    import shutil

    pytest.importorskip("pypdfium2")
    _scanned_pdf(tmp_path / "scan.pdf", "Remote work reduces commuting\nemissions by eighteen percent.\n"
                                       "Workers cycle more often now.\nThe city counted every trip.")
    if not shutil.which("tesseract"):
        pytest.skip("tesseract is not installed")
    st = ingest.ingest(ws, [tmp_path / "scan.pdf"], ocr=ingest.OcrOptions("off"))
    assert st["added"] == 0 and st["scanned_no_text"] == ["scan.pdf"] and st["empty_files"] == ["scan.pdf"]
    st = ingest.ingest(ws, [tmp_path / "scan.pdf"])  # automatic: recognised on this computer
    assert st["added"] == 1 and st["ocr"] == ["scan.pdf"]
    text = ws.scalar("SELECT text FROM documents")
    assert "commuting" in text.lower() and "eighteen" in text.lower()
    from enlens_claims.sources import document_overview
    d = document_overview(ws)[0]
    assert d["ocr_pages"] == 1 and d["scanned"].startswith("Scanned: text recognised (OCR)")


# ---------------------------------------------------------------------------------------------
# 0.6: look before checking, scope notes, reading one excerpt again, the second sorting pass
# ---------------------------------------------------------------------------------------------
class Scripted:
    """A backend that answers each task kind from a list of prepared answers (the last one repeats)."""
    name = "scripted"

    def __init__(self, **answers):
        self.answers = {k: list(v) for k, v in answers.items()}
        self.calls: list[tuple[str, str]] = []

    def complete(self, kind, system, user, model, payload):
        self.calls.append((kind, user))
        got = self.answers.get(kind) or [{}]
        a = got.pop(0) if len(got) > 1 else got[0]
        return json.dumps(a(payload) if callable(a) else a)


def _quote_first(verdict, conf, scope=None):
    def answer(payload):
        words = payload["passages"][0]["text"].split()
        return {"evidence": [{"passage": 1, "verdict": verdict, "confidence": conf, "quote": " ".join(words[:8]),
                              "scope_note": scope}], "rationale": "r"}
    return answer


def test_look_before_checking_proposes_parts_and_words_and_the_check_uses_the_edited_ones(loaded):
    be = Scripted(decompose=[{"parts": [{"text": "Hydrogen production increases water stress.", "role": "outcome"},
                                        {"text": "Hydrogen export hubs create local jobs.", "role": "outcome"}]}],
                  expand=[{"terms": ["water stress", "groundwater", "aquifer"]}, {"terms": ["employment", "jobs"]}],
                  check=[_quote_first("supports", 0.8)])
    n_before = {t: loaded.scalar(f"SELECT COUNT(*) FROM {t}") for t in ("hypotheses", "user_claims", "tasks", "results")}
    prep = claimcheck.prepare_check(loaded, CLAIM + " It also creates jobs.", be, "m", decompose=True)
    assert [p["text"] for p in prep["parts"]] == ["Hydrogen production increases water stress.",
                                                  "Hydrogen export hubs create local jobs."]
    assert prep["split"] and prep["terms"][0] == ["groundwater", "aquifer"] and "employment" in prep["terms"][1]
    assert {t: loaded.scalar(f"SELECT COUNT(*) FROM {t}") for t in n_before} == n_before  # nothing kept
    assert claimcheck.history(loaded) == []

    be.calls.clear()
    hid = claimcheck.check_hypothesis(loaded, CLAIM + " It also creates jobs.", be, "m", decompose=True, k=2,
                                      parts=["Hydrogen production increases water stress in arid regions.",
                                             {"text": "Hydrogen projects create local jobs for workers.", "role": "outcome"},
                                             "  "],
                                      terms=[["desalination"], []], edited={"parts": True, "terms": True})
    assert [k for k, _u in be.calls] == ["check", "check"]  # no split, no extra words: the person chose them
    r = claimcheck.hypothesis_result(loaded, hid)
    assert [p["claim"] for p in r["parts"]] == ["Hydrogen production increases water stress in arid regions.",
                                                "Hydrogen projects create local jobs for workers."]
    assert r["parts"][0]["terms"] == ["desalination"] and r["parts"][1]["terms"] == []
    assert r["options"]["prepared"] == {"parts": True, "terms": True, "edited_parts": True, "edited_terms": True}
    assert r["decomposed"] and r["parts"][1]["role"] == "outcome"
    one = claimcheck.check_hypothesis(loaded, CLAIM, be, "m", k=2, parts=[CLAIM], terms=["aquifer"])
    assert claimcheck.hypothesis_result(loaded, one)["parts"][0]["claim_id"] == one  # one part: the hypothesis itself
    with pytest.raises(ValueError, match="Part 1 is too short"):
        claimcheck.check_hypothesis(loaded, CLAIM, be, "m", parts=["short"])


def test_an_excerpt_about_another_group_is_flagged_and_supports_at_most_partly(loaded):
    be = Scripted(expand=[{"terms": ["water"]}], check=[_quote_first("supports", 0.9, "about fully remote workers, not hybrid")])
    hid = claimcheck.check_hypothesis(loaded, CLAIM, be, "m", k=2)
    item = next(x for x in claimcheck.result(loaded, hid)["passages"] if x["ai_verdict"])
    assert item["scope_note"] == "about fully remote workers, not hybrid" and item["ai_verdict"] == "partially_supports"
    assert loaded.scalar("SELECT prompt_version FROM tasks WHERE task_id=?", (f"check:{hid}",)) == "check-v4"
    from enlens_claims import prompts
    assert "scope_note" in prompts.system_prompt("check") and "scope_note" not in prompts.system_prompt("verify_set")


def test_the_ai_can_read_one_excerpt_again_and_its_earlier_judgement_is_kept(loaded):
    be = Scripted(expand=[{"terms": ["water"]}], check=[_quote_first("supports", 0.75)],
                  reread=[_quote_first("partially_supports", 0.59)])
    hid = claimcheck.check_hypothesis(loaded, CLAIM, be, "m", k=2)
    item = next(x for x in claimcheck.result(loaded, hid)["passages"] if x["ai_verdict"])
    claimcheck.set_judgement(loaded, hid, item["passage_id"], "refutes", "the next sentence says the opposite")
    be.calls.clear()
    r = claimcheck.reread_excerpt(loaded, hid, item["passage_id"], be, "m")
    assert r["ok"] and r["before"] == {"verdict": "supports", "confidence": 0.75}
    assert r["after"]["verdict"] == "partially_supports" and r["after"]["confidence"] == 0.59
    kind, user = be.calls[0]
    assert kind == "reread" and "P2 |" not in user and "A reader noted" in user and "the opposite" in user
    now = next(x for x in claimcheck.result(loaded, hid)["passages"] if x["passage_id"] == item["passage_id"])
    assert now["reread"] and now["ai_verdict"] == "partially_supports" and now["confidence"] == 0.59
    assert now["ai_history"][0]["verdict"] == "supports" and now["ai_history"][0]["confidence"] == 0.75
    assert now["user_verdict"] == "refutes" and now["verdict"] == "refutes"  # the person's judgement still counts
    claimcheck.reread_excerpt(loaded, hid, item["passage_id"], be, "m", note="")
    assert len(next(x for x in claimcheck.result(loaded, hid)["passages"]
                    if x["passage_id"] == item["passage_id"])["ai_history"]) == 2
    assert "A reader noted" not in be.calls[-1][1]  # an empty comment: no hint
    with pytest.raises(KeyError):
        claimcheck.reread_excerpt(loaded, hid, "nope:0001", be, "m")
    claimcheck.delete_hypothesis(loaded, hid)
    assert not loaded.scalar("SELECT COUNT(*) FROM tasks WHERE kind='reread'")


def test_the_second_sorting_pass_prefers_passages_that_word_the_claim(ws, tmp_path):
    from enlens_claims.search import lexical_score, passage_index
    d = tmp_path / "s"
    d.mkdir()
    (d / "a.txt").write_text("Water is scarce. Stress at work rose. Arid soil cracks. Regions differ a lot. Hydrogen "
                             "burns cleanly. Production rose. Green paint dries. Increase noted.")  # the words, apart
    (d / "b.txt").write_text("In arid regions, green hydrogen production can increase water stress for farmers. " +
                             "Other remarks about the weather, the seasons, the markets, the elections and the football "
                             "results follow here at length. " * 3)                                  # the claim, said
    ingest.ingest(ws, [d])
    idx = passage_index(ws)
    q = tokens(CLAIM)
    a = next(r for r in idx.rows if r["text"].startswith("Water is"))
    b = next(r for r in idx.rows if r["text"].startswith("In arid"))
    assert lexical_score(q, idx.bm25.idf, b["text"]) > lexical_score(q, idx.bm25.idf, a["text"])
    first = rank_passages(ws, CLAIM, [], k=1, per_doc=1)
    second = rank_passages(ws, CLAIM, [], k=1, per_doc=1, second_pass=True)
    assert first[0]["passage_id"] == a["passage_id"] and second[0]["passage_id"] == b["passage_id"]
    hid = claimcheck.check_hypothesis(ws, CLAIM, BaselineBackend(), "baseline", k=1, expand=False, second_pass=True)
    assert claimcheck.hypothesis_result(ws, hid)["options"]["second_pass"] is True


# ---------------------------------------------------------------------------------------------
# 0.6: evidence map, calibration, documents added since, second model, project zip
# ---------------------------------------------------------------------------------------------
def _grade_all(verdicts):
    """A check answer grading passage i with verdicts[i] (supports / refutes / …), quoting its first words."""
    def answer(payload):
        ev = []
        for i, p in enumerate(payload["passages"], 1):
            v = verdicts.get(p["doc_id"][:0] + p["title"][:6], verdicts.get("*"))
            if v:
                ev.append({"passage": i, "verdict": v[0], "confidence": v[1], "quote": " ".join(p["text"].split()[:6]),
                           "scope_note": None})
        return {"evidence": ev, "rationale": "r"}
    return answer


def test_evidence_map_says_which_documents_decide_the_answers(loaded):
    titles = {r["title"][:6]: r["doc_id"] for r in loaded.query("SELECT doc_id, title FROM documents")}
    water = next(t for t in titles if t.startswith("Water"))
    be = Scripted(expand=[{"terms": ["x"]}], check=[_grade_all({water: ("supports", 0.9), "*": ("refutes", 0.8)})])
    h1 = claimcheck.check_hypothesis(loaded, CLAIM, be, "m", k=4, per_doc=1)
    be2 = Scripted(expand=[{"terms": ["x"]}], check=[_grade_all({water: ("supports", 0.9)})])
    h2 = claimcheck.check_hypothesis(loaded, "Hydrogen export hubs create local jobs in Morocco.", be2, "m", k=4, per_doc=1)
    m = claimcheck.evidence_map(loaded)
    rows = {h["hyp_id"]: h for h in m["hypotheses"]}
    wdoc = titles[water]
    assert list(m["cells"][h2]) == [wdoc] and m["cells"][h2][wdoc]["side"] == "supports" and rows[h2]["single_document"]
    assert m["cells"][h1][wdoc]["side"] == "supports" and any(c["side"] == "refutes" for c in m["cells"][h1].values())
    assert rows[h1]["status"] == "contested" and not rows[h1]["single_document"]
    w = next(d for d in m["documents"] if d["doc_id"] == wdoc)
    assert w["answers"] == 2 and w["decides"] == 2 and w["key"]  # without it: contradicted, and nothing found
    for x in claimcheck.result(loaded, h2)["passages"]:
        if x["doc_id"] == wdoc:
            claimcheck.set_judgement(loaded, h2, x["passage_id"], "not_relevant")
    assert h2 not in [h for h, c in claimcheck.evidence_map(loaded)["cells"].items() if c]  # a person's judgement counts
    table = claimcheck.map_table(m)
    assert table[0]["Hypothesis"] and any(str(v).startswith("+ ") for v in table[0].values()) and "Key" in table[-1]["Hypothesis"]


def test_calibration_counts_agreement_by_the_ais_certainty(loaded):
    be = Scripted(expand=[{"terms": ["x"]}], check=[_grade_all({"*": ("supports", 0.65)})])
    hids = [claimcheck.check_hypothesis(loaded, f"{CLAIM} Variant {i}.", be, "m", k=3) for i in range(8)]
    items = [(h, x["passage_id"]) for h in hids for x in claimcheck.result(loaded, h)["passages"] if x["ai_verdict"]]
    assert claimcheck.calibration(loaded)["total"] == 0
    for n, (h, pid) in enumerate(items):
        claimcheck.set_judgement(loaded, h, pid, "supports" if n % 2 else "not_relevant")
    c = claimcheck.calibration(loaded)
    b = next(b for b in c["bins"] if b["from"] == 0.6)
    assert b["judged"] == len(items) and b["agreed"] == len(items) // 2 and c["total"] == len(items)
    if len(items) >= 20:
        assert c["suggestion"] is None and "less than 80%" in c["note"]
    for h, pid in items:
        claimcheck.set_judgement(loaded, h, pid, "supports")
    c = claimcheck.calibration(loaded)
    assert c["total"] < 20 or c["suggestion"]["threshold"] == 0.6


def test_documents_added_since_a_check_are_read_without_running_it_again(loaded, tmp_path):
    be = Scripted(expand=[{"terms": ["water"]}], check=[_grade_all({"*": ("supports", 0.8)})])
    hid = claimcheck.check_hypothesis(loaded, CLAIM, be, "m", k=2)
    pid = claimcheck.result(loaded, hid)["passages"][0]["passage_id"]
    claimcheck.set_judgement(loaded, hid, pid, "refutes", "mine")
    assert claimcheck.history(loaded)[0]["docs_added_since"] == 0
    import time
    time.sleep(1.1)  # added_at has whole seconds
    (tmp_path / "new").mkdir()
    (tmp_path / "new" / "desal.txt").write_text("Desalination plants supply water for hydrogen production in arid "
                                                "regions, so water stress need not increase. " * 3)
    (tmp_path / "new" / "unrelated.txt").write_text("A recipe for bread with flour, yeast and salt, baked for an hour. " * 3)
    ingest.ingest(loaded, [tmp_path / "new"])
    assert claimcheck.history(loaded)[0]["docs_added_since"] == 2
    new = claimcheck.new_documents(loaded, hid)
    assert sorted((d["title"][:5], d["matching"]) for d in new) == [("A rec", False), ("Desal", True)]
    r = claimcheck.read_new_documents(loaded, hid, be, "m", k=4)
    assert r["documents"] == 2 and r["matching"] == 1 and r["passages"] >= 1
    res = claimcheck.result(loaded, hid)
    assert {x["title"][:5] for x in res["passages"] if x["round"] == 2} == {"Desal"}
    assert next(x for x in res["passages"] if x["passage_id"] == pid)["user_verdict"] == "refutes"  # review kept
    assert claimcheck.history(loaded)[0]["docs_added_since"] == 0 and claimcheck.new_documents(loaded, hid) == []


def test_project_zip_carries_everything_and_imports_elsewhere(loaded, tmp_path):
    from enlens_claims import bundle
    from enlens_claims.cli import main

    outside = tmp_path / "elsewhere"
    outside.mkdir()
    (outside / "field.txt").write_text("Field notes on hydrogen. Water stress rises near the plant in summer. " * 3)
    ingest.ingest(loaded, [outside / "field.txt"])
    (loaded.root / "frameworks").mkdir(exist_ok=True)
    (loaded.root / "frameworks" / "mine.csv").write_text("code,description\nH1,Water use rises\n")
    hid = claimcheck.check_hypothesis(loaded, CLAIM, BaselineBackend(), "baseline", k=3)
    claimcheck.set_judgement(loaded, hid, claimcheck.result(loaded, hid)["passages"][0]["passage_id"], "refutes", "n")
    before = claimcheck.hypothesis_result(loaded, hid)
    st = bundle.export_project(loaded, tmp_path / "p.zip")
    assert st["originals"] >= 1 and st["documents"] == 4 and not st["missing"]
    with zipfile.ZipFile(tmp_path / "p.zip") as zf:
        names = set(zf.namelist())
    assert {"enlens.sqlite", "config.json", "bundle.json", "frameworks/mine.csv"} <= names
    assert any(n.startswith("originals/") and n.endswith("field.txt") for n in names)
    (outside / "field.txt").unlink()  # the copy travels with the project
    ws2 = bundle.import_project(tmp_path / "p.zip", tmp_path / "copy")
    after = claimcheck.hypothesis_result(ws2, hid)
    assert after["status"] == before["status"] and after["reviewed"] == 1
    assert [x["quote"] for p in after["parts"] for x in p["passages"]] == [x["quote"] for p in before["parts"]
                                                                          for x in p["passages"]]
    from enlens_claims.sources import original_file
    assert all(original_file(ws2, d)["available"] and original_file(ws2, d)["inside"]
               for (d,) in ws2.query("SELECT doc_id FROM documents"))
    with pytest.raises(ValueError, match="not empty"):
        bundle.import_project(tmp_path / "p.zip", tmp_path / "copy")
    (tmp_path / "notzip.zip").write_bytes(b"PK no")
    with pytest.raises(ValueError, match="not a zip"):
        bundle.import_project(tmp_path / "notzip.zip", tmp_path / "x")
    with zipfile.ZipFile(tmp_path / "evil.zip", "w") as zf:
        zf.writestr("enlens.sqlite", (tmp_path / "copy" / "enlens.sqlite").read_bytes())
        zf.writestr("../../escaped.txt", "x")
    bundle.import_project(tmp_path / "evil.zip", tmp_path / "evil")
    assert not (tmp_path / "escaped.txt").exists() and not (tmp_path.parent / "escaped.txt").exists()
    assert main(["--ws", str(loaded.root), "export-project", str(tmp_path / "cli.zip")]) == 0
    assert main(["--ws", str(tmp_path / "cli-copy"), "import-project", str(tmp_path / "cli.zip")]) == 0
    assert Workspace.open(tmp_path / "cli-copy").scalar("SELECT COUNT(*) FROM hypotheses") == 1


def test_a_comment_alone_is_not_a_judgement_and_rereads_say_why(loaded):
    be = Scripted(expand=[{"terms": ["water"]}], check=[_grade_all({"*": ("supports", 0.75)})],
                  reread=[lambda pl: {**_quote_first("refutes", 0.6)(pl), "rationale": "The note is right: it says the opposite."}])
    hid = claimcheck.check_hypothesis(loaded, CLAIM, be, "m", k=2)
    pid = next(x["passage_id"] for x in claimcheck.result(loaded, hid)["passages"] if x["ai_verdict"])
    claimcheck.set_judgement(loaded, hid, pid, None, "check the next page")
    x = next(x for x in claimcheck.result(loaded, hid)["passages"] if x["passage_id"] == pid)
    assert x["user_note"] == "check the next page" and not x["reviewed"] and x["user_verdict"] is None
    assert claimcheck.calibration(loaded)["total"] == 0  # a comment is not agreement
    r = claimcheck.reread_excerpt(loaded, hid, pid, be, "m")  # the saved comment goes along as the hint
    assert r["after"]["rationale"] == "The note is right: it says the opposite."
    assert "whether the reader's note changed your judgement" in be.calls[-1][1]
    x = next(x for x in claimcheck.result(loaded, hid)["passages"] if x["passage_id"] == pid)
    assert x["reread_info"]["hint"] == "check the next page" and x["reread_info"]["rationale"].startswith("The note")
    claimcheck.set_judgement(loaded, hid, pid, "refutes", "check the next page")  # agrees with the re-read…
    c = claimcheck.calibration(loaded)
    assert c["total"] == 1 and c["agreed"] == 0  # …but the AI's own first judgement was "supports"
    claimcheck.set_judgement(loaded, hid, pid, None, "still unsure")  # back to the AI's: the comment stays
    x = next(x for x in claimcheck.result(loaded, hid)["passages"] if x["passage_id"] == pid)
    assert x["user_note"] == "still unsure" and x["user_verdict"] is None
