"""Test fixtures: a minimal PDF writer and a small VCA workbook in the real layout."""
from __future__ import annotations

import textwrap
from pathlib import Path


def _esc(s: str) -> str:
    return s.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")


def make_pdf(path: Path, pages: list[str], title: str | None = None) -> Path:
    """Write a simple text PDF (Helvetica, one Tj per line) that pypdf can read back."""
    objs: list[bytes] = []
    n_pages = len(pages)
    # 1 catalog, 2 pages, 3 font, then (page, content) pairs, then optional info
    kids = " ".join(f"{4 + 2 * i} 0 R" for i in range(n_pages))
    objs.append(b"<< /Type /Catalog /Pages 2 0 R >>")
    objs.append(f"<< /Type /Pages /Kids [{kids}] /Count {n_pages} >>".encode())
    objs.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    for i, text in enumerate(pages):
        lines = []
        for para in text.split("\n"):
            lines += textwrap.wrap(para, 95) or [""]
        ops = ["BT /F1 10 Tf 12 TL 50 780 Td"]
        for ln in lines:
            ops.append(f"({_esc(ln)}) Tj T*")
        ops.append("ET")
        stream = "\n".join(ops).encode("latin-1", errors="replace")
        objs.append(f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 842] /Resources << /Font << /F1 3 0 R >> >> "
                    f"/Contents {5 + 2 * i} 0 R >>".encode())
        objs.append(b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream")
    info_ref = ""
    if title:
        objs.append(f"<< /Title ({_esc(title)}) >>".encode())
        info_ref = f" /Info {len(objs)} 0 R"
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for i, o in enumerate(objs, start=1):
        offsets.append(len(out))
        out += f"{i} 0 obj\n".encode() + o + b"\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objs) + 1}\n0000000000 65535 f \n".encode()
    for off in offsets:
        out += f"{off:010d} 00000 n \n".encode()
    out += f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R{info_ref} >>\nstartxref\n{xref}\n%%EOF\n".encode()
    path.write_bytes(bytes(out))
    return path


RED, LIGHT_RED, LIGHT_GREEN = "FFFF0000", "FFEAC4BC", "FFC6EFCE"


def make_vca_workbook(path: Path) -> Path:
    """Three justified targets plus one empty target, in the layout of the real supplementary table."""
    import openpyxl
    from openpyxl.styles import PatternFill

    wb = openpyxl.Workbook()
    wb.active.title = "(T1) Methodological Steps"
    ws = wb.create_sheet("(T3) SDGs & Targets - Analyses")
    ws.cell(5, 2, "Goal"), ws.cell(5, 3, "Target")
    refs_a = ("(1) Author A, Author B. Water footprint of electrolytic hydrogen in arid export regions. "
              "Int J Test 2023;1:1-10. https://doi.org/10.9999/papera.2023 "
              "(2) IRENA. Hydrogen jobs outlook for exporting countries. 2022")
    rows = [
        ("6. Clean water and sanitation", "6.4", {5: RED}, "Short-term (by 2030)", "One segment",
         "Electrolysis needs large amounts of fresh water. Hydrogen production can increase water stress in arid regions.", refs_a),
        ("8. Decent work and economic growth", "8.5", {8: LIGHT_GREEN}, "Short-term", "No Reciprocal Interdependence (Prod+Trans)",
         "Green hydrogen production can create local jobs in exporting countries.",
         "(1) IRENA. Hydrogen jobs outlook for exporting countries. 2022"),
        ("1. No poverty", "1.1", {8: LIGHT_RED}, "Medium to Long-term", "One segment",
         "Dependence on imported equipment limits local economic benefits.",
         "(1) Author C. A paper nobody uploaded about imported equipment and poverty. Energy Policy 2021. https://doi.org/10.9999/missing.2021"),
        ("7. Affordable and clean energy", "7.1", {}, "", "", "", ""),
    ]
    for i, (goal, tgt, fills, hz, inter, just, refs) in enumerate(rows):
        r = 6 + i
        ws.cell(r, 2, goal), ws.cell(r, 3, tgt), ws.cell(r, 4, "description")
        for col, rgb in fills.items():
            ws.cell(r, col).fill = PatternFill("solid", fgColor=rgb)
        ws.cell(r, 11, hz), ws.cell(r, 12, inter), ws.cell(r, 13, just), ws.cell(r, 14, refs)
    wb.save(path)
    return path


PAPER_A_PAGES = [
    "Water footprint of electrolytic hydrogen in arid export regions. https://doi.org/10.9999/papera.2023\n"
    "Abstract. Green hydrogen is produced by electrolysis powered by renewable electricity. "
    "It requires large amounts of fresh water, and the fuel is often exported. "
    "Its production can increase water stress in arid regions such as Morocco. "
    "Electrolysis can also compete with farming for groundwater.",
    "Results. In the short term, export hubs could create new jobs for local workers in Morocco by 2030. "
    "Using desalinated seawater avoids drawing on freshwater in coastal plants. "
    "Hydrogen projects can reduce emissions in the steel industry of importing countries. https://doi.org/10.9999/papera.2023",
    "References. Smith J, Jones K. Hydrogen and water. J Hydrogen 2020;12(3):45-60. https://doi.org/10.1111/a.1 "
    "Brown L et al. Water stress indicators. Water Res 2019;3:1-9. https://doi.org/10.1111/b.2 "
    "Green P et al. Export hubs. Energy 2021;5:7-9.",
]

REPORT_B_TEXT = """IRENA. Hydrogen jobs outlook for exporting countries. 2022

Hydrogen production in exporting countries can create local jobs in construction and operation of electrolysers.
Training programmes are needed so that local workers benefit from hydrogen projects.
Without local manufacturing, most of the value of hydrogen equipment is captured abroad.
"""

UNRELATED_HTML = """<html><head><title>Bird migration notes</title></head><body>
<p>Swallows migrate south in autumn. They fly long distances over the Sahara every year.</p>
<p>Ringing studies show that many birds return to the same nest site in spring, often within a few metres.</p>
<script>var x = 1;</script></body></html>"""
