"""The data folder (settings, projects, key) and the local web app's API."""
import json
import os
import stat
import threading
import time
import urllib.error
import urllib.request

import pytest

import fake_servers
from enlens_claims.llm import BackendError, BaselineBackend, LocalBackend, OllamaBackend

CLAIM = "Green hydrogen production can increase water stress in arid regions."


@pytest.fixture
def home(tmp_path, monkeypatch):
    monkeypatch.setenv("ENLENS_HOME", str(tmp_path / "home"))
    monkeypatch.delenv("ANTHROPIC_API_KEY", raising=False)
    from enlens_claims import home
    return home


def test_settings_projects_and_key(home, tmp_path):
    s = home.load_settings()
    # 0.9: subscription first: a new installation answers with Claude Code; the Claude API stays opt-in
    assert s["backend"] == "claude_code" and s["allow_claude_code"] is True and s["allow_cloud"] is False
    assert s["num_ctx"] == "auto"
    s.update(backend="ollama", ollama_model="qwen3:8b", unknown="dropped")
    home.save_settings(s)
    assert home.load_settings()["ollama_model"] == "qwen3:8b" and "unknown" not in home.load_settings()
    with pytest.raises(ValueError, match="Allow cloud models first"):
        home.update_settings({"backend": "anthropic"})
    with pytest.raises(ValueError, match="threshold must be between"):
        home.update_settings({"threshold": 2})
    assert home.update_settings({"allow_cloud": True, "backend": "anthropic"})["backend"] == "anthropic"

    home.remember_key("sk-test")
    p = home.home() / "credentials.json"
    assert home.stored_key() == "sk-test" and stat.S_IMODE(os.stat(p).st_mode) == 0o600
    home.forget_key()
    assert home.stored_key() == ""

    ws = home.create_project("Green / hydrogen: Morocco!")
    assert ws.root.name == "Green hydrogen Morocco" and home.list_projects() == ["Green hydrogen Morocco"]
    with pytest.raises(ValueError, match="already exists"):
        home.create_project("Green hydrogen Morocco")
    with pytest.raises(ValueError, match="Give the project a name"):
        home.create_project("///")
    folder = home.save_uploads(ws, [("../../evil.txt", b"x" * 10), ("..\\..\\win.txt", b"y"), ("a.pdf", b"%PDF")])
    assert sorted(p.name for p in folder.iterdir()) == ["a.pdf", "evil.txt", "win.txt"]
    assert folder.parent.parent == ws.root

    other = tmp_path / "pilot"
    from enlens_claims.store import Workspace
    Workspace.create(other)
    assert home.link_existing(str(other)) == "pilot" and "pilot" in home.list_projects()
    assert home.project_path("pilot") == other.resolve() and home.is_linked("pilot")
    with pytest.raises(ValueError, match="No workspace"):
        home.link_existing(str(tmp_path))
    assert home.remove_project("pilot") == "unlinked" and (other / "enlens.sqlite").exists()
    assert home.remove_project("Green hydrogen Morocco").startswith("moved to") and home.list_projects() == []


def test_old_settings_keep_a_chosen_cloud_backend(home):
    (home.home() / "settings.json").write_text(json.dumps({"backend": "anthropic"}))
    assert home.load_settings()["allow_cloud"] is False  # no key: not silently allowed
    home.remember_key("sk-x")
    assert home.load_settings()["allow_cloud"] is True


def test_make_backend(home):
    s = dict(home.DEFAULTS)
    s.update(backend="anthropic")
    with pytest.raises(BackendError, match="Cloud models are switched off"):
        home.make_backend(s)
    s.update(backend="ollama")
    with pytest.raises(BackendError, match="Choose an Ollama model"):
        home.make_backend(s)
    s.update(ollama_model="qwen3:8b", num_ctx=32768, think="off")
    be, model = home.make_backend(s)
    assert isinstance(be, OllamaBackend) and be.num_ctx == 32768 and be.think is False and model == "qwen3:8b"
    s.update(backend="baseline")
    assert isinstance(home.make_backend(s)[0], BaselineBackend)
    assert home.backend_label(s).startswith("Offline demo")
    s.update(backend="auto", ollama_url="http://127.0.0.1:9", openai_url="http://127.0.0.1:9/v1")
    from enlens_claims import llm
    llm._DETECTED.clear()
    with pytest.raises(BackendError, match="No local model server is running"):
        home.make_backend(s)


def test_default_model_is_one_in_memory_then_the_newest(home):
    pick = home.pick_chat_model
    models = ["b-model:7b", "a-model:7b", "nomic-embed-text:latest", "c-model:7b"]
    assert pick(models) == "a-model:7b"                                            # nothing known: alphabetical
    assert pick(models, modified={"c-model:7b": "2026-09-01T10:00:00Z", "a-model:7b": "2026-01-01T00:00:00Z"}) == "c-model:7b"
    assert pick(models, loaded=["nomic-embed-text:latest", "b-model:7b"],
                modified={"c-model:7b": "2026-09-01T10:00:00Z"}) == "b-model:7b"  # in memory wins; embedders never
    assert pick(["nomic-embed-text:latest"]) == "" and home._first_chat_model(models) == "a-model:7b"

    srv, url = fake_servers.start("ollama")
    try:
        from enlens_claims import llm
        srv.modified = {fake_servers.SLOW: "2026-10-01T00:00:00Z"}
        llm._DETECTED.clear()
        be, model = home.make_backend(dict(home.DEFAULTS, backend="auto", ollama_url=url))
        assert model == fake_servers.SLOW                                          # the newest pulled
        srv.loaded = [fake_servers.FAKE_EMBED, fake_servers.LOOPER]
        llm._DETECTED.clear()
        assert home.make_backend(dict(home.DEFAULTS, backend="auto", ollama_url=url))[1] == fake_servers.LOOPER
        # an Ollama model that is not installed (or none chosen): the same choice, and the app says so
        be, model = home.make_backend(dict(home.DEFAULTS, backend="ollama", ollama_url=url, ollama_model="gone:1b"))
        assert model == fake_servers.LOOPER and "gone:1b is not installed" in be.model_note
        assert home.make_backend(dict(home.DEFAULTS, backend="ollama", ollama_url=url))[1] == fake_servers.LOOPER
        assert home.make_backend(dict(home.DEFAULTS, backend="ollama", ollama_url=url,
                                      ollama_model=fake_servers.FAKE_MODEL))[1] == fake_servers.FAKE_MODEL
        assert OllamaBackend(base_url=url).unload(fake_servers.LOOPER) and srv.unloads == [fake_servers.LOOPER]
        assert fake_servers.LOOPER not in OllamaBackend(base_url=url).loaded()
        assert LocalBackend(base_url=url + "/v1").unload("x") is False             # no standard way: nothing done
    finally:
        srv.shutdown()


def test_context_length_is_the_smallest_that_fits_and_stays_put(home, tmp_path):
    plan = home.context_plan
    s = dict(home.DEFAULTS)
    assert plan(s)["num_ctx"] == 8192 and plan(s)["setting"] == "auto"            # 8 excerpts, no reasoning
    assert plan({**s, "think": "on"})["num_ctx"] == 12288                          # room for the reasoning
    assert plan({**s, "k": 20, "think": "on"})["num_ctx"] == 16384
    assert plan(s, kinds=("extract",))["num_ctx"] == 16384                         # the pack's long prompt
    assert plan({**s, "num_ctx": 4096})["too_small"] is True
    assert home.update_settings({"num_ctx": "auto"})["num_ctx"] == "auto"
    assert home.update_settings({"num_ctx": "12288"})["num_ctx"] == 12288
    with pytest.raises(ValueError, match="auto or a whole number"):
        home.update_settings({"num_ctx": "big"})
    # a settings file of 0.5 holds the old default, 16384: it becomes "auto"; one saved by 0.6 keeps its number
    (home.home() / "settings.json").write_text(json.dumps({"num_ctx": 16384}))
    assert home.load_settings()["num_ctx"] == "auto"
    home.update_settings({"num_ctx": 16384})
    assert home.load_settings()["num_ctx"] == 16384

    srv, url = fake_servers.start("ollama")
    try:
        from enlens_claims import claimcheck, ingest
        from enlens_claims.store import Workspace
        from helpers import PAPER_A_PAGES, make_pdf
        ws = Workspace.create(tmp_path / "ctx")
        make_pdf(tmp_path / "a.pdf", PAPER_A_PAGES)
        ingest.ingest(ws, [tmp_path / "a.pdf"])
        for think, want in (("off", 8192), ("default", 8192)):                     # the fake model does not reason
            be, model = home.make_backend(dict(home.DEFAULTS, backend="ollama", ollama_url=url,
                                               ollama_model=fake_servers.FAKE_MODEL, think=think))
            assert be.num_ctx == want, think
        srv.contexts.clear()
        claimcheck.check_hypothesis(ws, CLAIM, be, model, decompose=True, k=3)
        assert srv.contexts and set(srv.contexts) == {8192}                        # one length: no reload
    finally:
        srv.shutdown()


def test_auto_backend_finds_a_local_server(home):
    srv, url = fake_servers.start("ollama")
    try:
        from enlens_claims import llm
        llm._DETECTED.clear()
        s = dict(home.DEFAULTS, backend="auto", ollama_url=url)
        be, model = home.make_backend(s)
        assert isinstance(be, OllamaBackend) and model in fake_servers.MODELS and model != fake_servers.FAKE_EMBED
        found = llm.detect_local(ollama_url=url)
        assert found[0]["kind"] == "ollama" and fake_servers.FAKE_EMBED in found[0]["models"]
    finally:
        srv.shutdown()


# ---------------------------------------------------------------------------------------------
# the API
# ---------------------------------------------------------------------------------------------
class Api:
    def __init__(self, base, token):
        self.base, self.token = base, token

    def __call__(self, method, path, body=None, raw=None, headers=None, token=True):
        h = {"X-Enlens-Token": self.token} if token else {}
        h.update(headers or {})
        data = raw if raw is not None else (json.dumps(body).encode() if body is not None else None)
        if body is not None:
            h["Content-Type"] = "application/json"
        req = urllib.request.Request(self.base + path, data=data, method=method, headers=h)
        try:
            with urllib.request.urlopen(req, timeout=30) as r:
                b = r.read()
                return r.status, (json.loads(b) if r.headers.get_content_type() == "application/json" else b)
        except urllib.error.HTTPError as e:
            return e.code, json.loads(e.read() or b"{}")

    def wait(self, job_id, timeout=60):
        t0 = time.time()
        while time.time() - t0 < timeout:
            st, j = self("GET", f"/api/jobs/{job_id}")
            if j["state"] not in ("queued", "running"):
                return j
            time.sleep(0.05)
        raise AssertionError("job did not finish")


@pytest.fixture
def api(home):
    from enlens_claims.gui.server import make_server

    srv, app = make_server("127.0.0.1", 0)
    threading.Thread(target=srv.serve_forever, daemon=True).start()
    yield Api(f"http://127.0.0.1:{srv.server_address[1]}", app.token)
    srv.shutdown()
    srv.server_close()


def test_api_guards(api):
    st, body = api("GET", "/api/state", token=False)
    assert st == 401 and "token" in body["error"]
    assert api("GET", "/api/state", headers={"Host": "evil.example"})[0] == 403  # DNS rebinding
    assert api("GET", "/api/state", headers={"Origin": "http://evil.example"})[0] == 403
    assert api("GET", "/api/nothing")[0] == 404
    assert api("DELETE", "/api/state")[0] == 405
    st, page = api("GET", "/", token=False)
    assert st == 200 and api.token.encode() in page and b"{{ENLENS_TOKEN}}" not in page
    assert api("GET", "/static/../server.py", token=False)[0] == 404
    st, s = api("GET", "/api/state")
    assert st == 200 and s["settings"]["backend"] == "claude_code" and s["projects"] == [] and s["statuses"]
    st, body = api("PUT", "/api/settings", {"backend": "anthropic"})
    assert st == 400 and "Allow cloud" in body["error"]


def test_api_project_sources_check_and_results(api, corpus):
    assert api("PUT", "/api/settings", {"backend": "baseline"})[0] == 200
    st, p = api("POST", "/api/projects", {"name": "Demo"})
    assert st == 200 and p["project"]["name"] == "Demo"
    assert api("GET", "/api/state")[1]["project"] == "Demo"
    for f in sorted(corpus.iterdir()):
        st, up = api("POST", f"/api/projects/Demo/uploads?batch=b1&name={f.name}", raw=f.read_bytes())
        assert st == 200 and up["name"] == f.name
    st, r = api("POST", "/api/projects/Demo/ingest", {"batch": "b1"})
    job = api.wait(r["job"]["id"])
    assert job["state"] == "done" and job["result"]["added"] == 3 and len(job["result"]["documents_added"]) == 3
    docs = api("GET", "/api/projects/Demo/documents")[1]["documents"]
    assert len(docs) == 3 and all(d["file"] for d in docs)

    st, bad = api("POST", "/api/projects/Demo/checks", {"text": "short"})
    assert st == 400 and "at least 10" in bad["error"]
    st, bad = api("POST", "/api/projects/Demo/checks", {"text": CLAIM, "framework": "sdg", "code": "99.9"})
    assert st == 400 and "Unknown SDG target" in bad["error"]
    st, r = api("POST", "/api/projects/Demo/checks", {"text": CLAIM + " It also creates local jobs.", "decompose": True,
                                                       "framework": "sdg", "code": "6.4", "k": 3})
    job = api.wait(r["job"]["id"])
    assert job["state"] == "done", job["error"]
    assert job["progress"]["decompose"]["state"] == "done" and len(job["progress"]["parts"]) == 2
    hid = job["result"]["hyp_id"]
    st, h = api("GET", f"/api/projects/Demo/hypotheses/{hid}")
    assert st == 200 and h["decomposed"] and len(h["parts"]) == 2 and h["code"] == "6.4"
    assert h["code_label"] == "SDG target" and h["status"] in ("supported", "partially_supported", "contested")
    graded = [x for p in h["parts"] for x in p["passages"] if x["verdict"]]
    assert graded and all(x["text"][x["quote_start"]:x["quote_end"]] == x["quote"] for x in graded)

    lst = api("GET", "/api/projects/Demo/hypotheses")[1]
    assert lst["hypotheses"][0]["hyp_id"] == hid and lst["hypotheses"][0]["parts"] == 2
    assert api("GET", "/api/projects/Demo/hypotheses?q=nomatch")[1]["hypotheses"] == []
    pid = graded[0]["passage_id"]
    ctx = api("GET", f"/api/projects/Demo/passages/{pid}?around=1")[1]
    assert any(x["current"] for x in ctx["passages"])
    st, csv_bytes = api("GET", f"/api/projects/Demo/hypotheses/{hid}/export.csv")
    assert st == 200 and b"hypothesis,status,part" in csv_bytes
    st, _ = api("GET", f"/api/projects/Demo/export.csv?token={api.token}", token=False)
    assert st == 200
    assert api("PUT", f"/api/projects/Demo/hypotheses/{hid}/note", {"note": "check page 2"})[0] == 200
    assert api("GET", f"/api/projects/Demo/hypotheses/{hid}")[1]["note"] == "check page 2"

    st, r = api("POST", f"/api/projects/Demo/claims/{h['parts'][0]['claim_id']}/more", {"k": 2, "terms": "jobs, workers"})
    more = api.wait(r["job"]["id"])
    assert more["state"] == "done" and more["result"]["round"] == 2

    st, prev = api("POST", "/api/projects/Demo/batches/preview", {"text": "# notes\n" + CLAIM + "\nshort\n"})
    assert prev["count"] == 2
    st, r = api("POST", "/api/projects/Demo/batches", {"text": "hypothesis,code\n\"" + CLAIM + "\",6.4\nshort,\n", "k": 2})
    b = api.wait(r["job"]["id"])
    assert b["state"] == "done" and len(b["result"]["checked"]) == 1 and b["result"]["errors"][0]["index"] == 2
    assert len(api("GET", f"/api/projects/Demo/hypotheses?batch={b['result']['batch_id']}")[1]["hypotheses"]) == 1

    st, fw = api("POST", "/api/projects/Demo/frameworks", {"name": "Mine", "csv": "code,description\nH1,Water use rises\n"})
    assert st == 200 and fw["framework"]["name"] == "mine" and fw["framework"]["codes"] == 1
    st, bad = api("POST", "/api/projects/Demo/frameworks", {"name": "x", "csv": "code\nH1\n"})
    assert st == 400 and "description" in bad["error"]
    assert api("GET", "/api/projects/Demo/suggest?text=water+stress+in+arid+regions")[1]["codes"][0]["code"] == "6.4"

    doc = docs[0]["doc_id"]
    assert api("GET", f"/api/projects/Demo/documents/{doc}")[1]["passages"]
    assert api("DELETE", f"/api/projects/Demo/documents/{doc}")[0] == 200
    assert len(api("GET", "/api/projects/Demo/documents")[1]["documents"]) == 2
    assert api("GET", f"/api/projects/Demo/hypotheses/{hid}")[0] == 200  # past checks keep their evidence
    assert api("DELETE", f"/api/projects/Demo/hypotheses/{hid}")[0] == 200
    assert api("GET", f"/api/projects/Demo/hypotheses/{hid}")[0] == 404
    assert api("DELETE", "/api/projects/Demo", {"confirm": "wrong"})[0] == 400


def test_api_check_with_a_local_server_semantic_index_and_cancel(api, corpus):
    srv, url = fake_servers.start("ollama")
    try:
        api("PUT", "/api/settings", {"backend": "ollama", "ollama_url": url, "ollama_model": fake_servers.FAKE_MODEL,
                                     "embed_backend": "ollama", "embed_model": fake_servers.FAKE_EMBED})
        assert api("POST", "/api/servers/test")[1]["embedding"]["dimensions"] == 64
        api("POST", "/api/projects", {"name": "P"})
        st, r = api("POST", "/api/projects/P/ingest", {"path": str(corpus)})
        job = api.wait(r["job"]["id"])
        assert job["result"]["semantic_index"]["complete"] is True
        assert api("GET", "/api/projects/P")[1]["search"]["embedded"] > 0
        st, r = api("POST", "/api/projects/P/checks", {"text": CLAIM, "k": 3})
        job = api.wait(r["job"]["id"])
        h = api("GET", f"/api/projects/P/hypotheses/{job['result']['hyp_id']}")[1]
        assert h["options"]["search"] == "keywords+meaning"
        assert {p["found_by"] for p in h["parts"][0]["passages"]} & {"both", "meaning"}

        srv.delay = 0.2  # slow enough to stop it midway
        st, r = api("POST", "/api/projects/P/checks", {"text": CLAIM, "k": 3})
        time.sleep(0.5)
        api("POST", f"/api/jobs/{r['job']['id']}/cancel")
        job = api.wait(r["job"]["id"])
        assert job["state"] == "cancelled"
        assert len(api("GET", "/api/projects/P/hypotheses")[1]["hypotheses"]) == 1  # the stopped one left nothing
    finally:
        srv.shutdown()


def test_backend_problem_is_reported_as_such(api):
    api("PUT", "/api/settings", {"backend": "ollama", "ollama_url": "http://127.0.0.1:9", "ollama_model": "m"})
    api("POST", "/api/projects", {"name": "E"})
    st, body = api("POST", "/api/servers/test")
    assert st == 502 and body["kind"] == "backend" and "Start the Ollama app" in body["error"]


def test_progress_tracker_and_model_line():
    from enlens_claims.gui.progress import Tracker, model_line

    now = [100.0]
    t = Tracker(expand=True, clock=lambda: now[0])
    t.step("decompose", "Splitting", {})
    now[0] += 3
    t.step("decompose_done", "2 claims to check", {"parts": ["A claim one.", "A claim two."]})
    t.step("terms", "Finding search terms", {"part": 1, "parts": 2, "claim": "A claim one."})
    t.model_report({"phase": "thinking", "seconds": 4.04, "thinking_tokens": 140, "rate": 11.66, "tail": "x " * 300})
    st = t.state()
    assert st["decompose"]["seconds"] == 3 and st["parts"][0]["steps"][0]["state"] == "running"
    assert st["model"]["rate"] == 11.7 and len(st["model"]["tail"]) <= 240 and "140 reasoning tokens" in st["model"]["line"]
    t.step("terms_done", "9 terms", {"part": 1, "parts": 2, "terms": ["a"]})
    t.step("terms", "Finding search terms", {"part": 2, "parts": 2, "claim": "A claim two."})
    st = t.state()
    assert [p["part"] for p in st["parts"]] == [1, 2] and st["parts"][0]["steps"][0]["info"] == {"terms": ["a"]}
    t.finish()
    now[0] += 50
    assert t.state()["elapsed"] == 3
    assert model_line({"phase": "done", "stats": {"load_seconds": 2, "prompt_tokens": 2700, "prompt_seconds": 18,
                                                  "output_tokens": 300, "output_seconds": 25}}).startswith(
        "Model finished: loaded the model in 2 s, read 2,700 prompt tokens in 18 s (150/s), wrote 300 tokens in 25 s (12.0/s)")


def test_openai_compatible_server_falls_back_and_streams(loaded):
    """LM Studio rejects json_object; llama.cpp streams SSE. The backend copes with both."""
    from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer

    seen = []

    class H(BaseHTTPRequestHandler):
        def log_message(self, *a):
            pass

        def do_GET(self):
            body = json.dumps({"data": [{"id": "local-model"}, {"id": "text-embedding-nomic"}]}).encode()
            self.send_response(200)
            self.send_header("Content-Length", str(len(body)))
            self.end_headers()
            self.wfile.write(body)

        def do_POST(self):
            req = json.loads(self.rfile.read(int(self.headers["Content-Length"])))
            seen.append(req.get("response_format", {}).get("type"))
            if req.get("response_format", {}).get("type") == "json_schema":
                body = json.dumps({"error": {"message": "response_format json_schema not supported"}}).encode()
                self.send_response(400)
                self.send_header("Content-Length", str(len(body)))
                self.end_headers()
                self.wfile.write(body)
                return
            answer = '<think>The user wants terms.</think>{"terms": ["water", "stress"]}'
            self.send_response(200)
            self.send_header("Content-Type", "text/event-stream")
            self.end_headers()
            for i in range(0, len(answer), 7):
                chunk = {"choices": [{"delta": {"content": answer[i:i + 7]}, "finish_reason": None}]}
                self.wfile.write(f"data: {json.dumps(chunk)}\n\n".encode())
            self.wfile.write(b"data: [DONE]\n\n")

    srv = ThreadingHTTPServer(("127.0.0.1", 0), H)
    threading.Thread(target=srv.serve_forever, daemon=True).start()
    try:
        url = f"http://127.0.0.1:{srv.server_address[1]}/v1"
        be = LocalBackend(base_url=url)
        assert be.check("local-model")["model_available"]
        from enlens_claims import claimcheck, tasks
        out = be.complete("expand", "You prepare a search", "CLAIM:\nwater", "local-model", {})
        assert tasks.parse_json(out) == {"terms": ["water", "stress"]}
        assert seen == ["json_schema", "json_object"] and be._format == "json_object"
        cid = claimcheck.check_claim(loaded, CLAIM, be, "local-model", k=2)
        assert claimcheck.result(loaded, cid)["terms"] == ["water", "stress"]
    finally:
        srv.shutdown()
    with pytest.raises(BackendError, match="Could not reach the local server"):
        LocalBackend(base_url="http://127.0.0.1:9/v1").models()


def test_remove_all_refuses_when_pipeline_results_would_go(api, loaded, home):
    name = home.link_existing(str(loaded.root), "Pipeline")
    from enlens_claims import tasks
    tasks.make_split(loaded, model="m")
    tasks.run(loaded, "split", BaselineBackend())
    tasks.make_verify_vca(loaded, model="m", k=3)
    st, body = api("DELETE", f"/api/projects/{name}/documents", {"confirm": True})
    assert st == 409 and "SDG pipeline" in body["error"] and "reset-corpus" in body["error"]
    assert loaded.scalar("SELECT COUNT(*) FROM documents") == 3


def test_upload_names_cannot_leave_the_project(api):
    api("POST", "/api/projects", {"name": "U"})
    for name in ("..%2F..%2Fevil.txt", "%5C..%5C..%5Cevil.txt", "C%3A%5Cevil.txt", ".hidden"):
        st, r = api("POST", f"/api/projects/U/uploads?batch=b&name={name}", raw=b"x")
        assert st == 200 and "/" not in r["name"] and "\\" not in r["name"] and ":" not in r["name"], r
    assert api("DELETE", "/api/projects/U/frameworks/..%2F..%2Fx")[0] == 404


def test_network_mode_hands_out_no_token(home):
    from enlens_claims.gui.server import make_server

    srv, app = make_server("0.0.0.0", 0)
    threading.Thread(target=srv.serve_forever, daemon=True).start()
    try:
        port = srv.server_address[1]
        a = Api(f"http://127.0.0.1:{port}", app.token)
        st, page = a("GET", "/", token=False)
        assert st == 200 and app.token.encode() not in page
        assert a("GET", "/api/state", headers={"Host": f"rebind.attacker.example:{port}"})[0] == 403
        assert a("GET", "/api/state", headers={"Host": f"192.168.1.20:{port}"})[0] == 200  # an address of this computer
        assert a("GET", "/api/state", token=False)[0] == 401
    finally:
        srv.shutdown()
        srv.server_close()


def test_an_installer_can_switch_the_cloud_off_for_good(api, home, monkeypatch):
    assert api("GET", "/api/state")[1]["local_only"] is False
    monkeypatch.setenv("ENLENS_LOCAL_ONLY", "1")
    st, s = api("GET", "/api/state")
    assert s["local_only"] is True
    st, body = api("PUT", "/api/settings", {"allow_cloud": True, "backend": "anthropic"})
    assert st == 400 and "switched off on this computer" in body["error"]
    with pytest.raises(BackendError, match="switched off on this computer"):
        home.make_backend({**home.load_settings(), "backend": "anthropic", "allow_cloud": True})
    monkeypatch.delenv("ENLENS_LOCAL_ONLY")
    (home.home() / "local_only").write_text("")
    assert home.local_only()


def test_a_document_reads_by_paragraph_or_page(api, corpus):
    assert api("PUT", "/api/settings", {"backend": "baseline"})[0] == 200
    api("POST", "/api/projects", {"name": "Docs"})
    for f in sorted(corpus.iterdir()):
        api("POST", f"/api/projects/Docs/uploads?batch=b1&name={f.name}", raw=f.read_bytes())
    api.wait(api("POST", "/api/projects/Docs/ingest", {"batch": "b1"})[1]["job"]["id"])
    docs = api("GET", "/api/projects/Docs/documents")[1]["documents"]
    for d in docs:
        st, v = api("GET", f"/api/projects/Docs/documents/{d['doc_id']}?view=blocks&limit=500")
        assert st == 200 and v["blocks"] and v["total_blocks"] == len(v["blocks"])
        assert v["blocks"][0]["label"].startswith("Page" if v["n_pages"] else ("Paragraph", "The whole text"))
        assert all(b["text"] for b in v["blocks"])


def test_the_original_file_opens_safely_and_a_moved_file_is_said_plainly(api, corpus, tmp_path):
    api("PUT", "/api/settings", {"backend": "baseline"})
    api("POST", "/api/projects", {"name": "F"})
    for f in sorted(corpus.iterdir()):
        api("POST", f"/api/projects/F/uploads?batch=b1&name={f.name}", raw=f.read_bytes())
    api.wait(api("POST", "/api/projects/F/ingest", {"batch": "b1"})[1]["job"]["id"])
    outside = tmp_path / "outside"
    outside.mkdir()
    (outside / "field_notes.txt").write_text("Notes kept outside the project.\n\nHydrogen needs water in deserts.\n" * 3)
    api.wait(api("POST", "/api/projects/F/ingest", {"path": str(outside / "field_notes.txt")})[1]["job"]["id"])
    docs = {d["file"]: d["doc_id"] for d in api("GET", "/api/projects/F/documents")[1]["documents"]}

    def get(doc, extra=""):
        req = urllib.request.Request(f"{api.base}/api/projects/F/documents/{doc}/file?token={api.token}{extra}")
        with urllib.request.urlopen(req, timeout=10) as r:
            return r.headers, r.read()

    h, body = get(docs["paper_a.pdf"])
    assert h["Content-Type"] == "application/pdf" and body.startswith(b"%PDF") and h["Content-Disposition"].startswith("inline")
    h, body = get(docs["birds.html"])
    assert h["Content-Type"].startswith("text/html") and "sandbox" in h["Content-Security-Policy"]
    assert "default-src 'none'" in h["Content-Security-Policy"] and b"<script>" in body  # served, but cannot run
    h, body = get(docs["field_notes.txt"])
    assert h["Content-Type"].startswith("text/plain") and body.startswith(b"Notes kept outside")
    from urllib.parse import quote as urlq
    h, body = get(docs["field_notes.txt"], "&quote=" + urlq("Hydrogen needs water in deserts."))
    assert h["Content-Type"].startswith("text/html") and "sandbox" in h["Content-Security-Policy"]
    assert b'<mark id="quote">Hydrogen needs water in deserts.</mark>' in body
    assert api("GET", f"/api/projects/F/documents/{docs['field_notes.txt']}/file", token=False)[0] == 401
    st, chk = api("GET", f"/api/projects/F/documents/{docs['field_notes.txt']}/file?check=1")
    assert chk["available"] is True and chk["name"] == "field_notes.txt"
    (outside / "field_notes.txt").rename(outside / "moved_notes.txt")
    st, chk = api("GET", f"/api/projects/F/documents/{docs['field_notes.txt']}/file?check=1")
    assert chk["available"] is False and "no longer at" in chk["reason"]
    h, body = get(docs["field_notes.txt"])
    assert body.decode().startswith("The original file is no longer at") and b"Hydrogen needs water" in body
    (outside / "moved_notes.txt").write_text("changed")
    (outside / "moved_notes.txt").rename(outside / "field_notes.txt")
    assert "has changed since it was added" in api("GET", f"/api/projects/F/documents/{docs['field_notes.txt']}/file?check=1")[1]["reason"]
    v = api("GET", f"/api/projects/F/documents/{docs['paper_a.pdf']}?view=blocks&limit=500")[1]
    assert v["original"]["available"] and v["blocks"][0]["label"].startswith("Page 1")
    assert all(b["text"] and b["page"] for b in v["blocks"])
    st, s = api("GET", "/api/state")
    assert set(s["ocr"]) >= {"available", "missing", "languages"} and s["context"]["setting"] == "auto"


def _social_project(api, url, name="S"):
    """A project with the remote-work demo corpus and the fake model that grades by word overlap."""
    import sys
    from pathlib import Path
    sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "tools"))
    from dev_server import social_documents
    from helpers import make_pdf
    api("PUT", "/api/settings", {"backend": "ollama", "ollama_url": url, "ollama_model": fake_servers.FAKE_MODEL,
                                 "embed_backend": "off", "embed_model": ""})
    api("POST", "/api/projects", {"name": name})
    import tempfile
    folder = social_documents(Path(tempfile.mkdtemp()) / "docs", make_pdf)
    job = api.wait(api("POST", f"/api/projects/{name}/ingest", {"path": str(folder)})[1]["job"]["id"])
    assert job["state"] == "done", job
    return folder


def test_api_look_before_checking_scope_notes_and_reading_again(api):
    srv, url = fake_servers.start("ollama")
    srv.varied = True
    try:
        _social_project(api, url)
        hyp = "Hybrid work improves wellbeing and increases loneliness among young employees."
        st, r = api("POST", "/api/projects/S/prepare", {"text": hyp, "decompose": True})
        assert st == 200 and r["job"]["kind"] == "prepare"
        prep = api.wait(r["job"]["id"])["result"]
        assert len(prep["parts"]) == 2 and len(prep["terms"]) == 2 and prep["split"]
        assert api("GET", "/api/projects/S/hypotheses")[1]["hypotheses"] == []  # nothing kept
        st, bad = api("POST", "/api/projects/S/checks", {"text": hyp, "parts": ["short"]})
        assert st == 400 and "Part 1" in bad["error"]
        parts = [prep["parts"][0]["text"], prep["parts"][1]["text"]]
        st, r = api("POST", "/api/projects/S/checks", {"text": hyp, "parts": parts, "terms": [prep["terms"][0], []],
                                                       "edited": {"terms": True}, "k": 4})
        job = api.wait(r["job"]["id"])
        assert job["state"] == "done", job["error"]
        h = api("GET", f"/api/projects/S/hypotheses/{job['result']['hyp_id']}")[1]
        assert h["options"]["prepared"]["edited_terms"] is True and h["parts"][1]["terms"] == []
        # the fake flags "fully remote" excerpts for a hypothesis about hybrid work
        flagged = [x for p in h["parts"] for x in p["passages"] if x["scope_note"]]
        assert flagged and all(x["ai_verdict"] != "supports" for x in flagged)
        assert "fully remote" in flagged[0]["scope_note"]

        x = next(x for p in h["parts"] for x in p["passages"] if x["ai_verdict"])
        cid = next(p["claim_id"] for p in h["parts"] if x in p["passages"])
        st, r = api("POST", f"/api/projects/S/claims/{cid}/reread/{x['passage_id']}", {"note": "It says the opposite."})
        res = api.wait(r["job"]["id"])
        assert res["state"] == "done" and res["result"]["after"]["verdict"] == "refutes"
        h2 = api("GET", f"/api/projects/S/hypotheses/{h['hyp_id']}")[1]
        y = next(y for p in h2["parts"] for y in p["passages"] if y["passage_id"] == x["passage_id"] and y["n"] == x["n"])
        assert y["reread"] and y["ai_history"][0]["verdict"] == x["ai_verdict"] and y["ai_verdict"] == "refutes"
        assert api("POST", f"/api/projects/S/claims/{cid}/reread/nope", {})[0] == 404
    finally:
        srv.shutdown()


def test_api_second_model_keeps_one_model_in_memory_and_shows_disagreements(api):
    srv, url = fake_servers.start("ollama")
    srv.varied = True
    try:
        _social_project(api, url, "M")
        api("PUT", "/api/settings", {"k": 4})
        ids = []
        for text in ("Working from home reduces commuting emissions.", "Hybrid work improves wellbeing."):
            job = api.wait(api("POST", "/api/projects/M/checks", {"text": text})[1]["job"]["id"])
            ids.append(job["result"]["hyp_id"])
        st, bad = api("POST", "/api/projects/M/second-opinion", {"hyp_ids": ids})
        assert st == 400 and "second model" in bad["error"]
        api("PUT", "/api/settings", {"second_model": fake_servers.SECOND})
        assert srv.loaded[0] == fake_servers.FAKE_MODEL and srv.peak_loaded == 1
        st, r = api("POST", "/api/projects/M/second-opinion", {"hyp_ids": ids})
        job = api.wait(r["job"]["id"])
        assert job["state"] == "done", job["error"]
        assert job["result"]["graded"] == 2 and job["result"]["unloaded"] == [fake_servers.FAKE_MODEL, fake_servers.SECOND]
        assert srv.peak_loaded == 1 and not [m for m in srv.loaded if "embed" not in m]  # never two at once; none left
        h = api("GET", f"/api/projects/M/hypotheses/{ids[0]}")[1]
        items = [x for p in h["parts"] for x in p["passages"]]
        assert all(x["second"] and x["second"]["model"] == fake_servers.SECOND for x in items)
        sup = [x for x in items if x["ai_verdict"] == "supports"]
        assert sup and all(x["second"]["verdict"] == "partially_supports" and not x["second"]["agrees"] for x in sup)
        assert h["disagreements"] == len(sup) and h["second_models"] == [fake_servers.SECOND]
        row = next(r for r in api("GET", "/api/projects/M/hypotheses")[1]["hypotheses"] if r["hyp_id"] == ids[0])
        assert row["disagreements"] == len(sup)
        again = api.wait(api("POST", "/api/projects/M/second-opinion", {"hyp_ids": ids})[1]["job"]["id"])
        assert again["result"]["skipped"] == 2 and again["result"]["graded"] == 0  # not asked twice
        st, x = api("GET", f"/api/projects/M/hypotheses/{ids[0]}/export.xlsx")
        import io
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(x))
        head = [c.value for c in wb["Excerpts"][1]]
        assert "Second model's judgement" in head and "The two models agree" in head
    finally:
        srv.shutdown()


def test_api_map_calibration_new_documents_and_project_zip(api, tmp_path):
    srv, url = fake_servers.start("ollama")
    srv.varied = True
    try:
        folder = _social_project(api, url, "Z")
        api("PUT", "/api/settings", {"k": 4})
        hids = [api.wait(api("POST", "/api/projects/Z/checks", {"text": t})[1]["job"]["id"])["result"]["hyp_id"]
                for t in ("Working from home reduces commuting emissions.", "Remote work increases loneliness.")]
        m = api("GET", "/api/projects/Z/map")[1]
        assert {h["hyp_id"] for h in m["hypotheses"]} == set(hids) and m["documents"] and m["cells"]
        assert all(set(c) <= {d["doc_id"] for d in m["documents"]} for c in m["cells"].values())
        assert all(v["side"] in ("supports", "partly", "refutes", "both") for c in m["cells"].values() for v in c.values())
        cal = api("GET", "/api/projects/Z/calibration")[1]
        assert cal["total"] == 0 and cal["suggestion"] is None and len(cal["bins"]) == 6 and "20" in cal["note"]

        import time
        time.sleep(1.1)
        (tmp_path / "late").mkdir()
        (tmp_path / "late" / "late_commute_study.txt").write_text(
            "A late study of commuting. Working from home reduces commuting emissions by a fifth in our sample. " * 3)
        api.wait(api("POST", "/api/projects/Z/ingest", {"path": str(tmp_path / "late")})[1]["job"]["id"])
        rows = {r["hyp_id"]: r for r in api("GET", "/api/projects/Z/hypotheses")[1]["hypotheses"]}
        assert rows[hids[0]]["docs_added_since"] == 1
        h = api("GET", f"/api/projects/Z/hypotheses/{hids[0]}")[1]
        assert h["docs_added_since"][0]["matching"] and h["docs_added_since_matching"] == 1
        job = api.wait(api("POST", f"/api/projects/Z/hypotheses/{hids[0]}/read-new", {})[1]["job"]["id"])
        assert job["state"] == "done" and job["result"]["passages"] >= 1
        rows = {r["hyp_id"]: r for r in api("GET", "/api/projects/Z/hypotheses")[1]["hypotheses"]}
        assert rows[hids[0]]["docs_added_since"] == 0
        job = api.wait(api("POST", "/api/projects/Z/read-new", {})[1]["job"]["id"])  # all that have new documents
        assert job["state"] == "done" and [d["hyp_id"] for d in job["result"]["hypotheses"]] == [hids[1]]
        assert api("POST", "/api/projects/Z/read-new", {})[0] == 400  # nothing left

        import io
        import openpyxl
        st, x = api("GET", "/api/projects/Z/export.xlsx")
        assert "Map" in openpyxl.load_workbook(io.BytesIO(x)).sheetnames
        st, z = api("GET", f"/api/projects/Z/export.zip?token={api.token}", token=False)
        assert st == 200 and z[:2] == b"PK"
        st, r = api("POST", "/api/projects/import", raw=z, headers={"Content-Type": "application/zip"})
        assert st == 200 and r["project"]["name"] == "Z (2)" and r["project"]["hypotheses"] == 2
        st, r = api("POST", "/api/projects/import?name=Z", raw=z)
        assert st == 409 and "already exists" in r["error"]
        st, r = api("POST", "/api/projects/import?name=Colleague", raw=b"not a zip")
        assert st == 400 and "not a zip" in r["error"]
        docs = api("GET", "/api/projects/Colleague/documents")[0]
        assert docs == 404  # nothing was left behind
        d2 = api("GET", "/api/projects/Z (2)/documents".replace(" ", "%20"))[1]["documents"]
        st, chk = api("GET", f"/api/projects/Z%20(2)/documents/{d2[0]['doc_id']}/file?check=1")
        assert chk["available"] is True
    finally:
        srv.shutdown()


def test_a_body_no_endpoint_reads_does_not_spoil_the_next_request(api):
    """Keep-alive: an unread body must not be taken for the start of the next request."""
    import http.client
    api("POST", "/api/projects", {"name": "K"})
    host, port = api.base.split("//")[1].split(":")
    cx = http.client.HTTPConnection(host, int(port), timeout=10)
    hdr = {"X-Enlens-Token": api.token, "Content-Type": "application/json"}
    for _ in range(2):
        cx.request("GET", "/api/projects", body=b'{"unused": true}', headers=hdr)  # 200, keeps the connection
        assert cx.getresponse().read()
        cx.request("GET", "/api/state", headers=hdr)
        r = cx.getresponse()
        assert r.status == 200 and json.loads(r.read())["version"]
    cx.close()


def test_the_icon_starts_one_copy_without_a_terminal_and_quit_stops_it(home, tmp_path):
    import plistlib
    import socket
    import subprocess
    import sys
    from enlens_claims.gui import launcher

    # the icons: a macOS app, a Linux menu entry, a Windows command file when PowerShell is not there
    mac = launcher.make_launcher(str(tmp_path / "data"), True, [tmp_path / "mac"], python=sys.executable, platform="darwin")[0]
    info = plistlib.loads((mac / "Contents" / "Info.plist").read_bytes())
    script = (mac / "Contents" / "MacOS" / "ENLENS").read_text()
    assert info["CFBundleExecutable"] == "ENLENS" and os.access(mac / "Contents" / "MacOS" / "ENLENS", os.X_OK)
    assert "--launcher" in script and "--local-only" in script and str(tmp_path / "data") in script and "&" in script
    lin = launcher.make_launcher(None, False, [tmp_path / "lin"], python=sys.executable, platform="linux")[0]
    entry = lin.read_text()
    assert "Terminal=false" in entry and "enlens_claims.gui.launch --launcher" in entry and sys.executable in entry
    win = launcher.make_launcher(None, False, [tmp_path / "win"], python=sys.executable, platform="win32")[0]
    assert win.suffix in (".lnk", ".cmd") and (win.suffix == ".lnk" or "--launcher" in win.read_text())
    assert launcher.draw_icon(64) is None or launcher.draw_icon(64).size == (64, 64)

    with socket.socket() as sk:
        sk.bind(("127.0.0.1", 0))
        port = sk.getsockname()[1]
    env = {**os.environ, "ENLENS_HOME": str(home.home())}
    cmd = [sys.executable, "-m", "enlens_claims.gui.launch", "--launcher", "--headless", "--no-start-ollama",
           "--port", str(port)]
    first = subprocess.Popen(cmd, env=env, stdin=subprocess.DEVNULL, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    try:
        for _ in range(100):
            if launcher.running():
                break
            time.sleep(0.1)
        live = launcher.running()
        assert live and live["port"] == port and live["pid"] == first.pid
        again = subprocess.run(cmd, env=env, capture_output=True, text=True, timeout=30)  # a second double-click
        assert again.returncode == 0 and "already running" in again.stdout
        page = urllib.request.urlopen(live["url"], timeout=5).read().decode()
        token = page.split('name="enlens-token" content="')[1].split('"')[0]
        a = Api(live["url"].rstrip("/"), token)
        st, s = a("GET", "/api/state")
        assert s["launcher"] is True and s["can_quit"] is True
        assert a("POST", "/api/quit", {})[1]["stopping"] is True
        assert first.wait(timeout=15) == 0
        assert not (home.home() / "app.json").exists() and launcher.running() is None
        log = (home.home() / "logs" / "enlens.log").read_text()
        assert "started by the launcher" in log and "ENLENS stopped." in log
    finally:
        if first.poll() is None:
            first.kill()


def test_reasoning_only_when_judging_and_within_a_budget(home, tmp_path):
    """0.6.1: a model that would reason for ever is stopped at its budget and answers from its notes; with
    "judging", it reasons only when it judges excerpts."""
    from enlens_claims.llm import JUDGING

    srv, url = fake_servers.start("ollama")
    try:
        be = OllamaBackend(base_url=url, think=True, think_budget=300, num_ctx=8192)
        out = be.complete("expand", "You prepare a search", "CLAIM:\nwater stress", fake_servers.THINKER, {})
        assert json.loads(out)["terms"] and be.last_stats["budget_reached"] is True
        chats = [r for p, r in srv.requests if p == "/api/chat"]
        assert chats[0]["think"] is True and chats[0]["options"]["num_predict"] == 400 + 300 + 64
        assert chats[1]["think"] is False and "YOUR REASONING SO FAR" in chats[1]["messages"][1]["content"]
        assert be.last_stats["thinking_tokens"] == 300

        srv.requests.clear()
        be = OllamaBackend(base_url=url, think=True, think_kinds=JUDGING, think_budget=200, num_ctx=12288)
        from enlens_claims import claimcheck, ingest
        from enlens_claims.store import Workspace
        from helpers import PAPER_A_PAGES, make_pdf
        ws = Workspace.create(tmp_path / "w")
        make_pdf(tmp_path / "a.pdf", PAPER_A_PAGES)
        ingest.ingest(ws, [tmp_path / "a.pdf"])
        hid = claimcheck.check_hypothesis(ws, CLAIM, be, fake_servers.THINKER, k=2)
        thinks = {("expand" if "prepare a search" in r["messages"][0]["content"] else "check"): r.get("think")
                  for p, r in srv.requests if p == "/api/chat" and "YOUR REASONING" not in r["messages"][1]["content"]}
        assert thinks == {"expand": False, "check": True}
        assert claimcheck.hypothesis_result(ws, hid)["parts"][0]["round_status"] == "done"
    finally:
        srv.shutdown()

    s = dict(home.DEFAULTS)
    assert home.context_plan({**s, "think": "judging"})["num_ctx"] == 12288   # room to reason, and to answer from notes
    assert home.context_plan({**s, "think": "off"})["num_ctx"] == 8192
    assert home.update_settings({"think": "judging", "think_budget": 1000})["think"] == "judging"
    with pytest.raises(ValueError, match="think_budget"):
        home.update_settings({"think_budget": 50})
    be, _m = home.make_backend(dict(home.DEFAULTS, backend="ollama", ollama_url="http://127.0.0.1:9", ollama_model="m",
                                    think="judging"))
    assert be.think_kinds == JUDGING and be.think_budget == 1500 and be.think is True
